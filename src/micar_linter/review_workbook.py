"""Build a reviewer's labelling workbook (.xlsx) from their blinded review sheet, and read it back.

``micar-review-workbook`` writes one workbook for one reviewer from ``human-review-<id>.csv``:

- sheet "Instructions": how to label, pointing to ``codebook.md`` and ``review-protocol.md``;
- sheet "Review": one row per cell in the sheet's recorded order, so workbook row N is line N of
  the CSV. The cell-identity and requirement columns are locked. ``human_label`` has a dropdown
  restricted to ``review_validation.LABEL_ORDER``; ``reviewer_id`` accepts only this reviewer's
  id; ``reviewed_at`` accepts only a date from the day the findings were generated to today, shown
  as YYYY-MM-DD; ``source_pinpoint`` and ``review_note`` are free text. A locked ``entry_check``
  column shows, once the spreadsheet application computes it, what a row still needs;
- a very hidden sheet "_meta": the reviewer, the recorded sheet hash, the findings hash, the seed
  and a digest of the cell identities, which ``micar-review-merge`` checks on import.

The workbook is built from the blinded sheet only, so it carries no detector outcome. Review values
already in the CSV are the reviewer's own and are copied as they are; the generator never fills a
label. Sheet and workbook-structure protection have no password: they stop accidental edits, not a
determined user, so ``micar-review-merge import`` re-checks cell identity, row order and every value.

Do not commit the workbook; ``.gitignore`` excludes ``studies/**/*.xlsx``. The commands need the
optional ``review`` extra (openpyxl): ``uv run --extra review micar-review-workbook ...``.
"""

from __future__ import annotations

import argparse
import hashlib
import sys
import zipfile
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from micar_linter.review_merge import (
    EXIT_IDENTITY,
    EXIT_OK,
    EXIT_OVERWRITE,
    EXIT_USAGE,
    MergeRefused,
    Study,
    csv_rows,
    load_study,
    read_sheet,
)
from micar_linter.review_sheets import (
    BLINDED_FIELDS,
    CELL_FIELDS,
    REVIEW_FIELDS,
    render_csv,
    sheet_filename,
    structure_problems,
)
from micar_linter.review_validation import LABEL_ORDER, LABEL_REQUIRED_FIELDS, PENDING_LABEL

WORKBOOK_SCHEMA = "micar-study-review-workbook.v1"
INSTRUCTIONS_SHEET = "Instructions"
REVIEW_SHEET = "Review"
META_SHEET = "_meta"
CHECK_COLUMN = "entry_check"
WORKBOOK_COLUMNS: tuple[str, ...] = (*BLINDED_FIELDS, CHECK_COLUMN)
COLUMN_WIDTHS: dict[str, int] = {
    "study_doc_id": 11,
    "document_hash_sha256": 20,
    "rule_id": 30,
    "annex_item": 22,
    "rule_label": 28,
    "requirement": 60,
    "human_label": 19,
    "reviewer_id": 11,
    "reviewed_at": 13,
    "source_pinpoint": 36,
    "review_note": 48,
    CHECK_COLUMN: 26,
}
WRAPPED_COLUMNS = {"rule_label", "requirement", "source_pinpoint", "review_note"}
META_KEYS: tuple[str, ...] = (
    "schema_version",
    "reviewer_id",
    "sheet_file",
    "record_sheet_sha256",
    "findings_sha256",
    "seed",
    "rows",
    "identity_sha256",
    "sheet_sha256_at_generation",
)
# Every meta key except the CSV hash at generation must match the study on import. That hash only
# documents which CSV the workbook came from; the CSV may have changed since (see the import's
# overwrite check).
CHECKED_META_KEYS = META_KEYS[:-1]

_HEADER_FILL = "D9D9D9"
_EDITABLE_FILL = "FFF2CC"


def identity_digest(rows: list[dict[str, str]]) -> str:
    """SHA-256 over the cell-identity columns of the rows, in order."""
    lines = ("\x1f".join(row.get(name, "") for name in CELL_FIELDS) for row in rows)
    return hashlib.sha256("\n".join(lines).encode("utf-8")).hexdigest()


def workbook_meta(study: Study, reviewer_id: str, sheet_text: str) -> dict[str, str]:
    """The identifying values stored in the workbook's hidden _meta sheet."""
    expected_rows = csv_rows(study.expected_sheet(reviewer_id))
    record_sha = hashlib.sha256(study.expected_sheet(reviewer_id).encode("utf-8")).hexdigest()
    return {
        "schema_version": WORKBOOK_SCHEMA,
        "reviewer_id": reviewer_id,
        "sheet_file": sheet_filename(reviewer_id),
        "record_sheet_sha256": record_sha,
        "findings_sha256": study.findings_sha256,
        "seed": str(study.seed),
        "rows": str(len(expected_rows)),
        "identity_sha256": identity_digest(expected_rows),
        "sheet_sha256_at_generation": hashlib.sha256(sheet_text.encode("utf-8")).hexdigest(),
    }


def label_prompt() -> str:
    """Short input prompt for the label column, generated from LABEL_REQUIRED_FIELDS."""
    groups: dict[tuple[str, ...], list[str]] = {}
    for label, fields in LABEL_REQUIRED_FIELDS.items():
        groups.setdefault(fields, []).append(label)
    parts = [f"{', '.join(labels)}: {' and '.join(fields)} required." for fields, labels in groups.items()]
    return " ".join([*parts, "codebook.md governs."])


def build_workbook(study: Study, reviewer_id: str) -> Any:
    """Return an openpyxl Workbook for the reviewer's sheet, after checking the sheet's identity."""
    openpyxl = _openpyxl()
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.protection import WorkbookProtection
    from openpyxl.worksheet.datavalidation import DataValidation

    sheet_text, rows = read_sheet(study, reviewer_id)
    meta = workbook_meta(study, reviewer_id, sheet_text)
    letters = {name: get_column_letter(index) for index, name in enumerate(WORKBOOK_COLUMNS, start=1)}
    last_row = len(rows) + 1

    workbook = openpyxl.Workbook()
    instructions = workbook.active
    instructions.title = INSTRUCTIONS_SHEET
    review = workbook.create_sheet(REVIEW_SHEET)
    meta_sheet = workbook.create_sheet(META_SHEET)

    # Instructions.
    instructions.column_dimensions["A"].width = 110
    for row_number, line in enumerate(_instruction_lines(study, reviewer_id), start=1):
        cell = instructions.cell(row=row_number, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if row_number == 1:
            cell.font = Font(bold=True, size=13)
    instructions.protection.sheet = True

    # Review sheet: header.
    header_fill = PatternFill("solid", fgColor=_HEADER_FILL)
    editable_fill = PatternFill("solid", fgColor=_EDITABLE_FILL)
    for index, name in enumerate(WORKBOOK_COLUMNS, start=1):
        cell = review.cell(row=1, column=index, value=name)
        cell.font = Font(bold=True)
        cell.fill = editable_fill if name in REVIEW_FIELDS else header_fill
        cell.alignment = Alignment(vertical="top")
        review.column_dimensions[letters[name]].width = COLUMN_WIDTHS[name]

    # Review sheet: one row per cell, in the CSV's order.
    unlocked = Protection(locked=False)
    for row_number, row in enumerate(rows, start=2):
        for index, name in enumerate(WORKBOOK_COLUMNS, start=1):
            if name == CHECK_COLUMN:
                value: Any = _check_formula(row_number, letters, reviewer_id)
            elif name == "reviewed_at":
                value = _date_or_text(row.get(name, ""))
            else:
                value = row.get(name, "") or None
            cell = review.cell(row=row_number, column=index, value=value)
            if name != CHECK_COLUMN and isinstance(value, str) and value.startswith("="):
                cell.data_type = "s"  # text from the CSV, never a formula
            cell.alignment = Alignment(wrap_text=name in WRAPPED_COLUMNS, vertical="top")
            if name in REVIEW_FIELDS:
                cell.protection = unlocked
                cell.fill = editable_fill
                cell.number_format = "yyyy-mm-dd" if name == "reviewed_at" else "@"

    # Validation for the editable columns.
    labels = DataValidation(
        type="list",
        formula1='"' + ",".join(LABEL_ORDER) + '"',
        allow_blank=True,
        showErrorMessage=True,
        errorStyle="stop",
        errorTitle="Label",
        error="Choose one of: " + ", ".join(LABEL_ORDER) + ".",
        showInputMessage=True,
        promptTitle="Label",
        prompt=label_prompt(),
    )
    reviewer = DataValidation(
        type="list",
        formula1=f'"{reviewer_id}"',
        allow_blank=True,
        showErrorMessage=True,
        errorStyle="stop",
        errorTitle="Reviewer",
        error=f"This workbook belongs to reviewer {reviewer_id}.",
    )
    earliest = max(study.earliest_review_date, date(1900, 1, 1))  # spreadsheet dates start in 1900
    review_date = DataValidation(
        type="date",
        operator="between",
        formula1=f"DATE({earliest.year},{earliest.month},{earliest.day})",
        formula2="TODAY()",
        allow_blank=True,
        showErrorMessage=True,
        errorStyle="stop",
        errorTitle="Review date",
        error=f"Enter the review date as YYYY-MM-DD, not before {earliest.isoformat()} and not after today.",
        showInputMessage=True,
        promptTitle="Review date",
        prompt="ISO 8601 date of review, YYYY-MM-DD.",
    )
    for validation, column in (
        (labels, "human_label"),
        (reviewer, "reviewer_id"),
        (review_date, "reviewed_at"),
    ):
        validation.add(f"{letters[column]}2:{letters[column]}{last_row}")
        review.add_data_validation(validation)

    review.freeze_panes = "B2"
    review.auto_filter.ref = f"A1:{letters[CHECK_COLUMN]}{last_row}"
    protection = review.protection
    protection.sheet = True
    protection.formatColumns = False
    protection.formatRows = False
    protection.autoFilter = False

    # Meta.
    for row_number, key in enumerate(META_KEYS, start=1):
        meta_sheet.cell(row=row_number, column=1, value=key)
        meta_sheet.cell(row=row_number, column=2, value=meta[key])
    meta_sheet.sheet_state = "veryHidden"
    meta_sheet.protection.sheet = True

    workbook.active = 0
    workbook.security = WorkbookProtection(lockStructure=True)
    workbook.properties.creator = "micar-review-workbook"
    workbook.properties.title = f"Review workbook for {sheet_filename(reviewer_id)}"
    return workbook


def write_workbook(study: Study, reviewer_id: str, out: Path, *, force: bool = False) -> int:
    """Write the workbook; return the number of rows. Refuses to overwrite unless ``force``."""
    if out.exists() and not force:
        raise MergeRefused(
            EXIT_OVERWRITE,
            f"{out} exists and may hold labels; nothing was written. Pass --force to overwrite it",
        )
    workbook = build_workbook(study, reviewer_id)
    out.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out)
    return workbook[REVIEW_SHEET].max_row - 1


def read_workbook(path: Path, study: Study, reviewer_id: str) -> list[dict[str, str]]:
    """Return the workbook's rows as CSV-style strings after checking its identity and row order.

    Raises ``MergeRefused`` (exit 3) if the workbook was generated for another reviewer, sheet or
    findings file, or if any cell-identity value or the row order differs from the generated sheet.
    """
    openpyxl = _openpyxl()
    from openpyxl.utils.exceptions import InvalidFileException

    expected = study.expected_sheet(reviewer_id)
    try:
        workbook = openpyxl.load_workbook(path, data_only=True)
    except (OSError, KeyError, ValueError, zipfile.BadZipFile, InvalidFileException) as exc:
        raise MergeRefused(EXIT_USAGE, f"cannot read the workbook {path}: {exc}") from exc
    name = path.name
    if REVIEW_SHEET not in workbook.sheetnames or META_SHEET not in workbook.sheetnames:
        raise MergeRefused(
            EXIT_IDENTITY,
            f"{name}: not a review workbook from micar-review-workbook (sheets {workbook.sheetnames})",
        )

    sheet_text, _ = read_sheet(study, reviewer_id)
    expected_meta = workbook_meta(study, reviewer_id, sheet_text)
    actual_meta = {
        _cell_text(key): _cell_text(value)
        for key, value, *_ in workbook[META_SHEET].iter_rows(min_col=1, max_col=2, values_only=True)
        if key is not None
    }
    problems = [
        f"{name}: _meta {key} is {actual_meta.get(key)!r}, expected {expected_meta[key]!r}"
        for key in CHECKED_META_KEYS
        if actual_meta.get(key) != expected_meta[key]
    ]

    values = [list(row) for row in workbook[REVIEW_SHEET].iter_rows(values_only=True)]
    header = [_cell_text(value) for value in (values[0] if values else [])]
    while header and not header[-1]:
        header.pop()
    if header != list(WORKBOOK_COLUMNS):
        problems.append(f"{name}: header is {header}, expected {list(WORKBOOK_COLUMNS)}")
        raise MergeRefused(
            EXIT_IDENTITY, f"{name}: workbook structure differs from the generated sheet", problems
        )

    data = values[1:]
    while data and all(value is None or _cell_text(value) == "" for value in data[-1]):
        data.pop()
    width = len(WORKBOOK_COLUMNS)
    rows = []
    for number, raw in enumerate(data, start=2):
        padded = [*raw, *([None] * (width - len(raw)))]
        if any(_cell_text(value) for value in padded[width:]):
            problems.append(f"{name}: row {number} has values beyond the {CHECK_COLUMN} column")
        rows.append(
            {column: _cell_text(value) for column, value in zip(WORKBOOK_COLUMNS, padded, strict=False)}
        )
    problems.extend(
        f"{name}: {problem}"
        for problem in structure_problems(expected, render_csv(rows, BLINDED_FIELDS), validate_review=False)
    )
    if problems:
        raise MergeRefused(
            EXIT_IDENTITY, f"{name}: cell identity or row order differs from the generated sheet", problems
        )
    return [{column: row[column] for column in BLINDED_FIELDS} for row in rows]


def _instruction_lines(study: Study, reviewer_id: str) -> list[str]:
    sheet = sheet_filename(reviewer_id)
    study_dir = study.directory.as_posix()
    labels = [
        f"    {label}: {' and '.join(LABEL_REQUIRED_FIELDS[label])} required."
        for label in LABEL_ORDER
        if label != PENDING_LABEL
    ]
    return [
        f"Review workbook for reviewer {reviewer_id}",
        f"This workbook holds the cells of {sheet}, one row per cell, in that sheet's recorded order. "
        "It contains no detector outcome.",
        "Label each cell under codebook.md and review-protocol.md in the study directory. The codebook "
        "governs; this sheet only summarises it.",
        "While labelling, do not open human-review-matrix.csv, findings-pseudonymous.json, "
        "findings-pseudonymous.csv, findings-summary.md or playground output for these documents.",
        "On the Review sheet, edit only the yellow columns: human_label, reviewer_id, reviewed_at, "
        "source_pinpoint and review_note. The other columns and the row order are locked. Do not sort, "
        "insert or delete rows; filtering is fine.",
        "Labels (choose from the dropdown; a blank label counts as pending):",
        *labels,
        f"    {PENDING_LABEL}: not yet reviewed. The import refuses while any cell is pending.",
        f"reviewer_id: {reviewer_id}. reviewed_at: the ISO 8601 date of review, YYYY-MM-DD.",
        "The entry_check column is a hint computed by the spreadsheet application. The import's "
        "validation decides.",
        "Sheet protection has no password. It guards against accidental edits only; leave it on.",
        "When every cell is labelled, import the workbook into the CSV sheet:",
        f"    uv run --extra review micar-review-merge import --study-dir {study_dir} "
        f"--reviewer {reviewer_id} --workbook <this file>",
        f"Check progress without writing: uv run --extra review micar-review-merge check --study-dir "
        f"{study_dir} --reviewer {reviewer_id} --workbook <this file>",
        "Do not commit this workbook. Commit the labelled CSV sheet on its own.",
    ]


def _check_formula(row_number: int, letters: dict[str, str], reviewer_id: str) -> str:
    label = f"{letters['human_label']}{row_number}"

    def cell(name: str) -> str:
        return f"{letters[name]}{row_number}"

    def any_of(labels: list[str]) -> str:
        return "OR(" + ",".join(f'{label}="{value}"' for value in labels) + ")"

    tests = [
        (f'OR(TRIM({label})="",{label}="{PENDING_LABEL}")', PENDING_LABEL),
        (f"NOT({any_of(list(LABEL_ORDER))})", "label not allowed"),
        (f'TRIM({cell("reviewer_id")})<>"{reviewer_id}"', f"reviewer_id must be {reviewer_id}"),
        (f'{cell("reviewed_at")}=""', "reviewed_at missing"),
    ]
    groups: dict[str, list[str]] = {}
    for value, fields in LABEL_REQUIRED_FIELDS.items():
        for field_name in fields:
            groups.setdefault(field_name, []).append(value)
    for field_name, values in groups.items():
        tests.append((f'AND({any_of(values)},TRIM({cell(field_name)})="")', f"{field_name} required"))
    formula = '"ok"'
    for condition, message in reversed(tests):
        formula = f'IF({condition},"{message}",{formula})'
    return "=" + formula


def _date_or_text(value: str) -> Any:
    value = value.strip()
    if not value:
        return None
    try:
        return date.fromisoformat(value) if len(value) == 10 else value
    except ValueError:
        return value


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _openpyxl() -> Any:
    try:
        import openpyxl
    except ModuleNotFoundError as exc:
        raise MergeRefused(
            EXIT_USAGE,
            "openpyxl is not installed; workbook commands need the review extra, for example "
            "`uv run --extra review micar-review-workbook ...` or "
            "`uv run --extra review micar-review-merge import ...`",
        ) from exc
    return openpyxl


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="micar-review-workbook",
        description=(
            "Write a reviewer's labelling workbook (.xlsx) from their blinded review sheet. The "
            "workbook carries no detector outcome and never pre-fills a label."
        ),
    )
    parser.add_argument("--study-dir", required=True, type=Path, help="study directory with the sheets")
    parser.add_argument("--reviewer", required=True, help="reviewer id, for example r1")
    parser.add_argument(
        "--out",
        type=Path,
        default=None,
        help="workbook to write (default: <study-dir>/human-review-<reviewer>.xlsx, which git ignores)",
    )
    parser.add_argument(
        "--findings",
        type=Path,
        default=None,
        help="findings file (default: the one named in human-review-sheets.json)",
    )
    parser.add_argument(
        "--force", action="store_true", help="overwrite an existing workbook (its labels are lost)"
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        study = load_study(args.study_dir, args.findings)
        out = args.out or args.study_dir / f"human-review-{args.reviewer}.xlsx"
        study.expected_sheet(args.reviewer)
        rows = write_workbook(study, args.reviewer, out, force=args.force)
    except MergeRefused as exc:
        print(f"micar-review-workbook: {exc}", file=sys.stderr)
        for line in exc.details:
            print(f"  {line}", file=sys.stderr)
        return exc.exit_code
    except OSError as exc:
        print(f"micar-review-workbook: {exc}", file=sys.stderr)
        return EXIT_USAGE
    print(
        f"Wrote {out}: {rows} cells for reviewer {args.reviewer} in the order of "
        f"{sheet_filename(args.reviewer)}. Do not commit the workbook.",
        file=sys.stderr,
    )
    return EXIT_OK


if __name__ == "__main__":
    raise SystemExit(main())
