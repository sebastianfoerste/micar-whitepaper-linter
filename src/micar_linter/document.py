"""PDF and DOCX document ingestion and section parsing."""

from __future__ import annotations

import math
import re
import statistics
from collections import Counter
from collections.abc import Callable, Iterable
from dataclasses import dataclass, replace
from functools import partial
from itertools import pairwise
from pathlib import Path

from micar_linter.placeholders import is_not_applicable, is_placeholder
from micar_linter.rules.annex_i import NOT_APPROVED_STATEMENT_EN

# Common patterns (in English and German) matching whitepaper section headings.
# Word boundaries are used to avoid false positives.
SECTION_PATTERNS = {
    # The Art. 6(3) / Art. 51(3) statement, labelled as in Implementing Regulation (EU) 2024/2984.
    "not_approved_statement": [
        r"\bstatement\s+in\s+accordance\s+with\s+art(?:icle|\.)\s*(?:6|51)\s*\(3\)",
        r"\berklärung\s+gemäß\s+art(?:ikel|\.)\s*(?:6|51)\s+abs(?:atz|\.)\s*3\b",
    ],
    "summary": [
        r"\bsummary\b",
        r"\bzusammenfassung\w*",
        r"\bplain-language\s+summary\b",
        r"\bkey\s+information\b",
    ],
    "risk_warning": [r"\brisk\s+warning\b", r"\bwarnhinweis\w*", r"\bmandatory\s+warning\b"],
    "management_statement": [
        r"\bmanagement\s+body\b",
        r"\bleitungsorgan\w*",
        r"\bmanagement\s+statement\b",
        r"\bstatement\b.*\bmanagement\b",
        # Field 03 of the summary of the Implementing Regulation (EU) 2024/2984 template, "Compliance
        # statement in accordance with Article 6(6)", holds the statement of the management body.
        r"\bcompliance\s+statement\b",
    ],
    "notification_date": [
        r"\bdate\s+of\s+notification\b",
        r"\bnotifizierungsdatum\w*",
        r"\bnotification\s+date\b",
        # The German texts: Art. 6 Abs. 8, Art. 19 Abs. 7 and Art. 51 Abs. 7 MiCAR ask for "das Datum seiner
        # Übermittlung", and field 01 (I.01) of the Implementing Regulation (EU) 2024/2984 template is
        # "Datum der Mitteilung". The compounds "Übermittlungsdatum" and "Mitteilungsdatum" name the same
        # date. Other fields of the template, such as "Art der Übermittlung" and "Methoden der Übermittlung
        # des Werts", do not open this section.
        r"\bdatum\s+der\s+(?:übermittlung|mitteilung|notifizierung)\b",
        r"\b(?:übermittlungs|mitteilungs)datum\w*",
    ],
    # "Plain-language summary" heads the summary; it would otherwise feed the summary to the language check.
    "language": [r"(?<!plain[- ])\blanguage\b", r"\bsprache\w*"],
    "offeror": [r"\bofferor\b", r"\banbieter\w*"],
    "issuer": [r"\bissuer\b", r"\bemittent\w*"],
    "trading_platform_operator": [
        r"\btrading\s+platform\b",
        r"\bhandelsplattform\w*",
        r"\boperator\b.*\btrading\b",
    ],
    "project": [r"\bproject\b", r"\bprojekt\w*", r"\bbusiness\s+model\b", r"\bteam\b"],
    "offer_or_admission": [
        r"\bpublic\s+offer\b",
        r"\bzulassung\w*",
        r"\badmission\s+to\s+trading\b",
        r"\boffer\b.*\badmission\b",
    ],
    "crypto_asset": [
        r"\bcrypto-asset\b",
        # The title of Part F, "Information about the crypto-assets", but not the plural in the titles of
        # Part E ("offer to the public of crypto-assets") and Part G ("attached to the crypto-assets").
        r"\b(?:about|on)\s+the\s+crypto-assets\b(?!-)",
        r"\bcrypto\s+asset\b",
        r"\bcharacteristics\b.*\bcrypto\b",
        r"\binformation\s+about\s+the\s+token\b",
        r"\bthe\s+token\b",
    ],
    "art": [
        r"\basset-referenced\s+token\b",
        r"\bstabilisation\s+mechanism\b",
        r"\breference\s+assets\b",
        r"\binformation\s+about\s+the\s+token\b",
        r"\bthe\s+token\b",
    ],
    "emt": [
        r"\be-money\s+token\b",
        r"\bdenomination\b",
        r"\breferenced\s+currency\b",
        r"\binformation\s+about\s+the\s+token\b",
        r"\bthe\s+token\b",
    ],
    "rights_and_obligations": [
        r"\brights\s+and\s+obligations\b",
        r"\brechte\s+und\s+pflichten\b",
        r"\bholder\s+rights\b",
    ],
    "technology": [
        r"\btechnology\b",
        r"\btechnologie\w*",
        r"\bunderlying\s+technology\b",
        r"\bdlt\b",
        r"\bconsensus\b",
    ],
    "risks": [r"\brisk\s+factors\b", r"\brisikofaktor\w*", r"\brisks\b", r"\brisiken\b"],
    "environmental_impact": [
        r"\benvironmental\b",
        r"\bclimate\b",
        r"\bsustainability\b",
        r"\bumwelt\w*",
        r"\bklima\w*",
    ],
    "reserve_of_assets": [r"\breserve\b.*\bassets\b", r"\breservevermögen\w*", r"\basset\s+reserve\b"],
    "redemption": [r"\bredemption\b", r"\brücknahme\w*", r"\brücktausch\w*", r"\bredemption\s+right\b"],
    "safeguarding": [r"\bsafeguarding\b", r"\babsicherung\w*", r"\bsicherung\w*.*\bgelder\w*"],
}


# Art. 6(3) and Art. 51(3) MiCAR put the statement on the first page without a heading of
# its own, so a paragraph that opens with it is routed by its content. The German opening
# is a routing heuristic only; the rule does not treat it as verified statutory wording.
_STATEMENT_START = re.compile(
    r"^\W*(?:this\s+crypto-asset\s+white\s+paper\s+has\s+not\s+been\s+approved"
    r"|dieses\s+kryptowerte-whitepaper\s+wurde\s+von\s+keiner)\b",
    re.IGNORECASE,
)
# The statement has two sentences: the one opening above, then the one naming who is solely
# responsible for the content. Each is about 20 words in Art. 6(3) and Art. 51(3).
_RESPONSIBILITY = re.compile(r"\b(?:responsib|verantwort)\w*", re.IGNORECASE)
# The responsibility sentence as Art. 6(3) and Art. 51(3) word it: the person it names, then the clause
# "is solely responsible for the content of this crypto-asset white paper". The German pair is a
# routing heuristic, like the German opening.
_RESPONSIBILITY_CLAUSE = re.compile(
    r"\bsolely\s+responsible\s+for\s+the\s+content\b|\balleinige\s+verantwortung\s+für\s+den\s+inhalt\b",
    re.IGNORECASE,
)
_RESPONSIBLE_PERSON = re.compile(
    r"\bthe\s+(?:offeror|issuer|person|operator)\b|\b(?:der|die)\s+(?:anbieter|emittent|betreiber|person)\b",
    re.IGNORECASE,
)
# The words the statutory sentence puts between the person and the clause: "of the crypto-asset is",
# "seeking admission to trading of the crypto-asset is", "of the e-money token is", "des Kryptowerts
# trägt die". A person named elsewhere, as in a part title above a sentence with a subject of its own
# ("Example AG is solely responsible ..."), is not the one the clause makes responsible.
_PERSON_TO_CLAUSE = re.compile(
    r"(?:\s+(?:seeking\s+admission\s+to\s+trading|of\s+the\s+trading\s+platform))?"
    r"(?:\s+of\s+the\s+[\w-]+(?:(?<=-)\s+[\w-]+)?(?:\s+tokens?)?)?\s+is\s+"
    r"|(?:\s+(?:des|der)\s+[\w-]+(?:(?<=-)\s+[\w-]+)?)?"
    r"(?:,\s+die\s+die\s+zulassung\s+zum\s+handel\s+beantragt,)?\s+trägt\s+die\s+",
    re.IGNORECASE,
)
# The words that end the clause, "of this crypto-asset white paper" ("dieses Kryptowerte-Whitepapers"), with
# any stop, quote or note mark attached, and however a line break or a hyphen at a line end splits them.
_CLAUSE_WORDS = re.compile(
    r"\s+(?:of\s+this\s+crypto-\s*asset\s+white\s+paper|dieses\s+kryptowerte-\s*whitepapers?)"
    # A letter may not follow ("papers"); a note mark stuck to the word may ("paper2", "paper\u00b2").
    r"(?![^\W\d_\u00b9\u00b2\u00b3\u2070-\u2079])\S*",
    re.IGNORECASE,
)
# Where the clause ends in other words, the rest of its sentence up to the stop.
_CLAUSE_TAIL_MAX_WORDS = 8
# A full stop may sit inside closing quotes, brackets or Markdown emphasis, or before a note mark:
# 'paper.”', 'Whitepapers.“', 'paper.**', 'paper.[^1]', 'paper.¹', and 'Union.1' in extracted PDF text.
# A note mark follows a word, never a field or outline number such as "A.1", "I.03" or "IV.2".
_STATEMENT_SENTENCE_END = re.compile(
    r"[.!?](?:[\"'\u2018\u2019\u201c\u201d\u00ab\u00bb)\]*_]|\[\^\w+\]|[\u00b9\u00b2\u00b3\u2070-\u2079])*"
    r"(?:(?<=[^\W\d_]{2}[.!?])(?<!\b[IVX]{2}[.!?])(?<!\b[IVX]{3}[.!?])(?<!\b[IVX]{4}[.!?])\d{1,2})?(?=\s|$)"
)
_STATEMENT_MAX_LINES = 12
_SENTENCE_MAX_WORDS = 30
_HEADING_MAX_WORDS = 12


def is_statement_start(text: str) -> bool:
    """True where a line or paragraph opens with the Art. 6(3) / Art. 51(3) statement."""
    return bool(_STATEMENT_START.match(text))


def looks_like_heading(text: str) -> bool:
    """True for a Markdown heading, or a short line without closing punctuation that names a section."""
    stripped = text.strip()
    if stripped.startswith("#"):
        return True
    if not stripped or len(stripped.split()) > _HEADING_MAX_WORDS or stripped[-1] in ".,;:!?":
        return False
    return bool(match_heading_to_section(stripped))


def statement_extent(lines: list[str], start: int, is_heading: Callable[[int], bool] | None = None) -> int:
    """Index after the last line of the statement that opens at lines[start].

    A wrapped statement in a PDF or a hard-wrapped Markdown file spans several lines, and
    its continuation lines often contain words such as "crypto-asset" or "offeror" that the
    heading heuristics would otherwise read as new section headings. The statement ends
    with its opening sentence, or with the next sentence where that one is short and says
    who is responsible. It never runs past a heading, so an incomplete or reworded
    statement cannot swallow the sections that follow it.

    is_heading(index) reports a heading the loader reads from the layout, such as a DOCX
    paragraph in a Heading style or a PDF line in a heading font. A statement may be set
    apart in bold, capitals or a larger font, so the lines of the statutory responsibility
    sentence, from the person it names to the end of the clause, are recognised by their
    words and never read as a heading, whatever their layout or spacing.
    """
    window = lines[start : start + _STATEMENT_MAX_LINES]
    text = " ".join(window)
    line_ends: list[int] = []
    offset = 0
    for line in window:
        line_ends.append(offset + len(line))
        offset += len(line) + 1

    def line_at(position: int) -> int:
        return next(i for i, line_end in enumerate(line_ends) if line_end >= position)

    def marked(i: int) -> bool:
        return window[i].lstrip().startswith("#") or (is_heading is not None and is_heading(start + i))

    own: range = range(0)
    # Whether the responsibility sentence holds every statutory word, so the statement is whole at its end.
    complete = False
    clause = _RESPONSIBILITY_CLAUSE.search(text)
    if clause:
        persons = [
            person
            for person in _RESPONSIBLE_PERSON.finditer(text, 0, clause.start())
            if _PERSON_TO_CLAUSE.fullmatch(text, person.end(), clause.start())
        ]
        own_end = line_at(clause.end())
        words = _CLAUSE_WORDS.match(text, clause.end())
        tail = _STATEMENT_SENTENCE_END.search(text, clause.end())
        if words:
            # The statutory words, however they are set, with or without the stop.
            own_end, complete = line_at(words.end()), True
        elif tail and len(text[clause.end() : tail.end()].split()) <= _CLAUSE_TAIL_MAX_WORDS:
            # The rest of the sentence ("of this crypto-asset white paper."), but never a line that opens a
            # heading or a numbered label: a sentence that has lost its stop ends before it.
            own_end = next(
                (
                    i - 1
                    for i in range(own_end + 1, line_at(tail.end()) + 1)
                    if marked(i) or _NUMBERING.match(window[i].lstrip())
                ),
                line_at(tail.end()),
            )
        own = range(line_at((persons[-1] if persons else clause).start() + 1), own_end + 1)

    def firm_heading(i: int) -> bool:
        return i not in own and marked(i)

    # Continuation lines of the opening sentence never name a section, so a line that ends
    # inside it and looks like a heading means the sentence was cut short. The line where
    # the sentence ends may already start the next one ("Union. The offeror"), so it is
    # not read as a heading.
    first = _STATEMENT_SENTENCE_END.search(text)
    first_end = first.end() if first else len(text) + 1
    for i in range(1, len(window)):
        if line_ends[i] < first_end and i not in own and (looks_like_heading(window[i]) or firm_heading(i)):
            return start + i
    last = line_at(first_end) if first else None
    # A responsibility sentence with every statutory word, right after the opening one, ends the statement
    # with its last word, whatever follows it: no stop, a heading, a label, a note or a run-on.
    if (
        complete
        and (last is None or own.start <= last + 1)
        and not any(firm_heading(i) for i in range(1, own.start))
    ):
        return start + own.stop
    if first is None or last is None:
        return start + 1

    # The responsibility sentence of a narrow PDF column can look like a heading
    # ("The offeror of the crypto-asset is solely"), so only firm headings end it early.
    second = _STATEMENT_SENTENCE_END.search(text, first_end)
    if second is None:
        return start + last + 1
    sentence = text[first_end : second.end()]
    if len(sentence.split()) > _SENTENCE_MAX_WORDS or not _RESPONSIBILITY.search(sentence):
        return start + last + 1
    closing = line_at(second.end())
    heading = next((i for i in range(last + 1, closing + 1) if firm_heading(i)), None)
    if heading is not None:
        # A responsibility sentence that has lost its stop ends with its own lines, before the heading.
        return start + (
            own.stop if own and own.start <= last + 1 and last < own.stop <= heading else last + 1
        )
    return start + closing + 1


def paragraph_statement_extent(
    paragraphs: list[str], start: int, is_heading: Callable[[int], bool] | None = None
) -> int:
    """statement_extent for a loader that reads paragraphs (DOCX, XHTML) rather than wrapped lines.

    An opening paragraph without a full stop was cut off, unless the paragraphs after it carry on the
    English opening sentence word for word, as a statement set a line per paragraph does. So the paragraph
    after a cut-off opening is not taken in, whatever its first sentence.
    """
    end = statement_extent(paragraphs, start, is_heading)
    for index in range(start + 1, end):
        if _STATEMENT_SENTENCE_END.search(" ".join(paragraphs[start:index])):
            break
        joined = re.sub(r"^\W+", "", " ".join(" ".join(paragraphs[start : index + 1]).split()).lower())
        if not (NOT_APPROVED_STATEMENT_EN.startswith(joined) or joined.startswith(NOT_APPROVED_STATEMENT_EN)):
            return index
    return end


def match_heading_to_section(heading: str) -> list[str]:
    """Matches a heading line to a list of known ruleset section keys."""
    normalized = heading.lower()
    matched = []
    for key, patterns in SECTION_PATTERNS.items():
        for pattern in patterns:
            if re.search(pattern, normalized):
                matched.append(key)
                break
    return matched


# Formatting alone does not make a heading: a body sentence that mentions a section word
# ("Note: purchasers bear market, liquidity and technology risks.") must not open that section.
_MAX_HEADING_WORDS = 12
# A label that its numbering or its font marks as a heading may run longer: the part titles of the
# Implementing Regulation (EU) 2024/2984 template have up to 20 words, with or without the part letter
# ("Information about the offer to the public of crypto-assets or their admission to trading").
_MAX_MARKED_HEADING_WORDS = 20
# A question mark may end a heading ("Risks: What Could Go Wrong?"); these marks end a sentence or a phrase.
_SENTENCE_END = re.compile(r"[.!;,]$")
# A word of four or more letters, a full stop and a capitalised word: "Risks. Holders may lose ...".
# "A.1 Name", "5.2 Risks" and "Ledger S.A. Risks" have no such word.
_SENTENCE_BREAK = re.compile(r"([^\W\d_]{4,})[.!?]\s+([^\W\d_])")
# Abbreviations that keep their full stop inside a heading: "5. Risks (incl. Liquidity Risks)",
# "Quillmoor Corp. Risk Factors", "Risiken bzgl. des Tokens".
_ABBREVIATIONS = frozenset(
    {
        "abschn", "abzgl", "approx", "assoc", "buchst", "bzgl", "corp", "dept", "einschl", "evtl", "excl", "incl",
        "inkl", "insb", "intl", "mgmt", "resp", "sect", "unterabs", "ziff", "zzgl",
    }
)  # fmt: skip
# A list item starts with a bullet (• ● ▪ ■ ◦ ○ ‣ ► ▸ ·), an asterisk, a hyphen, an en or em dash, or a
# symbol-font glyph that PDF text extraction maps to the private use area.
_BULLET = re.compile(
    r"^[*\-\u00b7\u2013\u2014\u2022\u2023\u25a0\u25aa\u25b8\u25ba\u25cb\u25cf\u25e6\uf000-\uf0ff]"
)
# The numbering of the Implementing Regulation (EU) 2024/2984 template: "Part I", "Teil I", the field
# numbers "A.1" and "I.01", and the fields 00 to 10 of the summary table ("01 Date of notification").
_TEMPLATE_NUMBERING = (
    r"(?:Part|PART|Teil|TEIL)\s+(?:[A-Z]|[IVX]{1,5}|\d{1,2})[.:)]?|[A-Z]\.\d{1,2}(?:\.\d{1,2})*\.?|0\d|10"
)
# Outline numbering, which headings and list items share: "1.", "5.2", "1)", "I.", "IV.", "A.".
_OUTLINE_NUMBERING = r"\d{1,2}(?:\.\d{1,2})*[.)]|\d{1,2}(?:\.\d{1,2})+|[A-Z]\.|[IVX]{1,5}\."
_NUMBERING = re.compile(rf"^(?:(?P<template>{_TEMPLATE_NUMBERING})|{_OUTLINE_NUMBERING})\s+")
# Words a title-case heading leaves in lower case: "Information about the Issuer", "Rechte und Pflichten".
# None of them ends a sentence or a heading ("... crypto-assets or their" runs on to "admission to trading").
_MINOR_WORD = re.compile(
    r"a|about|against|among|an|and|as|at|between|by|for|from|in|into|its|of|on|or|over|per|than|the|their|to"
    r"|under|via|with|within|without|am|auf|aus|bei|bzw|das|dem|den|der|des|die|ein|eine|einer|eines|für|im"
    r"|mit|nach|oder|sowie|über|und|von|vom|zu|zum|zur"
)
# Dashes and quotation marks around heading words: en dash, em dash, curly single and double quotes.
_DASHES = "-\u2013\u2014"
_QUOTES = "\"'\u2018\u2019\u201c\u201d"


@dataclass(frozen=True)
class DocumentSections:
    """The section texts of a DOCX or PDF draft and, apart from them, the text of its headings and its
    whole extracted text.

    Heading text is not section body, and some text sits in no section (before the first heading, or
    after a line that ends a section). Regime inference reads the whole text: a heading such as
    "Reserve of Assets" may be the only place that names the regime.
    """

    sections: dict[str, str]
    headings: tuple[str, ...]
    text: str = ""


def _words(text: str) -> list[str]:
    """The words of a line; a lone dash ("Part G \u2013 Information ...") is not a word."""
    return [word for word in text.split() if word.strip(_DASHES)]


def _max_words(text: str, marked: bool = False) -> int:
    """The word limit of a heading: 20 when it is numbered or marked (entirely bold or in capitals in a
    DOCX, set in a heading font in a PDF), 12 otherwise."""
    return _MAX_MARKED_HEADING_WORDS if marked or _NUMBERING.match(text) else _MAX_HEADING_WORDS


def _has_sentence_break(text: str) -> bool:
    return any(
        letter.isupper() and word.lower() not in _ABBREVIATIONS
        for word, letter in _SENTENCE_BREAK.findall(text)
    )


def _is_label(text: str, max_words: int = _MAX_HEADING_WORDS) -> bool:
    """A short label rather than a sentence or a list item: few words, no sentence punctuation after a
    word and no leading bullet."""
    return (
        len(_words(text)) <= max_words
        and not _SENTENCE_END.search(text)
        and not _has_sentence_break(text)
        and not _BULLET.match(text)
    )


def _is_title_like(text: str) -> bool:
    """All capitals, or every word other than a short connective or an abbreviation starts with a capital."""
    words = [word.strip(f"()[]{_QUOTES},:;") for word in text.split()]
    words = [word for word in words if word[:1].isalpha()]
    if not words or not words[0][0].isupper():
        return False
    return all(
        not word.islower() or _MINOR_WORD.fullmatch(word) or word.rstrip(".") in _ABBREVIATIONS
        for word in words[1:]
    )


def _ends_sentence(line: str) -> bool:
    return bool(re.search(r"[.!?][\"')\]\u2019\u201d]*$", line))


def _runs_on(line: str) -> bool:
    """The sentence runs on into the next line: the line ends with a comma, a hyphenated word part or a
    connective such as "and", "the" or "die"."""
    if re.search(r"(?:,|\w-)$", line):
        return True
    last = line.split()[-1]
    return len(last) > 1 and (last.islower() or line.isupper()) and bool(_MINOR_WORD.fullmatch(last.lower()))


def _heading_shaped(text: str, marked: bool = False) -> bool:
    """A label within the word limit whose first letter after any numbering is a capital."""
    numbering = _NUMBERING.match(text)
    title = text[numbering.end() :] if numbering else text
    return _is_label(text, _max_words(text, marked)) and next((c for c in title if c.isalpha()), "").isupper()


def _template_kind(text: str) -> str | None:
    """How the template numbers a line: "part" for a part title ("Part F"), "field" for a field of a part
    ("F.1") and "summary" for a field of the summary table ("01"); None for other lines."""
    numbering = _NUMBERING.match(text)
    number = numbering.group("template") if numbering else None
    if number is None:
        return None
    if number[:4].lower() in ("part", "teil"):
        return "part"
    return "field" if "." in number else "summary"


def _outline_number(text: str) -> str | None:
    """The outline number a heading starts with, without its closing mark: "12" in "12. Information on
    risks" and "12.1" in "12.1 Market risk". None for other headings."""
    numbering = _NUMBERING.match(text)
    number = numbering.group().strip().rstrip(".)") if numbering else ""
    return number if re.fullmatch(r"\d{1,2}(?:\.\d{1,2})*", number) else None


def _numbered_below(number: str | None, opener_number: str | None) -> bool:
    """Whether a heading number extends the number of the heading that opened the section, as "12.1"
    and "12.1.3" extend "12"."""
    return number is not None and opener_number is not None and number.startswith(f"{opener_number}.")


# A heading that names one risk ("Market risk", "Smart contract risk", "Marktrisiko") rather than the
# risk section; the risk warning is a section of its own.
_ONE_RISK = re.compile(r"\brisk\b(?!\s+warning)|risiko\b", re.IGNORECASE)


def _sub_heading(text: str, number: str | None, opener_number: str | None, keys: list[str]) -> bool:
    """Whether the number or the words of a heading that names no section place it under the heading
    that opened the current section, whose keys are keys: "12.1" under "12.", or a heading that names
    one risk in the risk section."""
    return _numbered_below(number, opener_number) or ("risks" in keys and bool(_ONE_RISK.search(text)))


def _ends_section(
    kind: str | None,
    level: float | None,
    opener: str | None,
    section_level: float | None,
    *,
    label: bool = False,
    sub_heading: bool = False,
) -> bool:
    """Whether a heading that names no section ends the current section.

    kind is the template numbering of the heading (see _template_kind) and opener that of the heading
    that opened the current section. level and section_level rank the heading and the heading that
    opened the section: a larger value is a lower level (a smaller PDF font, a deeper DOCX outline level).
    label tells whether only bold type or capitals mark the heading: a DOCX paragraph without a Heading
    style or an outline level, or a PDF line in bold at the size of the body text. sub_heading tells
    whether its number or its words place it under the heading that opened the section (see _sub_heading).

    A part title always ends the section, and a field of a part never does, as it belongs to its part.
    A field of the summary table ends the section of the summary field before it. A label and a heading
    that its number or words place under the heading that opened the section stay in the section, as
    "Market risk" and "12.1 Market risk" do under "12. Information on risks". Any other heading, one with
    a Heading style, an outline level or a larger font, ends the section unless it ranks below the
    heading that opened it.
    """
    if kind == "part":
        return True
    if kind == "field":
        return False
    if kind == "summary":
        return opener == "summary"
    if label or sub_heading:
        return False
    return level is None or section_level is None or level <= section_level


def _opening_level(level: float | None, section_level: float | None, in_section: bool) -> float | None:
    """The level of the section a heading that names a section opens. in_section tells whether the
    heading names a section that is open: such a sub-heading ("Market Risks" under "Risk Factors") keeps
    the level of the heading that opened it. A heading that names another section has its own level, so
    a large cover title that names a section ("Crypto-Asset White Paper") sets no level for the headings
    after it."""
    if in_section and section_level is not None and (level is None or level > section_level):
        return section_level
    return level


def _opening_number(number: str | None, opener_number: str | None) -> str | None:
    """The outline number of the section a heading opens or ends: a heading numbered below the heading
    that opened the current section ("12.2 Technology risks" under "12.") keeps that heading's number."""
    return opener_number if _numbered_below(number, opener_number) else number


def _field_answer(text: str, keys: list[str]) -> str:
    """The answer after the label of a numbered field on the same line, as pypdf reads a row of the
    summary table: "2 February 2026" in "01 Date of notification 2 February 2026"."""
    if _template_kind(text) not in ("field", "summary"):
        return ""
    ends = [
        match.end()
        for key in keys
        for pattern in SECTION_PATTERNS[key]
        if (match := re.search(pattern, text, re.IGNORECASE))
    ]
    return text[max(ends) :].strip() if ends else ""


@dataclass(frozen=True)
class _Line:
    """A non-empty text line and, where the PDF reports them, the font its letters are set in and where
    it sits on its page."""

    text: str
    size: float | None = None  # the size of its smallest letter, in points
    bold: bool = False  # every letter is set in a bold face
    page: int | None = None  # the page it is on, counted from 0
    y: float | None = None  # the height of its (last) baseline above the foot of the page, in points


_HEADING, _BODY, _BREAK = "heading", "body", "break"
# A list item in lower case after a heading: "a) Market risk", "(b) Liquidity risk", "iv. Custody risk".
# German abbreviations of single letters ("z. B.", "u. a.", "d. h.") continue a sentence instead.
_LIST_MARKER = re.compile(r"^\(?(?:[a-z]|[ivx]{1,4})[.)]\s(?![^\W\d_]\.)")


def _line_role(
    line: str,
    previous: str | None,
    previous_role: str | None,
    following: str | None,
    *,
    started: bool,
    styled: bool,
    previous_styled: bool,
    by_font: bool = False,
    paragraph: bool = False,
) -> str:
    """The role of a line that names a section: a heading, body text, or a break that ends the section.

    previous and following are the neighbouring lines (None at either end), previous_role is the role of
    previous, started tells whether a heading or break came before, and styled and previous_styled whether
    the line and the one before are set larger or bolder than the body text. by_font tells whether the
    draft sets its headings in a heading font and the font of this line is known. paragraph tells whether
    the line is the first line of a paragraph set in its font (see _starts_paragraph).
    """
    if not _heading_shaped(line, styled) or paragraph:
        return _BODY
    numbering = _NUMBERING.match(line)
    title = line[numbering.end() :] if numbering else line
    # The template's own part and field labels, sentence case included ("Part I: Information on risks").
    if numbering and numbering.group("template"):
        return _HEADING
    # A heading font marks a heading, whatever the lines around it, except in a numbered list set in that
    # font: after a numbered body line in a heading font, a numbered line is the next list item, or the
    # heading after an empty section, so it ends the section.
    if styled:
        if numbering and previous_role == _BODY and previous_styled and _NUMBERING.match(previous):
            return _BREAK
        return _HEADING
    title_like = _is_title_like(title)
    # A line ending with a colon is a lead-in, not a heading. In sentence case it stays in the current
    # section; in title case it may also be a heading written with a colon ("Risk Factors:"), so it ends
    # the section rather than handing that section to the one before.
    if line.endswith(":"):
        return _BREAK if title_like else _BODY
    # Running text stays in the current section: the line a lead-in introduces, and a line that a
    # sentence runs on into or out of. A list item in lower case ("a) Market risk") may follow a heading,
    # and so may an answer that the item does not apply ("n/a", "entfällt").
    if previous is not None and previous.endswith(":"):
        return _BODY
    continued = (
        following is not None
        and following[:1].islower()
        and not _LIST_MARKER.match(following)
        and not is_not_applicable(following)
    )
    if (previous_role == _BODY and _runs_on(previous)) or continued:
        return _BODY
    # A line that runs on into the next one is no heading, but it may be the first line of a wrapped
    # heading as well as of a paragraph, so it ends the section.
    if _runs_on(line):
        return _BREAK
    # In a capitals paragraph, a line after a long or unfinished capitals line is a wrapped line of it; a
    # long capitals line over more capitals may start one, or be a long title over one, so it ends the section.
    if line.isupper():
        if (
            previous_role == _BODY
            and previous.isupper()
            and not _ends_sentence(previous)
            and (len(_words(previous)) >= 8 or _runs_on(previous))
        ):
            return _BODY
        if following is not None and following.isupper() and len(_words(line)) >= 8:
            # Under a heading it starts the paragraph of that heading.
            return _BREAK if previous_role == _BODY else _BODY
    # After a numbered body line, a numbered line is the next list item. In title case it may also be the
    # heading after an empty section, so it ends the section.
    if numbering and previous_role == _BODY and _NUMBERING.match(previous):
        return _BREAK if title_like else _BODY
    finished = (
        not started
        or previous_role != _BODY
        or _ends_sentence(previous)
        or is_placeholder(previous)
        or is_not_applicable(previous)
    )
    if title_like and (not numbering or finished):
        # Where the draft sets its headings in a heading font, a line in the body font is no heading.
        return _BREAK if by_font else _HEADING
    # Shaped like a heading but in sentence case ("Risk factors"), or numbered after an unfinished line:
    # the text below is credited to no section rather than to the section before.
    return _BREAK


def _font(line: _Line) -> tuple[float, bool] | None:
    """The size, to a tenth of a point, and the weight of a line's font."""
    return None if line.size is None else (round(line.size, 1), line.bold)


def _body_font(lines: list[_Line]) -> tuple[float, bool] | None:
    """The font size and weight that most letters are set in."""
    letters: Counter[tuple[float, bool]] = Counter()
    for line in lines:
        font = _font(line)
        if font is not None:
            letters[font] += sum(char.isalpha() for char in line.text)
    return letters.most_common(1)[0][0] if letters else None


def _larger(line: _Line, body_font: tuple[float, bool] | None) -> bool:
    """Set at least 15 % larger than the body text."""
    return line.size is not None and body_font is not None and line.size >= body_font[0] * 1.15


def _styled(line: _Line, body_font: tuple[float, bool] | None) -> bool:
    """Set at least 15 % larger than the body text, or in bold where the body text is not."""
    if line.size is None or body_font is None:
        return False
    size, bold = body_font
    return _larger(line, body_font) or (line.bold and not bold and line.size >= size * 0.95)


# Wrapped lines of one paragraph sit at most this many times their font size apart. A heading right
# after another is set further down, by the space between paragraphs.
_LINE_SPACING = 1.75


def _joined(first: str, second: str) -> str:
    """Two lines of a wrapped heading as one line: "crypto-" and "assets" make "crypto-assets"."""
    return first + second if re.search(r"[^\W\d_]-$", first) else f"{first} {second}"


def _heights_hold(lines: list[_Line], body_font: tuple[float, bool] | None) -> bool:
    """Whether the heights pypdf reports place the lines where they are printed.

    The lines of a body paragraph sit 1 to 2.5 times their font size apart. pypdf scales the leading of
    the T* operator by the font size, so in a PDF laid out with T* the lines seem many times further apart.
    """
    gaps = [
        first.y - line.y
        for first, line in pairwise(lines)
        if first.page == line.page
        and first.y is not None
        and line.y is not None
        and first.y > line.y
        and _font(first) == _font(line) == body_font
    ]
    return body_font is None or not gaps or statistics.median(gaps) <= 2.5 * body_font[0]


# The first line of a heading that wraps without running on ("... From the Offeror or Person" over
# "Seeking Admission to Trading") fills most of the column: at least this share of a full line of body
# text, in characters scaled by the font sizes.
_FULL_LINE_SHARE = 0.6


def _full_line(lines: list[_Line], body_font: tuple[float, bool] | None) -> float | None:
    """The number of characters on a full line of body text: the upper quartile of the lengths of the
    lines in the body font, as the last line of a paragraph is shorter."""
    lengths = [len(line.text) for line in lines if body_font is not None and _font(line) == body_font]
    return statistics.quantiles(lengths, n=4)[2] if len(lengths) >= 2 else None


def _fills_line(line: _Line, full_line: float | None, body_font: tuple[float, bool] | None) -> bool:
    """Whether a line fills most of the column: at least 60 % of the characters of a full line of body
    text, scaled by the body size over the size of the line, as a larger font fits fewer on a line."""
    if line.size is None or full_line is None or body_font is None:
        return False
    return len(line.text) >= _FULL_LINE_SHARE * full_line * body_font[0] / line.size


def _next_line(first: _Line, line: _Line, heights: bool) -> bool:
    """Whether line can go on with the text of first: on the same page, not numbered, not a list item,
    and set one line spacing below first where the heights hold."""
    if (
        first.page != line.page
        or _NUMBERING.match(line.text)
        or _BULLET.match(line.text)
        or _LIST_MARKER.match(line.text)
    ):
        return False
    return not (
        heights
        and first.y is not None
        and line.y is not None
        and line.size is not None
        and not 0 < first.y - line.y <= _LINE_SPACING * line.size
    )


def _wraps(first: _Line, line: _Line, first_styled: bool, styled: bool, heights: bool, fills: bool) -> bool:
    """Whether line is the next line of the heading or label that first starts.

    A line in the same heading font, set one line spacing below (where the heights hold), is the next
    line of a wrapped heading where first does not end a sentence, the two lines together are still
    shaped like a heading, and first runs on or fills most of the column (fills, see _fills_line). Where
    a draft sets no space between paragraphs, a short heading is one line spacing above the paragraph or
    the heading after it, and these start something new. A line in the body font is the next line only
    where first runs on, line starts in lower case and the two lines together are a label that names a
    section ("08 Characteristics of the crypto-" and "asset"). A numbered line, a list item, and a line on
    another page or further down start something new.
    """
    if not _next_line(first, line, heights):
        return False
    joined = _joined(first.text, line.text)
    if first_styled or styled:
        return (
            first_styled
            and styled
            and _font(first) == _font(line)
            and not _ends_sentence(first.text)
            and _heading_shaped(joined, True)
            and (_runs_on(first.text) or fills)
        )
    return (
        _runs_on(first.text)
        and line.text[:1].islower()
        and _is_label(joined, _max_words(joined))
        and bool(match_heading_to_section(joined))
    )


def _starts_paragraph(line: _Line, following: _Line | None, fills: bool, heights: bool) -> bool:
    """Whether a line is the first line of a paragraph set in its font: it fills most of the column
    (fills, see _fills_line), and the line after it is in the same font, goes on with its text (see
    _next_line) and starts in lower case, as a sentence does that runs on ("... approved by any
    competent" over "authority in any Member State ..."). The lines of a heading that wraps are already
    joined into one (see _wraps), and a short heading over a lower-case line ("Risk Factors" over "and
    Mitigation Measures") does not fill the column, so it stays a heading."""
    return (
        fills
        and following is not None
        and _font(line) == _font(following)
        and following.text[:1].islower()
        and _next_line(line, following, heights)
    )


def _page_text(text: str) -> str:
    """The text of a line with its numbers left out, to compare page headers and footers across pages."""
    return " ".join(re.sub(r"\d+", " ", text).casefold().split())


def _running(lines: list[_Line]) -> list[bool]:
    """Whether each line is a running page header or footer: among the first or last two lines of its
    page, and with the same text, numbers aside, among the first or last two lines of another page."""
    pages: dict[int, list[int]] = {}
    for index, line in enumerate(lines):
        if line.page is not None:
            pages.setdefault(line.page, []).append(index)
    edges = {index for indices in pages.values() for index in (*indices[:2], *indices[-2:])}
    edge_pages: dict[str, set[int | None]] = {}
    for index in edges:
        edge_pages.setdefault(_page_text(lines[index].text), set()).add(lines[index].page)
    return [index in edges and len(edge_pages[_page_text(line.text)]) > 1 for index, line in enumerate(lines)]


def _merge_wrapped(
    lines: list[_Line], styled: list[bool], running: list[bool], fills: list[bool], heights: bool
) -> tuple[list[_Line], list[bool], list[bool], list[bool]]:
    """The lines with the lines of each wrapped heading merged into one, and whether each is styled, a
    running page header or footer, and ends with a line that fills most of the column (fills, see
    _fills_line). A running line is never merged."""
    merged: list[_Line] = []
    merged_styled: list[bool] = []
    merged_running: list[bool] = []
    # Whether the last line merged into each line fills most of the column.
    merged_fills: list[bool] = []
    for line, line_styled, line_running, line_fills in zip(lines, styled, running, fills, strict=True):
        if (
            merged
            and not line_running
            and not merged_running[-1]
            and _wraps(merged[-1], line, merged_styled[-1], line_styled, heights, merged_fills[-1])
        ):
            first = merged[-1]
            size = None if first.size is None or line.size is None else min(first.size, line.size)
            text = _joined(first.text, line.text)
            merged[-1] = _Line(text, size, first.bold and line.bold, line.page, line.y)
            merged_fills[-1] = line_fills
        else:
            merged.append(line)
            merged_styled.append(line_styled)
            merged_running.append(line_running)
            merged_fills.append(line_fills)
    return merged, merged_styled, merged_running, merged_fills


# The statement's first words as the answer of field 02 (Table 2) or I.02 (Table 4) of 2024/2984, where pypdf
# reads the field number, the label and the start of the answer as one line.
_STATEMENT_ANSWER = re.compile(
    r"\b(?:this\s+crypto-asset\s+white\s+paper|dieses\s+kryptowerte-whitepaper)\b", re.IGNORECASE
)


def _split_statement_fields(
    lines: list[_Line], styled: list[bool], running: list[bool], fills: list[bool]
) -> tuple[list[_Line], list[bool], list[bool], list[bool], set[int]]:
    """Split a summary-table line whose answer is the statement, such as "02 Statement in accordance with
    Article 6(3) of Regulation (EU) 2023/1114 This crypto-asset white paper has not been", into its label
    and the statement's first line. The statement is then routed as one that starts a line of its own,
    and its words ("crypto-asset") do not add a section to the label. Returns the lines, their flags and
    the indices of the statement lines split off."""
    result: tuple[list[_Line], list[bool], list[bool], list[bool]] = ([], [], [], [])
    starts: set[int] = set()
    for line, *flags in zip(lines, styled, running, fills, strict=True):
        answer = _STATEMENT_ANSWER.search(line.text)
        label = line.text[: answer.start()].rstrip() if answer else ""
        parts = [line]
        if (
            answer
            and _template_kind(line.text) in ("field", "summary")
            and "not_approved_statement" in match_heading_to_section(label)
        ):
            parts = [replace(line, text=label), replace(line, text=line.text[answer.start() :])]
            starts.add(len(result[0]) + 1)
        for part in parts:
            for values, value in zip(result, (part, *flags), strict=True):
                values.append(value)
    return (*result, starts)


def _line_sections(lines: list[_Line], statement: bool = False) -> DocumentSections:
    """The sections of a draft read line by line. With statement, the Art. 6(3) / Art. 51(3) statement is
    routed by its content to not_approved_statement (see statement_extent), as in the other loaders."""
    text = "\n".join(line.text for line in lines)
    body_font = _body_font(lines)
    styled = [_styled(line, body_font) for line in lines]
    full_line = _full_line(lines, body_font)
    fills = [_fills_line(line, full_line, body_font) for line in lines]
    heights = _heights_hold(lines, body_font)
    lines, styled, running, fills = _merge_wrapped(lines, styled, _running(lines), fills, heights)
    # The draft sets its headings in a heading font if at least two lines in one name a section.
    by_font = (
        sum(
            1
            for line, line_styled, line_running in zip(lines, styled, running, strict=True)
            if line_styled
            and not line_running
            and _heading_shaped(line.text, True)
            and match_heading_to_section(line.text)
        )
        >= 2
    )
    # Split after the heading-font count, which a statement in a heading font must not tip.
    statement_starts: set[int] = set()
    if statement:
        lines, styled, running, fills, statement_starts = _split_statement_fields(
            lines, styled, running, fills
        )
    # A running page header or footer is body text that is not read as context for the lines around it.
    content = [index for index, line_running in enumerate(running) if not line_running]
    following_index = dict(pairwise(content))
    sections: dict[str, list[str]] = {}
    headings: list[str] = []
    current_keys: list[str] = []
    # The template numbering, the outline number and the level of the heading that opened the current
    # section.
    opener: str | None = None
    opener_number: str | None = None
    section_level: float | None = None
    previous: str | None = None
    previous_role: str | None = None
    previous_styled = False
    line_texts = [line.text for line in lines]
    statement_end = 0
    statement_role = _BODY
    # The first and last two lines of each page, where a page header or footer sits (see _running).
    page_lines: dict[int, list[int]] = {}
    for position, page_line in enumerate(lines):
        if page_line.page is not None:
            page_lines.setdefault(page_line.page, []).append(position)
    page_edges = {i for indices in page_lines.values() for i in (*indices[:2], *indices[-2:])}
    line_counts = Counter(_page_text(page_line.text) for page_line in lines)

    def statement_heading(index: int, statement_line: _Line) -> bool:
        # A heading inside the statement's lines (see statement_extent): a template part or field title,
        # or a line shaped like a heading in a heading font other than the statement's own, since a
        # statement may be set in bold. Where the statement runs over a page break, a line at the edge
        # of the next page may be a page header that names the current part ("Part A ..."), not running
        # because it changes from page to page: one in the body font, one set smaller than the
        # statement, or one that repeats another line of the draft, as a header repeats its heading.
        line = lines[index]
        statement_font = _font(statement_line)
        if running[index]:
            return False
        later_page = line.page != statement_line.page
        smaller = (
            later_page
            and line.size is not None
            and statement_font is not None
            and round(line.size, 1) < statement_font[0]
        )
        page_header = (
            index in page_edges
            and later_page
            and (not styled[index] or line_counts[_page_text(line.text)] > 1)
        )
        if _template_kind(line.text) and _heading_shaped(line.text, styled[index]):
            return not (smaller or page_header)
        return (
            styled[index]
            and _font(line) != statement_font
            and _heading_shaped(line.text, True)
            and not page_header
        )

    for index, line in enumerate(lines):
        if statement and (
            index < statement_end or index in statement_starts or is_statement_start(line.text)
        ):
            if index >= statement_end:
                statement_end = statement_extent(
                    line_texts, index, partial(statement_heading, statement_line=line)
                )
                # A statement cut off before its first full stop ends at a line read as a heading, which
                # does not continue its sentence ("approved by the"). A statement that ends with a full
                # sentence may share its last line with the next sentence, which runs on as body text.
                cut = not _STATEMENT_SENTENCE_END.search(" ".join(line_texts[index:statement_end]))
                statement_role = _BREAK if cut else _BODY
            sections.setdefault("not_approved_statement", []).append(line.text)
            previous, previous_role, previous_styled = line.text, statement_role, styled[index]
            continue
        if running[index]:
            for key in current_keys:
                sections.setdefault(key, []).append(line.text)
            continue
        keys = match_heading_to_section(line.text)
        kind = _template_kind(line.text)
        number = _outline_number(line.text)
        # A larger heading font is a higher level.
        level = -round(line.size, 1) if styled[index] and line.size is not None else None
        following = lines[following_index[index]] if index in following_index else None
        # The first line of a paragraph set in a heading font, such as a bold risk warning, is body text.
        paragraph = styled[index] and _starts_paragraph(line, following, fills[index], heights)
        role = _BODY
        if keys:
            role = _line_role(
                line.text,
                previous,
                previous_role,
                following.text if following else None,
                started=bool(headings),
                styled=styled[index],
                previous_styled=previous_styled,
                by_font=by_font and line.size is not None,
                paragraph=paragraph,
            )
        elif (
            (styled[index] or kind)
            and not paragraph
            and _heading_shaped(line.text, styled[index])
            and _ends_section(
                kind,
                level,
                opener,
                section_level,
                label=styled[index] and not _larger(line, body_font),
                sub_heading=_sub_heading(line.text, number, opener_number, current_keys),
            )
        ):
            # A heading that names no section, such as "Glossary" in a larger heading font.
            role = _BREAK
        if role == _BODY:
            for key in current_keys:
                sections.setdefault(key, []).append(line.text)
        else:
            headings.append(line.text)
            if role == _HEADING:
                section_level = _opening_level(level, section_level, bool(set(keys) & set(current_keys)))
                current_keys = keys
                answer = _field_answer(line.text, keys)
                for key in keys if answer else ():
                    sections.setdefault(key, []).append(answer)
            else:
                section_level = level
                current_keys = []
            opener, opener_number = kind, _opening_number(number, opener_number)
        previous, previous_role, previous_styled = line.text, role, styled[index]
    return DocumentSections({k: "\n".join(v) for k, v in sections.items()}, tuple(headings), text)


def sections_from_lines(lines: Iterable[str]) -> dict[str, str]:
    """Group plain text lines, such as the extracted text of a PDF, under the section headings among them.

    Without font information a heading is recognised by its shape alone (see _line_role).
    """
    return _line_sections([_Line(line.strip()) for line in lines if line.strip()]).sections


def _style_chain(style):
    """The style and the styles it is based on, stopping where a malformed document loops back.

    A numbering (list) style has no base style in python-docx, so the chain ends there.
    """
    seen = set()
    while style is not None and style.element not in seen:
        seen.add(style.element)
        yield style
        style = getattr(style, "base_style", None)


def _outline_level(ppr) -> int | None:
    """The w:outlineLvl of a paragraph or style: 0 to 8 are heading levels 1 to 9, 9 is body text."""
    from docx.oxml.ns import qn

    level = ppr.find(qn("w:outlineLvl")) if ppr is not None else None
    value = level.get(qn("w:val")) if level is not None else None
    return int(value) if value is not None and value.isdigit() else None


# A bold or capitals label without an outline level ranks below the headings of levels 1 to 9 (0 to 8).
_LABEL_LEVEL = 9


def _docx_heading_level(paragraph) -> int | None:
    """The outline level of a heading, 0 for a level 1 heading, from a Heading style or an outline level
    set on the paragraph or on the style it is based on. None for body text."""
    level = _outline_level(paragraph._p.pPr)
    if level is not None:
        return level if level < 9 else None
    for style in _style_chain(paragraph.style):
        name = style.name or ""
        if name.startswith("Heading"):
            number = name.removeprefix("Heading").strip()
            return int(number) - 1 if number.isdigit() and 1 <= int(number) <= 9 else 0
        level = _outline_level(style.element.pPr)
        if level is not None:
            return level if level < 9 else None
    return None


def _style_bold(style) -> bool | None:
    """The bold setting of the first style in the chain that sets one; a numbering style has no font."""
    fonts = (getattr(s, "font", None) for s in _style_chain(style))
    return next((font.bold for font in fonts if font is not None and font.bold is not None), None)


def _run_bold(run, paragraph_bold: bool) -> bool:
    """Direct formatting first, then the character style, then the paragraph style."""
    for bold in (run.bold, _style_bold(run.style)):
        if bold is not None:
            return bold
    return paragraph_bold


def _docx_fully_bold(paragraph) -> bool:
    """True if every letter after a leading number sits in a bold run, hyperlinks included, not just a
    bold lead-in label. A number typed in plain type before a bold title ("5." then "Risk Factors") counts."""
    from docx.text.hyperlink import Hyperlink

    paragraph_bold = bool(_style_bold(paragraph.style))
    runs = [
        run
        for item in paragraph.iter_inner_content()
        for run in (item.runs if isinstance(item, Hyperlink) else (item,))
    ]
    bold_letters: list[bool] = []
    for run in runs:
        letters = sum(char.isalpha() for char in run.text)
        if letters:
            bold_letters += [_run_bold(run, paragraph_bold)] * letters
    numbering = _NUMBERING.match(paragraph.text.lstrip())
    bold_letters = bold_letters[sum(char.isalpha() for char in numbering.group()) if numbering else 0 :]
    return bool(bold_letters) and all(bold_letters)


def _docx_label_heading(paragraph, text: str) -> bool:
    """A label of up to 20 words and under 200 characters that is entirely bold or in capitals."""
    return (
        len(text) < 10 * _MAX_MARKED_HEADING_WORDS
        and _is_label(text, _max_words(text, marked=True))
        and (text.isupper() or _docx_fully_bold(paragraph))
    )


def parse_docx(path: Path) -> DocumentSections:
    """Parse a DOCX file and group text paragraphs under identified section headings.

    A heading has a Heading style or an outline level, or is a short label that is entirely bold
    or in capitals. Heading text is not section body. A heading that names no section ends the
    current section unless it is a sub-heading of the heading that opened it (see _ends_section).
    """
    try:
        import docx
    except ImportError as exc:
        raise SystemExit(
            "DOCX support requires the optional 'docx' extra: pip install micar-whitepaper-linter[docx]"
        ) from exc

    doc = docx.Document(path)
    sections: dict[str, list[str]] = {}
    headings: list[str] = []
    texts: list[str] = []
    current_keys: list[str] = []
    # The template numbering, the outline number and the outline level of the heading that opened the
    # current section.
    opener: str | None = None
    opener_number: str | None = None
    section_level: float | None = None
    paragraphs = [(para, para.text.strip()) for para in doc.paragraphs if para.text.strip()]
    paragraph_texts = [text for _, text in paragraphs]
    statement_end = 0

    for index, (para, text) in enumerate(paragraphs):
        texts.append(text)
        if index < statement_end or is_statement_start(text):
            if index >= statement_end:
                # A heading as parse_docx reads one below: a Heading style, an outline level, or a label
                # that is entirely bold or in capitals. A label marked as the statement's first paragraph
                # is one of its lines, as a statement may be set in bold or capitals a paragraph per line.
                marks = (text.isupper(), _docx_fully_bold(para))
                statement_end = paragraph_statement_extent(
                    paragraph_texts,
                    index,
                    lambda i, marks=marks: (
                        _docx_heading_level(paragraphs[i][0]) is not None
                        or (
                            _docx_label_heading(*paragraphs[i])
                            and (paragraphs[i][1].isupper(), _docx_fully_bold(paragraphs[i][0])) != marks
                        )
                    ),
                )
            sections.setdefault("not_approved_statement", []).append(text)
            continue

        # The Title style names the white paper, not one of its sections.
        if para.style is not None and para.style.name == "Title":
            continue

        level = _docx_heading_level(para)
        if level is None and _docx_label_heading(para, text):
            level = _LABEL_LEVEL
        if level is None:
            for k in current_keys:
                sections.setdefault(k, []).append(text)
            continue

        headings.append(text)
        keys = match_heading_to_section(text)
        kind = _template_kind(text)
        number = _outline_number(text)
        if keys:
            section_level = _opening_level(level, section_level, bool(set(keys) & set(current_keys)))
            current_keys, opener = keys, kind
            opener_number = _opening_number(number, opener_number)
        elif _ends_section(
            kind,
            level,
            opener,
            section_level,
            label=level == _LABEL_LEVEL,
            sub_heading=_sub_heading(text, number, opener_number, current_keys),
        ):
            current_keys, opener, section_level = [], kind, level
            opener_number = _opening_number(number, opener_number)

    return DocumentSections({k: "\n".join(v) for k, v in sections.items()}, tuple(headings), "\n".join(texts))


def load_from_docx(path: Path) -> dict[str, str]:
    return parse_docx(path).sections


_BOLD_FACE = re.compile(r"bold|black|heavy|demi", re.I)


def _resolved(value):
    return value.get_object() if hasattr(value, "get_object") else value


def _bold_face(font) -> bool:
    """Bold by name ("Helvetica-Bold", "ABCDEF+Arial-BoldMT") or by a font weight of 600 or more."""
    if font is None:
        return False
    if _BOLD_FACE.search(str(font.get("/BaseFont", ""))):
        return True
    for face in (font, *(_resolved(d) for d in _resolved(font.get("/DescendantFonts")) or ())):
        weight = _resolved((_resolved(face.get("/FontDescriptor")) or {}).get("/FontWeight"))
        if isinstance(weight, int | float) and weight >= 600:
            return True
    return False


_RunStyle = tuple[float, bool, float]


def _font_style(cm, tm, font, size) -> _RunStyle | None:
    """The rendered size of a text run (its Tf size scaled by the text and graphics matrices), its weight
    and the height of its baseline on the page."""
    try:
        scale = math.hypot(tm[2] * cm[0] + tm[3] * cm[2], tm[2] * cm[1] + tm[3] * cm[3])
        y = tm[4] * cm[1] + tm[5] * cm[3] + cm[5]
        return abs(float(size)) * scale, _bold_face(font), float(y)
    except Exception:
        return None


def _record_run(runs: list[tuple[str, _RunStyle | None]], text, cm, tm, font, size) -> None:
    runs.append((str(text), _font_style(cm, tm, font, size)))


def _pdf_lines(reader) -> list[_Line]:
    """The non-empty text lines of a PDF, with the font and height of each where pypdf reports them.

    The text visitor of pypdf reports each run of text with its font and position. Where the runs do not
    add up to the extracted page text (form XObjects, right-to-left text), that page goes by line shape
    alone.
    """
    lines: list[_Line] = []
    for number, page in enumerate(reader.pages):
        runs: list[tuple[str, _RunStyle | None]] = []
        text = page.extract_text(visitor_text=partial(_record_run, runs)) or ""
        styles = (
            [style for run, style in runs for _ in run] if "".join(run for run, _ in runs) == text else []
        )
        end = 0
        for raw in text.splitlines(keepends=True):
            start, end = end, end + len(raw)
            if not raw.strip():
                continue
            letters = [styles[i] for i in range(start, end) if text[i].isalpha()] if styles else []
            if letters and None not in letters:
                size = min(style[0] for style in letters)
                bold = all(style[1] for style in letters)
                lines.append(_Line(raw.strip(), size, bold, number, letters[0][2]))
            else:
                lines.append(_Line(raw.strip(), page=number))
    return lines


def parse_pdf(path: Path) -> DocumentSections:
    """Parse a PDF file and extract text grouped under identified section headings.

    A heading is a line that names a section and is set larger or bolder than the body text, or is
    shaped like a heading (see _line_role); where the draft sets its headings in a heading font, a line
    in the body font is a heading only if numbered as in the template. A short line that names a section
    but is not accepted as a heading ends the current section, and so does a heading that names no
    section unless it is a sub-heading (see _ends_section). A running page header or footer is body
    text. Heading text is not section body.
    """
    try:
        import pypdf
    except ImportError as exc:
        raise SystemExit(
            "PDF support requires the optional 'pdf' extra: pip install micar-whitepaper-linter[pdf]"
        ) from exc

    return _line_sections(_pdf_lines(pypdf.PdfReader(path)), statement=True)


def load_from_pdf(path: Path) -> dict[str, str]:
    return parse_pdf(path).sections


def load_from_markdown(path: Path) -> dict[str, str]:
    """Parse a Markdown file and extract text grouped under identified section headings."""
    try:
        content = path.read_text(encoding="utf-8")
    except Exception as exc:
        raise SystemExit(f"Error reading Markdown file {path}: {exc}") from exc

    sections: dict[str, list[str]] = {}
    current_keys = []
    lines = [line.strip() for line in content.splitlines() if line.strip()]
    statement_end = 0

    for index, line_clean in enumerate(lines):
        if index < statement_end or is_statement_start(line_clean):
            if index >= statement_end:
                statement_end = statement_extent(lines, index)
            sections.setdefault("not_approved_statement", []).append(line_clean)
            continue

        # Detect Markdown headings (e.g. # Heading 1, ## Heading 2)
        heading_match = re.match(r"^#{1,6}\s+(.+)$", line_clean)
        if heading_match:
            heading_text = heading_match.group(1).strip()
            matched_keys = match_heading_to_section(heading_text)
            # Reset current_keys if it's a heading, so we don't leak text into unrelated sections
            current_keys = matched_keys
        else:
            if current_keys:
                for k in current_keys:
                    sections.setdefault(k, []).append(line_clean)

    return {k: "\n".join(v) for k, v in sections.items()}
