"""Review workbook and merge tool.

Every label in this file is synthetic and lives in a study directory that the generator builds
under tmp_path, so the tests do not depend on the committed labelling progress.
The autouse fixture checks that no test changes a committed study file.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import shutil
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from micar_linter import review_merge, review_sheets, review_validation, review_workbook
from micar_linter.review_merge import (
    ADJUDICATION_FIELDS,
    EXIT_ADJUDICATION,
    EXIT_IDENTITY,
    EXIT_OK,
    EXIT_OVERWRITE,
    EXIT_REFUSED,
    EXIT_USAGE,
    SUMMARY_FILENAME,
    _ranges,
    load_study,
    read_sheet,
)
from micar_linter.review_merge import main as merge_main
from micar_linter.review_sheets import (
    CELL_FIELDS,
    DETECTOR_FIELDS,
    MATRIX_FILENAME,
    REVIEW_FIELDS,
    read_sheets_record,
)
from micar_linter.review_validation import LABEL_ORDER, LABEL_REQUIRED_FIELDS, summarize_rows
from micar_linter.review_workbook import (
    CHECK_COLUMN,
    META_SHEET,
    REVIEW_SHEET,
    WORKBOOK_COLUMNS,
    label_prompt,
)
from micar_linter.review_workbook import main as workbook_main

REPO = Path(__file__).resolve().parents[1]
STUDY = REPO / "studies" / "2026-07-title-ii-annex-i-whitepaper-study"
FINDINGS_FILE = "findings-pseudonymous.json"
TODAY = date(2026, 10, 14)
REVIEW_DATE = "2026-07-20"
FINAL_LABELS = [label for label in LABEL_ORDER if label != "pending"]
COLUMN = {name: index for index, name in enumerate(WORKBOOK_COLUMNS, start=1)}


def _hashes(directory: Path) -> dict[str, str]:
    return {
        path.name: hashlib.sha256(path.read_bytes()).hexdigest()
        for path in sorted(directory.iterdir())
        if path.is_file()
    }


@pytest.fixture(autouse=True)
def committed_study_untouched():
    before = _hashes(STUDY)
    yield
    assert _hashes(STUDY) == before, "a test changed a committed study file"


@pytest.fixture(autouse=True)
def fixed_today(monkeypatch):
    monkeypatch.setattr(review_merge, "_today", lambda: TODAY)


@pytest.fixture
def study_dir(tmp_path: Path) -> Path:
    """A pending copy of the study for reviewer r1, whatever the committed labelling progress.

    Only the findings file is copied. The sheets, the matrix and the generation record come from
    the generator with the committed seed, and the summary from micar-review-summary, so the
    committed labels (once they exist) never reach these tests. The r1 sheet and the matrix must
    hash to the values recorded in the committed human-review-sheets.json.
    """
    target = tmp_path / "study"
    target.mkdir()
    shutil.copy2(STUDY / FINDINGS_FILE, target / FINDINGS_FILE)
    record = read_sheets_record(STUDY)
    assert record is not None
    outputs = review_sheets.generate_review_sheets(
        target / FINDINGS_FILE, reviewers=("r1",), seed=int(record["seed"])
    )
    recorded = {sheet["file"]: sheet["sha256"] for sheet in record["sheets"]}
    recorded[record["matrix"]["file"]] = record["matrix"]["sha256"]
    for name in ("human-review-r1.csv", MATRIX_FILENAME):
        assert hashlib.sha256(outputs[name].encode("utf-8")).hexdigest() == recorded[name], name
    review_sheets.write_review_sheets(outputs, target)
    summary = ["--matrix", str(target / MATRIX_FILENAME), "--out", str(target / SUMMARY_FILENAME)]
    assert review_validation.main(summary) == 0
    return target


def _synthetic(index: int, reviewer: str, *, label: str | None = None, when: str = REVIEW_DATE) -> dict:
    """Synthetic, clearly marked test values for one cell. Never written outside tmp_path."""
    label = label or FINAL_LABELS[index % len(FINAL_LABELS)]
    values = {
        "human_label": label,
        "reviewer_id": reviewer,
        "reviewed_at": when,
        "source_pinpoint": "",
        "review_note": "",
    }
    for name in LABEL_REQUIRED_FIELDS[label]:
        values[name] = f"synthetic test value {index}"
    return values


def _csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def _write_csv(path: Path, rows: list[dict[str, str]]) -> None:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=list(rows[0]), lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    path.write_text(buffer.getvalue(), encoding="utf-8")


def _label_csv(path: Path, reviewer: str, overrides: dict[int, dict] | None = None, when: str = REVIEW_DATE):
    rows = _csv(path)
    for index, row in enumerate(rows):
        row.update(_synthetic(index, reviewer, when=when))
        row.update((overrides or {}).get(index + 2, {}))
    _write_csv(path, rows)


def _make_workbook(study_dir: Path, reviewer: str = "r1") -> Path:
    assert workbook_main(["--study-dir", str(study_dir), "--reviewer", reviewer]) == EXIT_OK
    return study_dir / f"human-review-{reviewer}.xlsx"


def _fill_workbook(
    path: Path,
    reviewer: str = "r1",
    overrides: dict[int, dict] | None = None,
    keep: frozenset[int] = frozenset(),
) -> None:
    """Fill the editable columns the way a reviewer would.

    ``overrides`` is keyed by sheet row; rows in ``keep`` stay as generated unless overridden.
    """
    workbook = openpyxl.load_workbook(path)
    sheet = workbook[REVIEW_SHEET]
    for row_number in range(2, sheet.max_row + 1):
        if row_number in keep and row_number not in (overrides or {}):
            continue
        values = _synthetic(row_number - 2, reviewer)
        values["reviewed_at"] = date.fromisoformat(values["reviewed_at"])
        values.update((overrides or {}).get(row_number, {}))
        for name, value in values.items():
            sheet.cell(row=row_number, column=COLUMN[name], value=value if value != "" else None)
    workbook.save(path)


def _import(study_dir: Path, workbook: Path, reviewer: str = "r1", *extra: str) -> int:
    args = ["import", "--study-dir", str(study_dir), "--reviewer", reviewer, "--workbook", str(workbook)]
    return merge_main([*args, *extra])


def _matrix(study_dir: Path, *extra: str) -> int:
    return merge_main(["matrix", "--study-dir", str(study_dir), *extra])


def _add_reviewer(study_dir: Path, reviewer: str) -> None:
    args = ["--findings", str(study_dir / "findings-pseudonymous.json"), "--out-dir", str(study_dir)]
    assert review_sheets.main([*args, "--reviewer", reviewer]) == 0


def _markdown_table(text: str, start: str, end: str) -> list[list[str]]:
    section = text.split(start, 1)[1].split(end, 1)[0]
    rows = [line for line in section.splitlines() if line.startswith("| `")]
    return [[cell.strip() for cell in row.strip("|").split("|")] for row in rows]


def test_label_vocabulary_matches_the_codebook_and_the_protocol():
    codebook = (STUDY / "codebook.md").read_text(encoding="utf-8")
    table = _markdown_table(codebook, "## 3. Labels", "**Order of decisions")
    assert [row[0].strip("`") for row in table] == list(LABEL_ORDER)
    for label, _use, required in table:
        for name in LABEL_REQUIRED_FIELDS.get(label.strip("`"), ()):
            assert f"`{name}`" in required, (label, name)

    protocol = (STUDY / "review-protocol.md").read_text(encoding="utf-8")
    protocol_labels = [row[0].strip("`") for row in _markdown_table(protocol, "## Labels", "## First review")]
    assert sorted(protocol_labels) == sorted(LABEL_ORDER)
    assert set(LABEL_REQUIRED_FIELDS) == set(LABEL_ORDER) - {"pending"}


def test_workbook_layout_validation_and_protection(study_dir: Path):
    path = _make_workbook(study_dir)
    workbook = openpyxl.load_workbook(path)

    assert workbook.sheetnames == ["Instructions", REVIEW_SHEET, META_SHEET]
    assert workbook[META_SHEET].sheet_state == "veryHidden"
    assert workbook.security.lockStructure
    instructions = "\n".join(str(cell.value) for cell in workbook["Instructions"]["A"] if cell.value)
    assert "codebook.md" in instructions and "review-protocol.md" in instructions

    sheet = workbook[REVIEW_SHEET]
    assert [cell.value for cell in sheet[1]] == list(WORKBOOK_COLUMNS)
    assert sheet.freeze_panes == "B2"
    assert all(sheet.column_dimensions[cell.column_letter].width > 8 for cell in sheet[1])

    committed = _csv(study_dir / "human-review-r1.csv")
    assert sheet.max_row == len(committed) + 1 == 151
    for row_number, row in enumerate(committed, start=2):
        assert [sheet.cell(row_number, COLUMN[name]).value for name in CELL_FIELDS] == [
            row[name] for name in CELL_FIELDS
        ]
        assert sheet.cell(row_number, COLUMN["human_label"]).value == "pending"
        assert all(sheet.cell(row_number, COLUMN[name]).value is None for name in REVIEW_FIELDS[1:])
        formula = str(sheet.cell(row_number, COLUMN[CHECK_COLUMN]).value)
        assert formula.startswith("=IF(")
        assert f'EXACT(TRIM(G{row_number}),"present")' in formula
        assert f'NOT(EXACT(TRIM(H{row_number}),"r1"))' in formula
        assert '="present"' not in formula and '<>"r1"' not in formula

    assert sheet.protection.sheet
    assert sheet.protection.sort and sheet.protection.insertRows and sheet.protection.deleteRows
    assert not sheet.protection.formatColumns and not sheet.protection.autoFilter
    assert all(sheet.cell(2, COLUMN[name]).protection.locked for name in (*CELL_FIELDS, CHECK_COLUMN))
    assert not any(sheet.cell(151, COLUMN[name]).protection.locked for name in REVIEW_FIELDS)

    validations = {str(item.sqref): item for item in sheet.data_validations.dataValidation}
    labels = validations["G2:G151"]
    assert labels.type == "list"
    assert labels.formula1 == '"' + ",".join(LABEL_ORDER) + '"'
    assert labels.errorStyle == "stop" and labels.showErrorMessage
    assert labels.prompt == label_prompt() and len(label_prompt()) <= 255
    assert validations["H2:H151"].formula1 == '"r1"'
    dates = validations["I2:I151"]
    assert (dates.type, dates.operator, dates.formula1, dates.formula2) == (
        "date",
        "between",
        "DATE(2026,7,6)",
        "TODAY()",
    )


def test_workbook_carries_no_detector_outcome(study_dir: Path):
    path = _make_workbook(study_dir)
    workbook = openpyxl.load_workbook(path)
    texts = [
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]
    contexts = {row["detector_context"] for row in _csv(study_dir / "human-review-matrix.csv")}
    contexts.discard("")
    assert contexts
    for text in texts:
        assert not any(context in text for context in contexts)
        assert "flagged" not in text
        assert not any(name in text for name in DETECTOR_FIELDS)


def test_round_trip_from_workbook_to_matrix(study_dir: Path):
    original_matrix = _csv(study_dir / "human-review-matrix.csv")
    workbook = _make_workbook(study_dir)
    _fill_workbook(workbook)

    assert _import(study_dir, workbook) == EXIT_OK
    sheet = _csv(study_dir / "human-review-r1.csv")
    for index, row in enumerate(sheet):
        assert {name: row[name] for name in REVIEW_FIELDS} == _synthetic(index, "r1")
    check = [
        "--findings",
        str(study_dir / "findings-pseudonymous.json"),
        "--out-dir",
        str(study_dir),
        "--check",
    ]
    assert review_sheets.main(check) == 0
    assert merge_main(["check", "--study-dir", str(study_dir), "--reviewer", "r1"]) == EXIT_OK

    assert _matrix(study_dir) == EXIT_OK
    matrix = _csv(study_dir / "human-review-matrix.csv")
    by_cell = {(row["study_doc_id"], row["rule_id"]): row for row in sheet}
    assert len(matrix) == len(original_matrix) == 150
    for merged, before in zip(matrix, original_matrix, strict=True):
        for name in (*CELL_FIELDS, *DETECTOR_FIELDS):
            assert merged[name] == before[name]
        labelled = by_cell[(merged["study_doc_id"], merged["rule_id"])]
        assert {name: merged[name] for name in REVIEW_FIELDS} == {
            name: labelled[name] for name in REVIEW_FIELDS
        }

    summary = json.loads((study_dir / "human-review-summary.json").read_text(encoding="utf-8"))
    assert summary == summarize_rows(matrix)
    assert summary["publication_ready"] is True
    assert summary["rows_pending"] == 0
    assert review_sheets.main(check) == 0
    # Re-running the merge on unchanged labels is a no-op.
    assert _matrix(study_dir) == EXIT_OK


def test_import_refuses_while_any_cell_is_pending(study_dir: Path, capsys):
    workbook = _make_workbook(study_dir)
    _fill_workbook(workbook, overrides={77: {"human_label": "pending"}, 9: {"human_label": None}})
    before = (study_dir / "human-review-r1.csv").read_bytes()

    assert _import(study_dir, workbook) == EXIT_REFUSED
    assert "2 pending, rows 9, 77" in capsys.readouterr().err
    assert (study_dir / "human-review-r1.csv").read_bytes() == before
    check = ["check", "--study-dir", str(study_dir), "--reviewer", "r1", "--workbook", str(workbook)]
    assert merge_main(check) == EXIT_REFUSED


@pytest.mark.parametrize(
    ("override", "message"),
    [
        ({"human_label": "presnt"}, "human_label 'presnt' is not one of"),
        ({"reviewed_at": "14.10.2026"}, "is not an ISO 8601 date"),
        ({"reviewed_at": "2026-02-30"}, "is not a calendar date"),
        ({"reviewed_at": "2026-10-15"}, "is after today (2026-10-14)"),
        ({"reviewed_at": "2026-07-01"}, "is before 2026-07-06"),
        ({"reviewed_at": None}, "reviewed_at is blank"),
        ({"reviewer_id": "r2"}, "reviewer_id 'r2' is not this sheet's reviewer 'r1'"),
        ({"reviewer_id": None}, "reviewer_id is blank"),
        ({"human_label": "present", "source_pinpoint": None}, "source_pinpoint is required for 'present'"),
        ({"human_label": "missing", "source_pinpoint": " "}, "source_pinpoint is required for 'missing'"),
        ({"human_label": "uncertain", "review_note": None}, "review_note is required for 'uncertain'"),
        (
            {"human_label": "extraction_failure", "review_note": None},
            "review_note is required for 'extraction_failure'",
        ),
    ],
)
def test_import_refuses_invalid_values(study_dir: Path, capsys, override: dict, message: str):
    workbook = _make_workbook(study_dir)
    _fill_workbook(workbook, overrides={40: override})
    before = (study_dir / "human-review-r1.csv").read_bytes()

    assert _import(study_dir, workbook) == EXIT_REFUSED
    err = capsys.readouterr().err
    assert re.search(r"row 40 \(WP-\d{3}, [A-Z0-9_]+\): [^\n]*" + re.escape(message), err), err
    assert (study_dir / "human-review-r1.csv").read_bytes() == before


def _swap_rule_ids(sheet, first: int, second: int) -> None:
    column = COLUMN["rule_id"]
    a, b = sheet.cell(first, column).value, sheet.cell(second, column).value
    sheet.cell(first, column).value, sheet.cell(second, column).value = b, a


@pytest.mark.parametrize(
    "tamper",
    [
        "swap rows",
        "edit hash",
        "drop last row",
        "edit requirement",
        "extra column",
        "meta digest",
        "meta seed",
    ],
)
def test_import_detects_workbook_identity_tampering(study_dir: Path, capsys, tamper: str):
    workbook_path = _make_workbook(study_dir)
    _fill_workbook(workbook_path)
    workbook = openpyxl.load_workbook(workbook_path)
    sheet = workbook[REVIEW_SHEET]
    if tamper == "swap rows":
        _swap_rule_ids(sheet, 2, 3)
    elif tamper == "edit hash":
        sheet.cell(10, COLUMN["document_hash_sha256"]).value = "0" * 64
    elif tamper == "drop last row":
        for column in range(1, len(WORKBOOK_COLUMNS) + 1):
            sheet.cell(151, column).value = None
    elif tamper == "edit requirement":
        sheet.cell(20, COLUMN["requirement"]).value = "Something else."
    elif tamper == "extra column":
        sheet.cell(30, len(WORKBOOK_COLUMNS) + 1).value = "stray"
    else:
        meta = workbook[META_SHEET]
        key = "identity_sha256" if tamper == "meta digest" else "seed"
        for row in meta.iter_rows():
            if row[0].value == key:
                row[1].value = "0"
    workbook.save(workbook_path)
    before = (study_dir / "human-review-r1.csv").read_bytes()

    assert _import(study_dir, workbook_path) == EXIT_IDENTITY
    assert "cell identity or row order differs" in capsys.readouterr().err
    assert (study_dir / "human-review-r1.csv").read_bytes() == before


def test_import_rejects_another_reviewers_workbook(study_dir: Path, capsys):
    workbook = _make_workbook(study_dir, "r1")
    _fill_workbook(workbook, "r2")
    _add_reviewer(study_dir, "r2")
    before = (study_dir / "human-review-r2.csv").read_bytes()

    assert _import(study_dir, workbook, "r2") == EXIT_IDENTITY
    assert "_meta reviewer_id is 'r1', expected 'r2'" in capsys.readouterr().err
    assert (study_dir / "human-review-r2.csv").read_bytes() == before


def test_changed_findings_or_reordered_sheet_are_refused(study_dir: Path, capsys):
    workbook = _make_workbook(study_dir)
    _fill_workbook(workbook)

    sheet = study_dir / "human-review-r1.csv"
    original = sheet.read_text(encoding="utf-8")
    lines = original.splitlines(keepends=True)
    sheet.write_text("".join([lines[0], lines[2], lines[1], *lines[3:]]), encoding="utf-8")
    assert _import(study_dir, workbook) == EXIT_IDENTITY
    assert "human-review-r1.csv: cell identity or row order differs" in capsys.readouterr().err
    sheet.write_text(original, encoding="utf-8")

    findings = study_dir / "findings-pseudonymous.json"
    findings.write_bytes(findings.read_bytes() + b"\n")
    assert _import(study_dir, workbook) == EXIT_IDENTITY
    assert _matrix(study_dir) == EXIT_IDENTITY
    assert "generated from a different findings file" in capsys.readouterr().err
    assert sheet.read_text(encoding="utf-8") == original


def test_import_and_workbook_refuse_to_overwrite(study_dir: Path):
    workbook = _make_workbook(study_dir)
    assert workbook_main(["--study-dir", str(study_dir), "--reviewer", "r1"]) == EXIT_OVERWRITE
    assert workbook_main(["--study-dir", str(study_dir), "--reviewer", "r3"]) == EXIT_USAGE
    _fill_workbook(workbook)
    assert _import(study_dir, workbook) == EXIT_OK
    assert _import(study_dir, workbook) == EXIT_OK  # same labels: nothing to overwrite

    _fill_workbook(workbook, overrides={5: {"review_note": "synthetic changed note"}})
    labelled = (study_dir / "human-review-r1.csv").read_bytes()
    assert _import(study_dir, workbook) == EXIT_OVERWRITE
    assert (study_dir / "human-review-r1.csv").read_bytes() == labelled
    assert _import(study_dir, workbook, "r1", "--replace") == EXIT_OK
    assert _csv(study_dir / "human-review-r1.csv")[3]["review_note"] == "synthetic changed note"

    # A workbook generated from the labelled sheet carries the reviewer's own values back.
    assert workbook_main(["--study-dir", str(study_dir), "--reviewer", "r1", "--force"]) == EXIT_OK
    assert _import(study_dir, workbook) == EXIT_OK


def _label_csv_rows(path: Path, values: dict[int, dict]) -> None:
    """Enter synthetic values in the CSV directly, as a reviewer may; ``values`` is keyed by row."""
    rows = _csv(path)
    for row_number, update in values.items():
        rows[row_number - 2].update(update)
    _write_csv(path, rows)


def test_import_continues_a_workbook_generated_from_the_current_csv(study_dir: Path):
    # Started in the CSV, continued in a workbook generated from it: no --replace needed.
    sheet = study_dir / "human-review-r1.csv"
    started = {2: _synthetic(0, "r1", label="uncertain"), 3: _synthetic(1, "r1", label="present")}
    _label_csv_rows(sheet, started)
    workbook = _make_workbook(study_dir)
    revised = {**started[3], "review_note": "synthetic note revised in the workbook"}
    _fill_workbook(workbook, keep=frozenset({2, 3}), overrides={3: revised})

    assert _import(study_dir, workbook) == EXIT_OK
    rows = _csv(sheet)
    assert {name: rows[0][name] for name in REVIEW_FIELDS} == _synthetic(0, "r1", label="uncertain")
    assert rows[1]["review_note"] == "synthetic note revised in the workbook"
    assert all(row["human_label"] != "pending" for row in rows)


def test_import_refuses_when_the_csv_changed_after_the_workbook(study_dir: Path, capsys):
    sheet = study_dir / "human-review-r1.csv"
    _label_csv_rows(sheet, {2: _synthetic(0, "r1", label="uncertain")})
    workbook = _make_workbook(study_dir)
    _fill_workbook(workbook, keep=frozenset({2}))
    # The reviewer labels another row in the CSV after generating the workbook.
    _label_csv_rows(sheet, {5: _synthetic(3, "r1", label="missing")})
    changed = sheet.read_bytes()

    assert _import(study_dir, workbook) == EXIT_OVERWRITE
    assert "changed after human-review-r1.xlsx was generated from it" in capsys.readouterr().err
    assert sheet.read_bytes() == changed
    assert _import(study_dir, workbook, "r1", "--replace") == EXIT_OK
    assert _csv(sheet)[3]["human_label"] == _synthetic(3, "r1")["human_label"]


def test_not_before_raises_the_earliest_review_date(study_dir: Path, capsys):
    _label_csv(study_dir / "human-review-r1.csv", "r1")  # every reviewed_at is 2026-07-20
    check = ["check", "--study-dir", str(study_dir), "--reviewer", "r1"]
    assert merge_main([*check, "--not-before", "2026-07-20"]) == EXIT_OK
    assert merge_main([*check, "--not-before", "2026-07-21"]) == EXIT_REFUSED
    assert "is before 2026-07-21, the --not-before date" in capsys.readouterr().err
    assert _matrix(study_dir, "--not-before", "2026-07-21") == EXIT_REFUSED
    assert merge_main([*check, "--not-before", "2026-07-05"]) == EXIT_USAGE
    assert "is before 2026-07-06, the date the findings were generated" in capsys.readouterr().err
    with pytest.raises(SystemExit) as exited:
        merge_main([*check, "--not-before", "2026-13-01"])
    assert exited.value.code == EXIT_USAGE

    args = ["--study-dir", str(study_dir), "--reviewer", "r1", "--not-before", "2026-07-21"]
    assert workbook_main(args) == EXIT_OK
    workbook = openpyxl.load_workbook(study_dir / "human-review-r1.xlsx")
    dates = {str(item.sqref): item for item in workbook[REVIEW_SHEET].data_validations.dataValidation}
    assert dates["I2:I151"].formula1 == "DATE(2026,7,21)"
    instructions = "\n".join(str(cell.value) for cell in workbook["Instructions"]["A"] if cell.value)
    assert "--reviewer r1 --not-before 2026-07-21 --workbook" in instructions


def test_workbook_copies_the_reviewers_own_values_verbatim(study_dir: Path):
    # A reviewer who started in the CSV gets their own values back; nothing else is filled.
    sheet = study_dir / "human-review-r1.csv"
    rows = _csv(sheet)
    rows[0].update(_synthetic(0, "r1", label="uncertain"))
    rows[0]["review_note"] = "=synthetic note that looks like a formula"
    rows[1].update(_synthetic(1, "r1", label="present"))
    _write_csv(sheet, rows)

    workbook = _make_workbook(study_dir)
    study = load_study(study_dir)
    read_back = review_workbook.read_workbook(workbook, study, "r1")
    assert [{name: row[name] for name in REVIEW_FIELDS} for row in read_back] == [
        {name: row[name] for name in REVIEW_FIELDS} for row in rows
    ]
    assert {row["human_label"] for row in read_back[2:]} == {"pending"}


def test_matrix_refuses_pending_or_invalid_sheets(study_dir: Path, capsys):
    before = _hashes(study_dir)
    assert _matrix(study_dir) == EXIT_REFUSED
    assert "human-review-r1.csv: 150 pending, rows 2-151" in capsys.readouterr().err

    _label_csv(study_dir / "human-review-r1.csv", "r1", overrides={12: {"reviewed_at": "2026/07/20"}})
    labelled = _hashes(study_dir)
    assert _matrix(study_dir) == EXIT_REFUSED
    assert "row 12" in capsys.readouterr().err
    assert _hashes(study_dir) == labelled
    assert {name: labelled[name] for name in before if name != "human-review-r1.csv"} == {
        name: before[name] for name in before if name != "human-review-r1.csv"
    }


def _two_reviewer_study(study_dir: Path) -> list[tuple[str, str]]:
    """Label r1 and r2 with synthetic values that differ on three cells; return those cells."""
    _add_reviewer(study_dir, "r2")
    _label_csv(study_dir / "human-review-r1.csv", "r1")
    r2_path = study_dir / "human-review-r2.csv"
    r1 = {(row["study_doc_id"], row["rule_id"]): row for row in _csv(study_dir / "human-review-r1.csv")}
    rows = _csv(r2_path)
    differing = [(rows[i]["study_doc_id"], rows[i]["rule_id"]) for i in (4, 50, 120)]
    for index, row in enumerate(rows):
        key = (row["study_doc_id"], row["rule_id"])
        label = r1[key]["human_label"]
        if key in differing:
            label = "missing" if label != "missing" else "present"
        row.update(_synthetic(index, "r2", label=label, when="2026-07-25"))
    _write_csv(r2_path, rows)
    return differing


def _adjudication(study_dir: Path, report: list[dict[str, str]], decide, *, bom: bool = False) -> Path:
    """Copy the disagreement report and fill the adjudication columns with synthetic decisions."""
    rows = [dict(row) for row in report]
    for row in rows:
        row.update(decide(row))
    path = study_dir.parent / "adjudication.csv"
    _write_csv(path, rows)
    if bom:  # as a spreadsheet application's "CSV UTF-8" export writes it
        path.write_bytes(b"\xef\xbb\xbf" + path.read_bytes())
    return path


def test_two_reviewers_need_an_adjudication_file(study_dir: Path, capsys):
    differing = _two_reviewer_study(study_dir)
    matrix_before = (study_dir / "human-review-matrix.csv").read_bytes()

    assert _matrix(study_dir) == EXIT_ADJUDICATION
    assert "3 cells have different labels from r1 and r2" in capsys.readouterr().err
    assert (study_dir / "human-review-matrix.csv").read_bytes() == matrix_before
    report = _csv(study_dir / review_merge.DISAGREEMENTS_FILENAME)
    assert [(row["study_doc_id"], row["rule_id"]) for row in report] == sorted(
        differing,
        key=[(r["study_doc_id"], r["rule_id"]) for r in _csv(study_dir / "human-review-matrix.csv")].index,
    )
    assert all(row["final_label"] == "pending" and row["r1_label"] != row["r2_label"] for row in report)
    assert not set(report[0]) & set(DETECTOR_FIELDS)
    assert list(report[0])[-len(ADJUDICATION_FIELDS) :] == list(ADJUDICATION_FIELDS)

    def decision(by: str, when: str = "2026-08-10"):
        def decide(row):
            label = row["r2_label"]
            return {
                "final_label": label,
                "adjudicated_by": by,
                "adjudicated_at": when,
                "final_pinpoint": "synthetic adjudication pinpoint"
                if label in ("present", "missing")
                else "",
                "adjudication_reason": "synthetic adjudication reason",
            }

        return decide

    refusals = {
        "pending": (lambda row: {}, "3 pending"),
        "one original reviewer alone": (decision("r1"), "names only some of the original reviewers"),
        "before the reviews": (decision("r1+r2", "2026-07-21"), "before 2026-07-25"),
        "bad adjudicator id": (decision("Dr. Example"), "use study IDs joined by '+'"),
    }
    for case, (decide, message) in refusals.items():
        path = _adjudication(study_dir, report, decide)
        assert _matrix(study_dir, "--adjudication", str(path)) == EXIT_ADJUDICATION, case
        assert message in capsys.readouterr().err, case
        assert (study_dir / "human-review-matrix.csv").read_bytes() == matrix_before, case

    # Stale: a reviewer changes a note on a disputed cell after the report was copied.
    valid = _adjudication(study_dir, report, decision("r1+r2"))
    r2_rows = _csv(study_dir / "human-review-r2.csv")
    r2_original = [dict(row) for row in r2_rows]
    target = next(row for row in r2_rows if (row["study_doc_id"], row["rule_id"]) == differing[0])
    target["review_note"] = target["review_note"] + " (synthetic edit)"
    if not target["source_pinpoint"]:
        target["source_pinpoint"] = "synthetic edit"
    _write_csv(study_dir / "human-review-r2.csv", r2_rows)
    assert _matrix(study_dir, "--adjudication", str(valid)) == EXIT_ADJUDICATION
    assert "differ from the current disagreement report" in capsys.readouterr().err
    _write_csv(study_dir / "human-review-r2.csv", r2_original)

    third_person = _adjudication(study_dir, report, decision("r3"), bom=True)
    assert _matrix(study_dir, "--adjudication", str(third_person)) == EXIT_OK
    matrix = {
        (row["study_doc_id"], row["rule_id"]): row for row in _csv(study_dir / "human-review-matrix.csv")
    }
    report_by_cell = {(row["study_doc_id"], row["rule_id"]): row for row in report}
    for key, row in matrix.items():
        if key in differing:
            assert row["human_label"] == report_by_cell[key]["r2_label"]
            assert row["reviewer_id"] == "r3"
            assert row["reviewed_at"] == "2026-08-10"
            assert row["review_note"].startswith(
                f"Adjudicated (r1: {report_by_cell[key]['r1_label']}, r2: {report_by_cell[key]['r2_label']}): "
            )
        else:
            assert row["reviewer_id"] == "r1+r2"
            assert row["reviewed_at"] == "2026-07-25"
    summary = json.loads((study_dir / "human-review-summary.json").read_text(encoding="utf-8"))
    assert summary["publication_ready"] is True

    # A different adjudication would change the merged labels: refused unless --replace.
    consensus = _adjudication(study_dir, report, decision("r1+r2"))
    assert _matrix(study_dir, "--adjudication", str(consensus)) == EXIT_OVERWRITE
    assert _matrix(study_dir, "--adjudication", str(consensus), "--replace") == EXIT_OK


def test_two_reviewers_who_agree_need_no_adjudication(study_dir: Path):
    _add_reviewer(study_dir, "r2")
    _label_csv(study_dir / "human-review-r1.csv", "r1")
    r1 = {(row["study_doc_id"], row["rule_id"]): row for row in _csv(study_dir / "human-review-r1.csv")}
    rows = _csv(study_dir / "human-review-r2.csv")
    for index, row in enumerate(rows):
        label = r1[(row["study_doc_id"], row["rule_id"])]["human_label"]
        row.update(_synthetic(index, "r2", label=label))
    _write_csv(study_dir / "human-review-r2.csv", rows)

    assert _matrix(study_dir) == EXIT_OK
    assert _csv(study_dir / review_merge.DISAGREEMENTS_FILENAME) == []
    assert {row["reviewer_id"] for row in _csv(study_dir / "human-review-matrix.csv")} == {"r1+r2"}


def test_committed_study_passes_the_identity_checks():
    # The only test that reads the committed sheets. It holds whatever the labelling progress:
    # load_study regenerates the sheets with the committed reviewers and seed, requires the
    # committed generation record byte for byte, and read_sheet ignores the review columns.
    study = load_study(STUDY)
    assert study.reviewers
    for reviewer in study.reviewers:
        read_sheet(study, reviewer)


def test_ranges():
    assert _ranges([]) == ""
    assert _ranges([7, 2, 3, 4, 9, 10]) == "2-4, 7, 9-10"
