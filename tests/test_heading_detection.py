"""A DOCX paragraph or PDF line opens a section only if it is shaped like a heading.

A body paragraph with a bold lead-in ("Note: ...") in a DOCX, or a short body line in a PDF, that
mentions "risks" used to open a risks section, so a draft with no risk section passed the ANNEX_I.I
blocker. In a PDF, list items, wrapped lines and lead-ins stay in their section, a line shaped like a
heading but not accepted as one ends the section before it, and a bold or larger font marks a heading.
Heading text is not section body, but regime inference still reads it. The PDFs are real files parsed
by pypdf.
"""

from __future__ import annotations

import textwrap
import threading
from html import escape
from pathlib import Path

import docx
import pytest
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from micar_linter.anchors import _xhtml_section_identifiers
from micar_linter.document import load_from_docx, load_from_markdown, load_from_pdf, match_heading_to_section
from micar_linter.linter import lint_whitepaper
from micar_linter.locations import _xhtml_heading_lines
from micar_linter.rules.base import Finding, Severity
from micar_linter.whitepaper import WhitepaperType, load_whitepaper
from micar_linter.xhtml_parser import load_from_xhtml

# An invented issuer. None of the body text below matches a section pattern except the note.
SUMMARY = (
    "The Quillmoor Utility Token (QMT) gives holders access to a shared ledger service for logistics firms."
)
PROJECT = "Quillmoor Ledger S.A. runs the shared ledger service for logistics firms in three Member States."
# Body text, not a heading: in the DOCX only "Note:" is bold, in the PDF it is one short line.
NOTE = "Note: purchasers bear market, liquidity and technology risks."
RISK_DETAIL = (
    "Holders may sell units only where a market exists, and prices can move sharply within a single day. "
    "Liquidity can disappear when few buyers are active, so holders may be unable to exit at a fair price. "
    "A change in regulatory treatment in any Member State could restrict transfers or require new licences. "
    "Operational failures at the service provider, such as outages or faulty software updates, can delay "
    "transfers for days. Keys held by holders can be lost or stolen, and lost keys cannot be recovered by "
    "anyone. Smart contract defects can freeze balances until a fix is deployed and accepted by the network. "
    "Holders should read this document in full and seek independent advice before they buy units. Past "
    "performance of similar services does not indicate future returns, and holders may lose the full amount."
)
TECHNOLOGY = "Transfers settle on a permissioned ledger operated by five validating nodes."


def _draft(*risk_heading: tuple[str, str]) -> list[tuple[str, str]]:
    """(kind, text) blocks; the risk discussion sits under the project unless a risk heading is given."""
    return [
        ("heading", "Summary"),
        ("body", SUMMARY),
        ("heading", "Information about the Project"),
        ("body", PROJECT),
        *risk_heading,
        ("note", NOTE),
        ("body", RISK_DETAIL),
        ("heading", "Underlying Technology"),
        ("body", TECHNOLOGY),
    ]


def _write_docx(path: Path, blocks: list[tuple[str, str]]) -> Path:
    doc = docx.Document()
    for kind, text in blocks:
        if kind == "title":
            doc.add_paragraph(text, style="Title")
        elif kind in ("heading", "large"):
            doc.add_heading(text, level=1)
        elif kind == "subheading":
            doc.add_heading(text, level=2)
        elif kind in ("bold", "bold body"):
            doc.add_paragraph().add_run(text).bold = True
        elif kind == "outline":
            # A Normal paragraph that Word lists in the navigation pane as a level 1 heading.
            level = OxmlElement("w:outlineLvl")
            level.set(qn("w:val"), "0")
            paragraph = doc.add_paragraph(text)
            paragraph._p.get_or_add_pPr().append(level)
        elif kind == "note":
            label, rest = text.split(" ", 1)
            paragraph = doc.add_paragraph()
            paragraph.add_run(label).bold = True
            paragraph.add_run(f" {rest}")
        elif kind == "numbered":
            # A number typed in plain type, then the title in bold.
            number, title = text.split(" ", 1)
            paragraph = doc.add_paragraph()
            paragraph.add_run(f"{number}\t")
            paragraph.add_run(title).bold = True
        elif kind == "link":
            # A bold title that is also an internal link, so its run sits inside w:hyperlink.
            run = OxmlElement("w:r")
            properties = OxmlElement("w:rPr")
            properties.append(OxmlElement("w:b"))
            run.append(properties)
            run_text = OxmlElement("w:t")
            run_text.text = text
            run.append(run_text)
            link = OxmlElement("w:hyperlink")
            link.set(qn("w:anchor"), "risks")
            link.append(run)
            doc.add_paragraph()._p.append(link)
        else:
            doc.add_paragraph(text)
    doc.save(path)
    return path


# 10 pt Helvetica for body text and plain lines; Helvetica-Bold in 10 pt for "bold" lines and "bold body"
# text, 11 pt for "subheading" lines, 14 pt for "large" ones and 20 pt for a "title". In a DOCX, "large" is
# a Heading 1, "subheading" a Heading 2 and "title" a Title-style paragraph.
PDF_FONTS = {
    "bold": "/F2 10 Tf",
    "bold body": "/F2 10 Tf",
    "subheading": "/F2 11 Tf",
    "large": "/F2 14 Tf",
    "title": "/F2 20 Tf",
}


def _shown(kind: str, line: str, next_line: str) -> str:
    """The content stream operators of one block line; a "space" block only moves 20 pt further down."""
    if kind == "space":
        return "0 -20 Td"
    escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    return f"{PDF_FONTS.get(kind, '/F1 10 Tf')} ({escaped}) Tj {next_line}"


def _write_pdf(path: Path, blocks: list[tuple[str, str]], next_line: str = "0 -13 Td") -> Path:
    """A PDF with one text line per block, 13 pt apart, and body text wrapped at 90 characters.

    The lines of a paragraph and the paragraphs follow each other with no space between them, as in a
    draft that sets none, unless a "space" block moves further down. Each line moves to the next with
    next_line: Td, as word processors write, or T* after a TL leading. A "page" block starts a new page.
    """
    pages: list[list[tuple[str, str]]] = [[]]
    for kind, text in blocks:
        if kind == "page":
            pages.append([])
        else:
            wrapped = textwrap.wrap(text, 90) if kind in ("body", "bold body") else [text]
            pages[-1].extend((kind, line) for line in wrapped)
    fonts = b"/Font << /F1 3 0 R /F2 4 0 R >>"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [%s] /Count %d >>"
        % (b" ".join(b"%d 0 R" % (5 + 2 * number) for number in range(len(pages))), len(pages)),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>",
    ]
    for number, lines in enumerate(pages):
        shown = " ".join(_shown(kind, line, next_line) for kind, line in lines)
        content = f"BT 13 TL 50 770 Td {shown} ET".encode("cp1252")
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 1000 792] /Resources << %s >> /Contents %d 0 R >>"
            % (fonts, 6 + 2 * number)
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(pdf)
    pdf += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    pdf += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    pdf += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    path.write_bytes(bytes(pdf))
    return path


WRITERS = {".docx": (_write_docx, load_from_docx), ".pdf": (_write_pdf, load_from_pdf)}


def _risks_finding(path: Path) -> Finding:
    report = lint_whitepaper(load_whitepaper(path, WhitepaperType.OTHER))
    return next(finding for finding in report.findings if finding.rule.rule_id == "ANNEX_I.I")


def _words(text: str) -> str:
    return " ".join(text.split())


def test_docx_bold_lead_in_paragraph_does_not_open_a_risk_section(tmp_path: Path):
    path = _write_docx(tmp_path / "quillmoor.docx", _draft())

    finding = _risks_finding(path)
    assert finding.rule.severity is Severity.BLOCKER
    assert finding.status == "missing"
    sections = load_from_docx(path)
    assert "risks" not in sections
    assert NOTE in sections["project"]


def test_pdf_body_line_mentioning_risks_does_not_open_a_risk_section(tmp_path: Path):
    path = _write_pdf(tmp_path / "quillmoor.pdf", _draft())

    finding = _risks_finding(path)
    assert finding.rule.severity is Severity.BLOCKER
    assert finding.status == "missing"
    sections = load_from_pdf(path)
    assert "risks" not in sections
    assert NOTE in sections["project"]


@pytest.mark.parametrize(
    "kind,heading",
    [("heading", "Risks"), ("bold", "Risk Factors"), ("body", "RISK FACTORS"), ("outline", "Risks")],
)
def test_docx_risk_heading_still_opens_the_risk_section(tmp_path: Path, kind: str, heading: str):
    path = _write_docx(tmp_path / "quillmoor.docx", _draft((kind, heading)))

    assert _words(RISK_DETAIL) in _words(load_from_docx(path)["risks"])
    assert _risks_finding(path).status == "pass"


@pytest.mark.parametrize(
    "heading", ["Risks", "3. Risk Factors", "RISK FACTORS", "Part I: Information on Risks"]
)
def test_pdf_risk_heading_still_opens_the_risk_section(tmp_path: Path, heading: str):
    path = _write_pdf(tmp_path / "quillmoor.pdf", _draft(("heading", heading)))

    assert _words(RISK_DETAIL) in _words(load_from_pdf(path)["risks"])
    assert _risks_finding(path).status == "pass"


@pytest.mark.parametrize("suffix", sorted(WRITERS))
def test_heading_text_is_not_section_body(tmp_path: Path, suffix: str):
    write, load = WRITERS[suffix]
    blocks = [
        ("heading", "Summary"),
        ("body", SUMMARY),
        ("heading", "Risks"),
        ("heading", "Underlying Technology"),
        ("body", TECHNOLOGY),
    ]
    path = write(tmp_path / f"quillmoor{suffix}", blocks)

    # A risks heading with nothing under it leaves the section empty.
    finding = _risks_finding(path)
    assert (finding.status, finding.word_count) == ("missing", 0)
    sections = load(path)
    assert "risks" not in sections
    assert _words(sections["summary"]) == SUMMARY
    assert _words(sections["technology"]) == TECHNOLOGY


# The part titles of Implementing Regulation (EU) 2024/2984 separate the part letter with an en dash.
DASH = chr(0x2013)
# A risk discussion that passes the ANNEX_I.I blocker wherever it is credited to the risks section.
RISK_TEXT = RISK_DETAIL + " Technology failures at the ledger can also stop transfers for days."
THIN_SUMMARY = "This summary gives the key information on the QMT token."


def _findings(
    path: Path, whitepaper_type: WhitepaperType | None = WhitepaperType.OTHER
) -> dict[str, Finding]:
    report = lint_whitepaper(load_whitepaper(path, whitepaper_type))
    return {finding.rule.rule_id: finding for finding in report.findings}


@pytest.mark.parametrize(
    "heading", ["Risk factors", "Beschreibung der wesentlichen Risiken", "5. Risks relating to the token"]
)
def test_pdf_line_rejected_as_a_heading_ends_the_section_before_it(tmp_path: Path, heading: str):
    # Shaped like a heading but not in title case: the risk text below it must not reach the summary.
    blocks = [
        ("heading", "Summary"),
        ("body", THIN_SUMMARY),
        ("heading", heading),
        ("body", RISK_TEXT),
        ("heading", "Underlying Technology"),
        ("body", TECHNOLOGY),
    ]
    path = _write_pdf(tmp_path / "quillmoor.pdf", blocks)

    sections = load_from_pdf(path)
    assert sorted(sections) == ["summary", "technology"]
    assert _words(sections["summary"]) == THIN_SUMMARY
    findings = _findings(path)
    assert findings["COMMON.SUMMARY"].rule.severity is Severity.BLOCKER
    assert (findings["COMMON.SUMMARY"].status, findings["COMMON.SUMMARY"].word_count) == ("review", 10)
    assert findings["ANNEX_I.I"].status == "missing"


@pytest.mark.parametrize("kind", ["bold", "large"])
@pytest.mark.parametrize("heading", ["Risk factors", "5. Risks relating to the token"])
def test_pdf_heading_font_marks_a_sentence_case_heading(tmp_path: Path, kind: str, heading: str):
    blocks = [("large", "Summary"), ("body", SUMMARY), (kind, heading), ("body", RISK_TEXT)]
    path = _write_pdf(tmp_path / "quillmoor.pdf", blocks)

    sections = load_from_pdf(path)
    assert "risks" in sections
    assert _words(sections["risks"]) == _words(RISK_TEXT)
    assert _findings(path)["ANNEX_I.I"].status == "pass"


# Lines in the project section that name risks without being a heading, and whether the risk text after
# them stays in the project section (body) or is credited to no section (the line ends the section).
PROJECT_LINES = {
    "numbered list item": (
        [
            "Key features of the project:",
            "1. Shared ledger for logistics firms",
            "2. Low fees compared with legacy systems",
            "3. Exposure to market, liquidity and technology risks",
        ],
        True,
    ),
    "bullet item": (["Holders should consider:", "• Market and Technology Risks"], True),
    "dash item": (["- Market and Technology Risks"], True),
    "wrapped capitals": (
        [
            "PURCHASERS SHOULD CAREFULLY CONSIDER THE MARKET, LIQUIDITY AND",
            "TECHNOLOGY RISKS DESCRIBED IN THIS DOCUMENT BEFORE THEY BUY ANY UNITS",
            "OF THE TOKEN.",
        ],
        True,
    ),
    "wrapped german nouns": (
        [
            "Quillmoor betreibt den Dienst in drei Mitgliedstaaten. Anleger sollten die",
            "Risiken der Token und der Plattform",
            "beachten, bevor sie Einheiten kaufen.",
        ],
        True,
    ),
    "capitals line running on": (
        ["HOLDERS BEAR TECHNOLOGY RISKS AND", "MAY LOSE THE FULL AMOUNT INVESTED."],
        False,
    ),
    "sentence running on after a german abbreviation": (
        ["Anleger tragen die Risiken", "z. B. Markt- und Liquiditätsrisiken des Tokens."],
        True,
    ),
    "sentence-case lead-in": (["Note on the risks:"], True),
    "title-case lead-in": (["Holders Should Note These Risks:"], False),
    "title-case numbered list": (
        [
            "The project has three features.",
            "1. Shared Ledger.",
            "2. Low Fees.",
            "3. Market and Technology Risks",
        ],
        False,
    ),
}


@pytest.mark.parametrize("case", sorted(PROJECT_LINES))
def test_pdf_list_and_wrapped_lines_open_no_section(tmp_path: Path, case: str):
    lines, stays_in_project = PROJECT_LINES[case]
    blocks = [
        ("heading", "Summary"),
        ("body", SUMMARY),
        ("heading", "Information about the Project"),
        ("body", PROJECT),
        *(("line", line) for line in lines),
        ("body", RISK_TEXT),
        ("heading", "Underlying Technology"),
        ("body", TECHNOLOGY),
    ]
    path = _write_pdf(tmp_path / "quillmoor.pdf", blocks)

    sections = load_from_pdf(path)
    assert sorted(sections) == ["project", "summary", "technology"]
    assert (_words(RISK_TEXT) in _words(sections["project"])) is stays_in_project
    findings = _findings(path)
    assert findings["ANNEX_I.I"].rule.severity is Severity.BLOCKER
    assert findings["ANNEX_I.I"].status == "missing"


# Part titles of the Implementing Regulation as bold Normal paragraphs, each with up to 20 words.
RIGHTS = "Holders have rights and obligations; each holder may enforce rights in court under the terms. " * 6
OFFER = (
    "The offer to the public runs from 2 March 2026 at a price of 1 EUR per token in three Member States. "
    * 4
)
ISSUER = (
    "The issuer is Quillmoor Ledger S.A., registered at 4 Rue Quillmoor, Luxembourg, as a public company. "
    * 3
)
PART_TITLES = {
    "B": (f"Part B {DASH} Information about the issuer, if different from the offeror or person seeking admission "
          "to trading", "issuer", ISSUER, "ANNEX_I.B", "review"),
    "E": (f"Part E {DASH} Information about the offer to the public of crypto-assets or their admission to "
          "trading", "offer_or_admission", OFFER, "ANNEX_I.E", "review"),
    "G": (f"Part G {DASH} Information on the rights and obligations attached to the crypto-assets",
          "rights_and_obligations", RIGHTS, "ANNEX_I.G", "pass"),
}  # fmt: skip


@pytest.mark.parametrize("part", sorted(PART_TITLES))
def test_docx_bold_part_title_of_the_template_is_a_heading(tmp_path: Path, part: str):
    title, key, text, rule_id, status = PART_TITLES[part]
    blocks = [
        ("heading", "Summary"),
        ("body", SUMMARY),
        ("bold", f"Part D {DASH} Information about the crypto-asset project"),
        ("body", PROJECT),
        ("bold", title),
        ("body", text),
        ("bold", f"Part H {DASH} Information on the underlying technology"),
        ("body", TECHNOLOGY),
    ]
    path = _write_docx(tmp_path / "quillmoor.docx", blocks)

    sections = load_from_docx(path)
    assert _words(sections[key]) == _words(text)
    assert _words(sections["project"]) == PROJECT
    findings = _findings(path)
    assert (findings[rule_id].status, findings[rule_id].word_count) == (status, len(text.split()))
    assert findings["ANNEX_I.D"].word_count == len(PROJECT.split())


# The same part titles without the part letter, as many drafts set them: 13 to 16 words, so only their
# bold type (DOCX) or heading font (PDF) marks them as headings.
GERMAN_OFFEROR = (
    "Der Anbieter ist die Quillmoor Ledger S.A. mit Sitz in Luxemburg. Sie ist im Handelsregister unter der "
    "Nummer B 123456 eingetragen und beschäftigt vierzig Mitarbeiter in Luxemburg und Hamburg. "
) * 2
UNNUMBERED_TITLES = {
    "B": ("Information about the issuer, if different from the offeror or person seeking admission to trading",
          ["issuer", "offer_or_admission", "offeror"], ISSUER, "ANNEX_I.B", Severity.MAJOR, "review"),
    "E": ("Information about the offer to the public of crypto-assets or their admission to trading",
          ["offer_or_admission"], OFFER, "ANNEX_I.E", Severity.BLOCKER, "review"),
    "A-de": ("Informationen über den Anbieter oder die Person, die die Zulassung zum Handel beantragt",
             ["offer_or_admission", "offeror"], GERMAN_OFFEROR, "ANNEX_I.A", Severity.BLOCKER, "review"),
}  # fmt: skip


@pytest.mark.parametrize(
    "suffix,kind", [(".docx", "bold"), (".pdf", "large")], ids=["docx-bold", "pdf-14pt-bold"]
)
@pytest.mark.parametrize("part", sorted(UNNUMBERED_TITLES))
def test_unnumbered_part_title_of_up_to_twenty_words_is_a_heading_when_marked(
    tmp_path: Path, suffix: str, kind: str, part: str
):
    title, keys, text, rule_id, severity, status = UNNUMBERED_TITLES[part]
    write, load = WRITERS[suffix]
    blocks = [
        (kind, "Summary"),
        ("body", SUMMARY),
        (kind, "Information about the crypto-asset project"),
        ("body", PROJECT),
        (kind, title),
        ("body", text),
        (kind, "Information on the underlying technology"),
        ("body", TECHNOLOGY),
    ]
    path = write(tmp_path / f"quillmoor{suffix}", blocks)

    sections = load(path)
    assert sorted(sections) == sorted({"summary", "project", "crypto_asset", "technology", *keys})
    for key in keys:
        assert _words(sections[key]) == _words(text)
    # The thin project section keeps its own text and does not pass on the part after it.
    assert _words(sections["project"]) == PROJECT
    findings = _findings(path)
    assert findings["ANNEX_I.D"].rule.severity is Severity.BLOCKER
    assert (findings["ANNEX_I.D"].status, findings["ANNEX_I.D"].word_count) == (
        "review",
        len(PROJECT.split()),
    )
    assert findings[rule_id].rule.severity is severity
    assert (findings[rule_id].status, findings[rule_id].word_count) == (status, len(text.split()))


# Headings in a 14 pt bold heading font that wrap onto a second line, as LibreOffice and Chrome export the
# long part titles: (first line, second line, sections, section text, rule, status of that rule).
WRAPPED_HEADINGS = {
    "numbered, lower case": (
        "5. Information about the offer to the public of crypto-assets or their",
        "admission to trading",
        ["offer_or_admission"], OFFER, "ANNEX_I.E", "review",
    ),
    "unnumbered, lower case": (
        "Information about the offer to the public of crypto-assets or their",
        "admission to trading",
        ["offer_or_admission"], OFFER, "ANNEX_I.E", "review",
    ),
    "title case, hyphen": (
        "7. Information on the Rights and Obligations Attached to the Crypto-",
        "Assets",
        ["rights_and_obligations"], RIGHTS, "ANNEX_I.G", "pass",
    ),
    "title case, capital": (
        f"Part B {DASH} Information About the Issuer, If Different From the Offeror or Person",
        "Seeking Admission to Trading",
        ["issuer", "offer_or_admission", "offeror"], ISSUER, "ANNEX_I.B", "review",
    ),
    "german": (
        "6. Informationen über den Anbieter oder die Person, die die Zulassung zum",
        "Handel beantragt",
        ["offer_or_admission", "offeror"], GERMAN_OFFEROR, "ANNEX_I.A", "review",
    ),
}  # fmt: skip


@pytest.mark.parametrize("next_line", ["0 -13 Td", "T*"])
@pytest.mark.parametrize("case", sorted(WRAPPED_HEADINGS))
def test_pdf_heading_wrapped_onto_two_lines_in_the_heading_font_is_one_heading(
    tmp_path: Path, case: str, next_line: str
):
    first, second, keys, text, rule_id, status = WRAPPED_HEADINGS[case]
    blocks = [
        ("large", "Summary"),
        ("body", SUMMARY),
        ("large", "4. Information about the crypto-asset project"),
        ("body", PROJECT),
        ("large", first),
        ("large", second),
        ("body", text),
        ("large", "8. Information on the underlying technology"),
        ("body", TECHNOLOGY),
    ]
    # pypdf misreports the heights of lines laid out with T*, so there the font alone joins the lines.
    path = _write_pdf(tmp_path / "quillmoor.pdf", blocks, next_line)

    sections = load_from_pdf(path)
    assert sorted(sections) == sorted({"summary", "project", "crypto_asset", "technology", *keys})
    for key in keys:
        assert _words(sections[key]) == _words(text)
    # The thin project section before the wrapped heading does not pass on the part under it.
    assert _words(sections["project"]) == PROJECT
    findings = _findings(path)
    assert findings["ANNEX_I.D"].rule.severity is Severity.BLOCKER
    assert (findings["ANNEX_I.D"].status, findings["ANNEX_I.D"].word_count) == (
        "review",
        len(PROJECT.split()),
    )
    assert (findings[rule_id].status, findings[rule_id].word_count) == (status, len(text.split()))


def test_pdf_headings_one_paragraph_apart_in_the_heading_font_stay_two_headings(tmp_path: Path):
    # A heading with nothing under it, then the next heading: the space between paragraphs keeps them
    # apart, so the risks section stays empty instead of passing on the text under the next heading
    # (risk-like text here, so that a merged heading would pass ANNEX_I.I).
    blocks = [
        ("large", "Summary"),
        ("body", SUMMARY),
        ("large", "Risk Factors"),
        ("space", ""),
        ("large", "Information on the Underlying Technology"),
        ("body", RISK_TEXT),
    ]
    path = _write_pdf(tmp_path / "quillmoor.pdf", blocks)

    sections = load_from_pdf(path)
    assert sorted(sections) == ["summary", "technology"]
    assert _words(sections["technology"]) == _words(RISK_TEXT)
    findings = _findings(path)
    assert findings["ANNEX_I.I"].rule.severity is Severity.BLOCKER
    assert findings["ANNEX_I.I"].status == "missing"


@pytest.mark.parametrize("next_line", ["0 -13 Td", "T*"])
def test_pdf_short_heading_one_line_above_the_next_heading_stays_a_heading(tmp_path: Path, next_line: str):
    # As above, but with no space between the paragraphs, and in a PDF laid out with T*, where the font
    # alone decides: "Risk Factors" neither runs on nor fills the column, so it is not joined with the
    # heading below it.
    blocks = [
        ("large", "Summary"),
        ("body", SUMMARY),
        ("large", "Risk Factors"),
        ("large", "Information on the Underlying Technology"),
        ("body", RISK_TEXT),
    ]
    path = _write_pdf(tmp_path / "quillmoor.pdf", blocks, next_line)

    sections = load_from_pdf(path)
    assert sorted(sections) == ["summary", "technology"]
    assert _words(sections["technology"]) == _words(RISK_TEXT)
    findings = _findings(path)
    assert findings["ANNEX_I.I"].rule.severity is Severity.BLOCKER
    assert findings["ANNEX_I.I"].status == "missing"


def test_pdf_short_heading_over_a_lower_case_line_in_its_font_stays_a_heading(tmp_path: Path):
    # A heading that wraps after a short first line, as in a narrow column. "Risk Factors" neither runs on
    # nor fills the column, so it is not joined with "and Mitigation Measures"; it does not start a
    # paragraph either, which fills the column. So it opens the risks, and the risk text below does not
    # pass the thin technology section before it.
    blocks = [
        ("large", "Summary"),
        ("body", SUMMARY),
        ("large", "Information on the Underlying Technology"),
        ("body", THIN_TECHNOLOGY),
        ("large", "Risk Factors"),
        ("large", "and Mitigation Measures"),
        ("body", RISK_TEXT),
    ]
    path = _write_pdf(tmp_path / "quillmoor.pdf", blocks)

    sections = load_from_pdf(path)
    assert sorted(sections) == ["risks", "summary", "technology"]
    assert _words(sections["technology"]) == _words(THIN_TECHNOLOGY)
    assert _words(RISK_TEXT) in _words(sections["risks"])
    findings = _findings(path)
    assert findings["ANNEX_I.H"].rule.severity is Severity.BLOCKER
    assert (findings["ANNEX_I.H"].status, findings["ANNEX_I.H"].word_count) == (
        "review",
        len(THIN_TECHNOLOGY.split()),
    )
    assert findings["ANNEX_I.I"].status == "pass"


# The opening statements of a white paper, set under bold headings in the body size. The summary alone is
# too thin to pass COMMON.SUMMARY.
KEY_SUMMARY = (
    "This summary gives the key information about the Quillmoor Utility Token (QMT) in plain language. The "
    "token gives holders access to a shared ledger service for logistics firms. Holders should read the "
    "whole white paper before they buy units."
)
WARNING = (
    "This crypto-asset white paper has not been reviewed or approved by any competent authority in any "
    "Member State of the European Union. The offeror of the crypto-asset is solely responsible for the "
    "content of this crypto-asset white paper. The crypto-asset may lose its value in part or in full, may "
    "not always be transferable and may not be liquid."
)
STATEMENT = (
    "The management body of Quillmoor Ledger S.A. confirms that this crypto-asset white paper complies "
    "with Title II of Regulation (EU) 2023/1114 and that, to the best of its knowledge, the information "
    "presented is fair, clear and not misleading and that the white paper makes no omission likely to "
    "affect its import."
)


@pytest.mark.parametrize("suffix", sorted(WRITERS))
@pytest.mark.parametrize("bold_statement", [False, True], ids=["bold-warning", "bold-warning-and-statement"])
def test_heading_over_a_bold_paragraph_with_no_space_between_stays_a_heading(
    tmp_path: Path, suffix: str, bold_statement: bool
):
    # With no space between paragraphs, the bold headings sit one line spacing above and below the bold
    # risk warning. They are not joined with it: each heading opens its own section and the bold paragraph
    # is the body of its heading, so the thin summary does not pass on the statements after it.
    write, load = WRITERS[suffix]
    blocks = [
        ("bold", "Summary"),
        ("body", KEY_SUMMARY),
        ("bold", "Risk warning"),
        ("bold body", WARNING),
        ("bold", "Statement by the management body"),
        ("bold body" if bold_statement else "body", STATEMENT),
        ("bold", "Information on risks"),
        ("body", RISK_TEXT),
    ]
    path = write(tmp_path / f"quillmoor{suffix}", blocks)

    sections = load(path)
    assert sorted(sections) == ["management_statement", "risk_warning", "risks", "summary"]
    assert _words(sections["summary"]) == KEY_SUMMARY
    assert _words(sections["risk_warning"]) == _words(WARNING)
    assert _words(sections["management_statement"]) == _words(STATEMENT)
    findings = _findings(path)
    for rule_id, text in [
        ("COMMON.SUMMARY", KEY_SUMMARY),
        ("COMMON.RISK_WARNING", WARNING),
        ("COMMON.MANAGEMENT_STATEMENT", STATEMENT),
    ]:
        assert findings[rule_id].rule.severity is Severity.BLOCKER
        assert findings[rule_id].word_count == len(text.split())
    assert findings["COMMON.SUMMARY"].status == "review"
    assert findings["COMMON.MANAGEMENT_STATEMENT"].status == "pass"
    assert findings["ANNEX_I.I"].status == "pass"


# A risk section written as a list whose items start in lower case.
RISK_ITEMS = [
    "a) Market risk: prices can move sharply within a single day and holders may lose the full amount.",
    "b) Liquidity risk: liquidity can disappear when few buyers are active, so holders may be unable to exit.",
    "c) Regulatory risk: a change in regulatory treatment in any Member State could restrict transfers.",
    "d) Operational risk: outages or faulty software updates at the service provider can delay transfers.",
    "e) Software risk: code defects can freeze balances until a fix is deployed and accepted.",
    "f) Custody risk: keys held by holders can be lost or stolen, and lost keys cannot be recovered by anyone.",
]
THIN_TECHNOLOGY = (
    "Transfers settle on a permissioned ledger with a Byzantine fault tolerant consensus protocol. Smart "
    "contracts implement the unit register, and an independent firm audits the code once a year."
)


@pytest.mark.parametrize(
    "kind,technology,heading",
    [
        ("large", "4. Underlying Technology", "Risk Factors"),
        ("large", "4. Technologie", "5. Risiken"),
        # Without a heading font the shape decides, and a list item after a line does not make it body text.
        ("heading", "Underlying Technology", "Risk Factors"),
    ],
    ids=["14pt-bold", "14pt-bold-german", "body-font"],
)
def test_pdf_heading_before_a_lower_case_list_item_is_a_heading(
    tmp_path: Path, kind: str, technology: str, heading: str
):
    blocks = [
        (kind, "Summary"),
        ("body", SUMMARY),
        (kind, technology),
        ("body", THIN_TECHNOLOGY),
        (kind, heading),
        *(("line", item) for item in RISK_ITEMS),
    ]
    path = _write_pdf(tmp_path / "quillmoor.pdf", blocks)

    sections = load_from_pdf(path)
    assert sorted(sections) == ["risks", "summary", "technology"]
    assert _words(sections["technology"]) == _words(THIN_TECHNOLOGY)
    assert _words(sections["risks"]) == _words(" ".join(RISK_ITEMS))
    findings = _findings(path)
    # The thin technology section does not pass on the risk list, and the risk section is not missing.
    assert findings["ANNEX_I.H"].rule.severity is Severity.BLOCKER
    assert (findings["ANNEX_I.H"].status, findings["ANNEX_I.H"].word_count) == (
        "review",
        len(THIN_TECHNOLOGY.split()),
    )
    assert findings["ANNEX_I.I"].rule.severity is Severity.BLOCKER
    assert findings["ANNEX_I.I"].status != "missing"


def test_pdf_line_running_on_into_a_lower_case_line_is_read_with_it_by_shape(tmp_path: Path):
    # Without a heading font: "or" runs on into "their admission to trading", and the two lines together
    # are the numbered template title of Part E.
    blocks = [
        ("heading", "Summary"),
        ("body", SUMMARY),
        ("line", f"Part D {DASH} Information about the crypto-asset project"),
        ("body", PROJECT),
        ("line", f"Part E {DASH} Information about the offer to the public of crypto-assets or"),
        ("line", "their admission to trading"),
        ("body", OFFER),
    ]
    path = _write_pdf(tmp_path / "quillmoor.pdf", blocks)

    sections = load_from_pdf(path)
    assert sorted(sections) == ["crypto_asset", "offer_or_admission", "project", "summary"]
    assert _words(sections["offer_or_admission"]) == _words(OFFER)
    findings = _findings(path)
    assert (findings["ANNEX_I.D"].status, findings["ANNEX_I.D"].word_count) == (
        "review",
        len(PROJECT.split()),
    )
    assert findings["ANNEX_I.E"].rule.severity is Severity.BLOCKER
    assert (findings["ANNEX_I.E"].status, findings["ANNEX_I.E"].word_count) == ("review", len(OFFER.split()))


def test_pdf_part_title_of_twenty_words_and_a_dash_is_a_heading(tmp_path: Path):
    part_j = (
        f"Part J {DASH} Information on the sustainability indicators in relation to adverse impact on the climate "
        "and other environment-related adverse impacts"
    )
    environment = (
        "The consensus mechanism uses little energy, so climate and environmental impacts are low. " * 8
    )
    blocks = [
        ("heading", f"Part I {DASH} Information on risks"),
        ("body", RISK_TEXT),
        ("heading", part_j),
        ("body", environment),
    ]
    path = _write_pdf(tmp_path / "quillmoor.pdf", blocks)

    sections = load_from_pdf(path)
    assert sorted(sections) == ["environmental_impact", "risks"]
    assert _words(sections["risks"]) == _words(RISK_TEXT)
    findings = _findings(path)
    assert (findings["ANNEX_I.I"].status, findings["ANNEX_I.J"].status) == ("pass", "pass")


@pytest.mark.parametrize("suffix", sorted(WRITERS))
def test_regime_is_inferred_from_heading_text(tmp_path: Path, suffix: str):
    write, _ = WRITERS[suffix]
    blocks = [
        ("heading", "Summary"),
        ("body", "The QAR token is described here in plain language for retail holders."),
        ("heading", "Reserve of Assets"),
        ("body", "The reserve is held at two credit institutions in cash and bonds."),
        ("heading", "Redemption"),
        ("body", "Holders may redeem at any time."),
    ]
    path = write(tmp_path / f"quillmoor{suffix}", blocks)

    # Only the headings name the regime; with no stated regime the draft is still read as an ART.
    whitepaper = load_whitepaper(path)
    assert whitepaper.type is WhitepaperType.ART
    assert sorted(whitepaper.sections) == ["redemption", "reserve_of_assets", "summary"]
    findings = _findings(path, None)
    assert findings["ANNEX_II.I"].rule.severity is Severity.BLOCKER
    assert findings["ANNEX_II.I"].status == "review"
    assert findings["ANNEX_II.G.COMPOSITION_DISCLOSURE"].status == "review"


@pytest.mark.parametrize("suffix", sorted(WRITERS))
def test_regime_is_inferred_from_text_in_no_section(tmp_path: Path, suffix: str):
    # Only the introduction before the first heading names the regime, and it belongs to no section.
    write, _ = WRITERS[suffix]
    blocks = [
        (
            "body",
            "This white paper describes the QAR, an asset-referenced token issued by Quillmoor Ledger S.A.",
        ),
        ("heading", "Summary"),
        ("body", "The QAR token is described here in plain language for retail holders."),
        ("heading", "Redemption"),
        ("body", "Holders may redeem at any time."),
    ]
    path = write(tmp_path / f"quillmoor{suffix}", blocks)

    whitepaper = load_whitepaper(path)
    assert whitepaper.type is WhitepaperType.ART
    assert sorted(whitepaper.sections) == ["redemption", "summary"]
    # The Annex II BLOCKERs run: the redemption check reviews its thin section.
    findings = _findings(path, None)
    assert findings["ANNEX_II.I"].rule.severity is Severity.BLOCKER
    assert (findings["ANNEX_II.I"].status, findings["ANNEX_II.I"].word_count) == ("review", 6)


@pytest.mark.parametrize(
    "field,answer,key",
    [
        (
            "01 Date of notification",
            "2 March 2026, the date of notification to the BaFin.",
            "notification_date",
        ),
        (
            "08 Characteristics of the crypto-asset",
            "QMT is a utility token on a permissioned ledger.",
            "crypto_asset",
        ),
    ],
    ids=["01", "08"],
)
def test_pdf_two_digit_summary_field_opens_its_section(tmp_path: Path, field: str, answer: str, key: str):
    blocks = [
        ("heading", "Summary"),
        ("body", SUMMARY),
        ("line", field),
        ("line", answer),
        ("heading", f"Part I {DASH} Information on risks"),
        ("body", RISK_TEXT),
    ]
    path = _write_pdf(tmp_path / "quillmoor.pdf", blocks)

    sections = load_from_pdf(path)
    assert sorted(sections) == sorted(["risks", "summary", key])
    assert sections[key] == answer
    assert _words(sections["summary"]) == SUMMARY
    findings = _findings(path)
    assert findings["COMMON.NOTIFICATION_DATE"].status == (
        "pass" if key == "notification_date" else "missing"
    )
    assert findings["ANNEX_I.I"].status == "pass"


@pytest.mark.parametrize("kind", ["numbered", "link"])
def test_docx_bold_heading_with_a_plain_number_or_a_link_is_a_heading(tmp_path: Path, kind: str):
    heading = "5. Risk Factors" if kind == "numbered" else "Risk Factors"
    blocks = [("heading", "Summary"), ("body", SUMMARY), (kind, heading), ("body", RISK_TEXT)]
    path = _write_docx(tmp_path / "quillmoor.docx", blocks)

    sections = load_from_docx(path)
    assert sorted(sections) == ["risks", "summary"]
    assert _findings(path)["ANNEX_I.I"].status == "pass"


@pytest.mark.parametrize("suffix,kind", [(".docx", "bold"), (".pdf", "line")], ids=["docx-bold", "pdf-line"])
@pytest.mark.parametrize(
    "heading",
    ["5. Risks (incl. Liquidity Risks)", "Quillmoor Corp. Risk Factors", "Risiken bzgl. des Tokens"],
)
def test_abbreviation_in_a_heading_is_not_a_sentence_break(
    tmp_path: Path, suffix: str, kind: str, heading: str
):
    write, load = WRITERS[suffix]
    path = write(
        tmp_path / f"quillmoor{suffix}",
        [("heading", "Summary"), ("body", SUMMARY), (kind, heading), ("body", RISK_TEXT)],
    )

    assert sorted(load(path)) == ["risks", "summary"]
    assert _findings(path)["ANNEX_I.I"].status == "pass"


# Part F: its official title names the crypto-assets in the plural. The plural in the titles of Part E
# and Part G does not make them crypto-asset sections.
CRYPTO_ASSETS = (
    "The QMT units are utility tokens on a permissioned ledger. Each unit entitles its holder to record one "
    "hundred transfers. The units are fungible and transferable between onboarded holders. "
) * 3


@pytest.mark.parametrize("suffix", sorted(WRITERS))
def test_part_f_title_with_crypto_assets_in_the_plural_opens_its_section(tmp_path: Path, suffix: str):
    write, load = WRITERS[suffix]
    blocks = [
        ("large", "Summary"),
        ("body", SUMMARY),
        (
            "large",
            f"Part E {DASH} Information about the offer to the public of crypto-assets or their admission",
        ),
        ("body", OFFER),
        ("large", f"Part F {DASH} Information about the crypto-assets"),
        ("body", CRYPTO_ASSETS),
        ("large", f"Part G {DASH} Information on the rights and obligations attached to the crypto-assets"),
        ("body", RIGHTS),
    ]
    path = write(tmp_path / f"quillmoor{suffix}", blocks)

    sections = load(path)
    assert sorted(sections) == ["crypto_asset", "offer_or_admission", "rights_and_obligations", "summary"]
    assert _words(sections["crypto_asset"]) == _words(CRYPTO_ASSETS)
    assert _words(sections["offer_or_admission"]) == _words(OFFER)
    findings = _findings(path)
    assert findings["ANNEX_I.E"].rule.severity is Severity.BLOCKER
    assert (findings["ANNEX_I.E"].status, findings["ANNEX_I.E"].word_count) == ("review", len(OFFER.split()))
    assert findings["ANNEX_I.F"].rule.severity is Severity.BLOCKER
    assert findings["ANNEX_I.F"].word_count == len(CRYPTO_ASSETS.split())


THIN_RISKS = "Holders may lose the full amount they paid for their units."


@pytest.mark.parametrize("suffix", sorted(WRITERS))
@pytest.mark.parametrize(
    "heading", ["Glossary", f"Part K {DASH} Glossary"], ids=["heading-font", "template-part"]
)
def test_heading_that_names_no_section_ends_the_section_before_it(tmp_path: Path, suffix: str, heading: str):
    # The text under "Glossary" reads like risk text, so it would pass ANNEX_I.I if it joined the risks.
    write, load = WRITERS[suffix]
    blocks = [
        ("large", "Summary"),
        ("body", SUMMARY),
        ("large", "Risk Factors"),
        ("body", THIN_RISKS),
        ("large", heading),
        ("body", RISK_TEXT),
    ]
    path = write(tmp_path / f"quillmoor{suffix}", blocks)

    sections = load(path)
    assert sorted(sections) == ["risks", "summary"]
    assert _words(sections["risks"]) == THIN_RISKS
    findings = _findings(path)
    assert findings["ANNEX_I.I"].rule.severity is Severity.BLOCKER
    assert (findings["ANNEX_I.I"].status, findings["ANNEX_I.I"].word_count) == ("review", 11)


@pytest.mark.parametrize("suffix", sorted(WRITERS))
def test_sub_heading_that_names_no_section_stays_in_its_section(tmp_path: Path, suffix: str):
    # A smaller heading font (PDF) or a Heading 2 (DOCX) under "Risk Factors" is part of the risks.
    write, load = WRITERS[suffix]
    blocks = [
        ("large", "Summary"),
        ("body", SUMMARY),
        ("large", "Risk Factors"),
        ("body", THIN_RISKS),
        ("subheading", "Market Volatility"),
        ("body", RISK_TEXT),
        ("large", "Glossary"),
        ("body", TECHNOLOGY),
    ]
    path = write(tmp_path / f"quillmoor{suffix}", blocks)

    sections = load(path)
    assert sorted(sections) == ["risks", "summary"]
    assert _words(RISK_TEXT) in _words(sections["risks"])
    assert TECHNOLOGY not in _words(sections["risks"])
    assert _findings(path)["ANNEX_I.I"].status == "pass"


# A risk section split into one paragraph per risk, under labels that name no section. Together the
# paragraphs pass ANNEX_I.I; the first alone does not.
RISK_PARTS = [
    (
        "Market risk",
        "Holders may sell units only where a market exists, and prices can move sharply within a single day. "
        "Past performance of similar services does not indicate future returns, and holders may lose the full "
        "amount. Holders should read this document in full and seek independent advice before they buy units.",
    ),
    (
        "Liquidity risk",
        "Liquidity can disappear when few buyers are active, so holders may be unable to exit at a fair price.",
    ),
    (
        "Regulatory risk",
        "A change in regulatory treatment in any Member State could restrict transfers or require new licences.",
    ),
    (
        "Operational risk",
        "Operational failures at the service provider, such as outages or faulty software updates, can delay "
        "transfers for days. Keys held by holders can be lost or stolen, and lost keys cannot be recovered.",
    ),
    (
        "Smart contract risk",
        "Smart contract defects can freeze balances until a fix is deployed and accepted by the network. "
        "Technology failures at the ledger can also stop transfers for days.",
    ),
]
GERMAN_RISK_LABELS = [
    "Marktrisiko",
    "Liquiditätsrisiko",
    "Regulatorisches Risiko",
    "Operationelles Risiko",
    "Smart-Contract-Risiko",
]
# The risk headings: the heading that opens the section and the labels of its paragraphs.
RISK_HEADINGS = {
    "numbered": (
        "12. Information on risks",
        [f"12.{n} {label}" for n, (label, _) in enumerate(RISK_PARTS, 1)],
    ),
    "unnumbered": ("Information on risks", [label for label, _ in RISK_PARTS]),
    "german": (
        "12. Informationen über die Risiken",
        [f"12.{n} {label}" for n, label in enumerate(GERMAN_RISK_LABELS, 1)],
    ),
}


def _risk_blocks(kind: str, opener: str, labels: list[str]) -> list[tuple[str, str]]:
    """A summary, the risk heading and the risk paragraphs under their labels, every heading in kind,
    then a "Glossary" in the larger heading font."""
    blocks = [(kind, "Summary"), ("body", SUMMARY), (kind, opener)]
    for label, (_, text) in zip(labels, RISK_PARTS, strict=True):
        blocks += [(kind, label), ("body", text)]
    return [*blocks, ("large", "Glossary"), ("body", TECHNOLOGY)]


def _assert_risk_parts_pass(path: Path, sections: dict[str, str]) -> None:
    """Every risk paragraph is in the risk section, the glossary text is not, and ANNEX_I.I passes."""
    assert sorted(sections) == ["risks", "summary"]
    risks = _words(sections["risks"])
    assert all(_words(text) in risks for _, text in RISK_PARTS)
    assert TECHNOLOGY not in risks
    findings = _findings(path)
    assert findings["ANNEX_I.I"].rule.severity is Severity.BLOCKER
    assert findings["ANNEX_I.I"].status == "pass"


@pytest.mark.parametrize("suffix", sorted(WRITERS))
@pytest.mark.parametrize("case", sorted(RISK_HEADINGS))
def test_bold_label_under_a_bold_heading_stays_in_its_section(tmp_path: Path, suffix: str, case: str):
    # "12.1 Market risk" or "Market risk" in bold under a bold "12. Information on risks": a bold DOCX
    # label and a PDF line in bold at the body size are sub-headings, not the next section.
    write, load = WRITERS[suffix]
    opener, labels = RISK_HEADINGS[case]
    path = write(tmp_path / f"quillmoor{suffix}", _risk_blocks("bold", opener, labels))

    _assert_risk_parts_pass(path, load(path))


@pytest.mark.parametrize("suffix", sorted(WRITERS))
@pytest.mark.parametrize("case", sorted(RISK_HEADINGS))
def test_heading_numbered_below_or_naming_one_risk_stays_in_the_risk_section(
    tmp_path: Path, suffix: str, case: str
):
    # The labels have the Heading 1 style or the 14 pt heading font of the risk heading, the same rank.
    # "12.1" extends the number of "12.", and "Market risk" names one risk, so each stays in the risks.
    write, load = WRITERS[suffix]
    opener, labels = RISK_HEADINGS[case]
    path = write(tmp_path / f"quillmoor{suffix}", _risk_blocks("large", opener, labels))

    _assert_risk_parts_pass(path, load(path))


@pytest.mark.parametrize("suffix", sorted(WRITERS))
def test_cover_title_that_names_a_section_sets_no_heading_level(tmp_path: Path, suffix: str):
    # A large cover title that names the crypto-asset section opens it in a PDF (a DOCX skips the Title
    # style). The headings after it keep their own level, so "Glossary" in the font of "Risk Factors"
    # still ends the risks, in the PDF as in the DOCX.
    write, load = WRITERS[suffix]
    blocks = [
        ("title", "Crypto-Asset White Paper for the Quillmoor Utility Token"),
        ("space", ""),
        ("large", "Summary"),
        ("body", SUMMARY),
        ("large", "Risk Factors"),
        ("body", THIN_RISKS),
        ("large", "Glossary"),
        ("body", RISK_TEXT),
    ]
    path = write(tmp_path / f"quillmoor{suffix}", blocks)

    sections = load(path)
    assert _words(sections["risks"]) == THIN_RISKS
    findings = _findings(path)
    assert findings["ANNEX_I.I"].rule.severity is Severity.BLOCKER
    assert (findings["ANNEX_I.I"].status, findings["ANNEX_I.I"].word_count) == ("review", 11)


# Rows of the summary table of Implementing Regulation (EU) 2024/2984 as pypdf reads them: the field
# number, the label and a one-line answer on one line, or a wrapped label followed by its answer.
SUMMARY_TABLE = [
    "01 Date of notification 2 February 2026",
    "02 Statement in accordance",
    "with Article 6(3) of",
    "Regulation (EU) 2023/1114",
    "This crypto-asset white paper has not been approved by any competent authority in any Member State.",
    "04 Statement in accordance",
    "with Article 6(5) of",
    "Regulation (EU) 2023/1114",
    "The crypto-asset may lose its value in part or in full and may not always be transferable.",
    "08 Characteristics of the crypto-",
    "asset",
    "The QMT units are utility tokens that give access to the shared ledger service.",
]


def test_pdf_summary_table_field_that_names_no_section_ends_the_field_before_it(tmp_path: Path):
    blocks = [
        ("large", "Summary"),
        *(("line", line) for line in SUMMARY_TABLE),
        ("large", f"Part A {DASH} Information about the offeror or the person seeking admission to trading"),
        ("line", "A.1 Name Quillmoor Ledger S.A."),
        ("line", "A.2 Legal form Societe anonyme"),
        ("line", "A.3 Registered address 12 Rue de la Gare, L-1611 Luxembourg"),
        ("large", f"Part I {DASH} Information on risks"),
        ("body", RISK_TEXT),
    ]
    path = _write_pdf(tmp_path / "quillmoor.pdf", blocks)

    sections = load_from_pdf(path)
    # The answer on the field line stays usable; the statements of fields 02 and 04 are no answer to it.
    assert sections["notification_date"] == "2 February 2026"
    assert sections["crypto_asset"] == SUMMARY_TABLE[-1]
    # The numbered fields of Part A belong to Part A.
    assert "A.3 Registered address 12 Rue de la Gare" in sections["offeror"]
    findings = _findings(path)
    assert findings["COMMON.NOTIFICATION_DATE"].rule.severity is Severity.MAJOR
    # The three words are a calendar date, which is all the field asks for.
    assert (findings["COMMON.NOTIFICATION_DATE"].status, findings["COMMON.NOTIFICATION_DATE"].word_count) == (
        "pass",
        3,
    )
    assert findings["ANNEX_I.A"].rule.severity is Severity.BLOCKER
    assert findings["ANNEX_I.A"].status != "missing"
    assert findings["ANNEX_I.I"].status == "pass"


# German drafts head the date of notification in the words of the German texts: Art. 6 Abs. 8, Art. 19
# Abs. 7 and Art. 51 Abs. 7 MiCAR ask for "das Datum seiner Übermittlung", and the Implementing Regulation
# (EU) 2024/2984 template labels field 01 (I.01 for ARTs and EMTs) "Datum der Mitteilung". "Datum der
# Notifizierung" spells out "Notifizierungsdatum", and "Übermittlungsdatum" and "Mitteilungsdatum" are the
# terms of the two texts as compounds.
GERMAN_NOTIFICATION_HEADINGS = [
    "Datum der Übermittlung",
    "Datum der Mitteilung",
    "Datum der Notifizierung",
    "Übermittlungsdatum",
    "Mitteilungsdatum",
]
# Labels of other fields of the German template that name an "Übermittlung" or a "Datum": the type of
# submission (F.5, B.8 and B.5), the methods of transferring value on reimbursement (E.25) and the
# publication date (F.10, B.12 and B.9).
GERMAN_OTHER_LABELS = [
    "Art der Übermittlung",
    "F.5 Die Art der Übermittlung",
    "E.25 Methoden der Übermittlung des Werts bei einer Rückerstattung",
    "F.10 Datum der Veröffentlichung",
]
GERMAN_SUMMARY = (
    "Der Quillmoor Utility Token (QMT) gibt seinen Inhabern Zugang zu einem gemeinsamen Ledger für "
    "Logistikunternehmen."
)
GERMAN_NOTIFICATION = (
    "Das Whitepaper wurde der BaFin am 2. März 2026 übermittelt, und dieses Datum gilt als Tag der "
    "Mitteilung."
)
GERMAN_TECHNOLOGY = "Übertragungen werden auf einem Ledger mit fünf validierenden Knoten abgewickelt."


@pytest.mark.parametrize(
    "heading",
    [
        *GERMAN_NOTIFICATION_HEADINGS,
        "DATUM DER ÜBERMITTLUNG",
        "Datum der Übermittlung des Kryptowerte-Whitepapers",
        "01 Datum der Mitteilung",
        "I.01 Datum der Mitteilung",
    ],
)
def test_german_notification_date_heading_names_the_notification_date(heading: str):
    assert match_heading_to_section(heading) == ["notification_date"]


# Marketing communications ("Marketingmitteilungen", Art. 7 MiCAR) are no notification, and the word
# boundary keeps "Mitteilungsdatum" from matching inside a longer word.
@pytest.mark.parametrize(
    "label", [*GERMAN_OTHER_LABELS, "Datum der Marketingmitteilung", "Marketingmitteilungsdatum"]
)
def test_other_german_template_label_does_not_name_the_notification_date(label: str):
    assert "notification_date" not in match_heading_to_section(label)


def _write_markdown(path: Path, blocks: list[tuple[str, str]]) -> Path:
    """A Markdown draft: each heading block a level 2 heading, each other block a paragraph."""
    markdown = "\n\n".join(f"## {text}" if kind == "heading" else text for kind, text in blocks)
    path.write_text(f"{markdown}\n", encoding="utf-8")
    return path


def _write_xhtml(path: Path, blocks: list[tuple[str, str]]) -> Path:
    """An XHTML draft with one element per line from line 4: each heading block an h2 whose id is
    "heading-" and its number, counted from 1, each other block a paragraph."""
    elements = []
    headings = 0
    for kind, text in blocks:
        if kind == "heading":
            headings += 1
            elements.append(f'<h2 id="heading-{headings}">{escape(text)}</h2>')
        else:
            elements.append(f"<p>{escape(text)}</p>")
    body = "\n".join(elements)
    path.write_text(
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        f'<html xmlns="http://www.w3.org/1999/xhtml" lang="de">\n<body>\n{body}\n</body>\n</html>\n',
        encoding="utf-8",
    )
    return path


# WRITERS, and Markdown and heading-based XHTML, which share the heading patterns.
ALL_WRITERS = {
    **WRITERS,
    ".md": (_write_markdown, load_from_markdown),
    ".xhtml": (_write_xhtml, load_from_xhtml),
}


def _german_draft(heading: str, text: str, summary: str = GERMAN_SUMMARY) -> list[tuple[str, str]]:
    """A German draft: a summary, text under heading, and a technology section that ends it."""
    return [
        ("heading", "Zusammenfassung"),
        ("body", summary),
        ("heading", heading),
        ("body", text),
        ("heading", "Technologie"),
        ("body", GERMAN_TECHNOLOGY),
    ]


@pytest.mark.parametrize("suffix", sorted(ALL_WRITERS))
@pytest.mark.parametrize("heading", GERMAN_NOTIFICATION_HEADINGS)
def test_german_notification_date_heading_opens_its_section(tmp_path: Path, suffix: str, heading: str):
    write, load = ALL_WRITERS[suffix]
    path = write(tmp_path / f"quillmoor{suffix}", _german_draft(heading, GERMAN_NOTIFICATION))

    sections = load(path)
    assert sorted(sections) == ["notification_date", "summary", "technology"]
    assert _words(sections["notification_date"]) == GERMAN_NOTIFICATION
    assert _words(sections["summary"]) == GERMAN_SUMMARY
    finding = _findings(path)["COMMON.NOTIFICATION_DATE"]
    assert finding.rule.severity is Severity.MAJOR
    assert (finding.status, finding.word_count) == ("pass", len(GERMAN_NOTIFICATION.split()))


@pytest.mark.parametrize("suffix", sorted(ALL_WRITERS))
def test_other_german_template_labels_open_no_notification_date_section(tmp_path: Path, suffix: str):
    # The type of submission ("Neu" for a new white paper), the methods of transferring value and the
    # publication date are not the date of notification, so the check still reports it missing.
    write, load = ALL_WRITERS[suffix]
    blocks = [
        ("heading", "Zusammenfassung"),
        ("body", GERMAN_SUMMARY),
        ("heading", "Art der Übermittlung"),
        ("body", "Neu"),
        ("heading", "Methoden der Übermittlung des Werts bei einer Rückerstattung"),
        ("body", "Die Rückerstattung erfolgt durch Überweisung auf das Konto des Käufers."),
        ("heading", "Datum der Veröffentlichung"),
        ("body", "Das Whitepaper wird am 16. März 2026 veröffentlicht."),
        ("heading", "Technologie"),
        ("body", GERMAN_TECHNOLOGY),
    ]
    path = write(tmp_path / f"quillmoor{suffix}", blocks)

    assert "notification_date" not in load(path)
    finding = _findings(path)["COMMON.NOTIFICATION_DATE"]
    assert (finding.status, finding.word_count) == ("missing", 0)


def test_pdf_german_summary_table_row_gives_its_answer_to_the_notification_date(tmp_path: Path):
    # Field 01 of the German summary table as pypdf reads it, with the answer on the label line. Field 02
    # names no section, so it ends field 01 and its statement is no answer to it.
    blocks = [
        ("large", "Zusammenfassung"),
        ("line", "01 Datum der Mitteilung 2026-03-02"),
        ("line", "02 Erklärung gemäß Artikel 6 Absatz 3 der Verordnung (EU) 2023/1114"),
        (
            "line",
            "Dieses Kryptowerte-Whitepaper wurde von keiner zuständigen Behörde eines Mitgliedstaats der "
            "Europäischen Union genehmigt.",
        ),
    ]
    path = _write_pdf(tmp_path / "quillmoor.pdf", blocks)

    assert load_from_pdf(path) == {"notification_date": "2026-03-02"}
    finding = _findings(path)["COMMON.NOTIFICATION_DATE"]
    assert finding.rule.severity is Severity.MAJOR
    # The answer is a calendar date, which is all the field asks for.
    assert (finding.status, finding.word_count) == ("pass", 1)


def test_xhtml_german_notification_date_heading_anchors_its_findings(tmp_path: Path):
    # Finding anchors and SARIF lines read the heading as the loader does: the second heading, on line 6.
    path = _write_xhtml(
        tmp_path / "quillmoor.xhtml", _german_draft("Datum der Mitteilung", GERMAN_NOTIFICATION)
    )

    assert _xhtml_section_identifiers(path)["notification_date"] == "xhtml:heading-2"
    assert _xhtml_heading_lines(path)["notification_date"] == 6


# German drafts call the redemption "Rücktausch", as the German text of MiCAR does: "Recht auf Rücktausch"
# (Art. 39), "Rücktauschplan" (Art. 47) and "dauerhaften Rücktauschrechte" (Anhang II Teil D Nr. 9).
# Implementing Regulation (EU) 2024/2984 labels field I.06 of the ART summary "Recht auf Rücktausch" and
# field D.31 "Form des Rücktauschs".
GERMAN_REDEMPTION_HEADINGS = ["Rücktausch", "Recht auf Rücktausch", "Dauerhafte Rücktauschrechte"]
GERMAN_ART_SUMMARY = (
    "Der QAR-Token der Quillmoor Ledger S.A. ist ein vermögenswertereferenzierter Token, der sich auf einen "
    "Korb aus Euro-Einlagen und Anleihen bezieht."
)
# Art. 39 Abs. 1 and 2 MiCAR: redemption at all times, here by paying the market value of the referenced
# assets. At 31 words the section is too thin for ANNEX_II.I (100 words), so the check reviews it whatever
# its terms.
GERMAN_REDEMPTION = (
    "Die Inhaber des QAR-Token können jederzeit den Rücktausch bei der Quillmoor Ledger S.A. verlangen. Der "
    "Emittent zahlt einen Geldbetrag in Höhe des Marktwerts der Vermögenswerte, auf die der Token Bezug nimmt."
)


@pytest.mark.parametrize(
    "heading",
    [
        *GERMAN_REDEMPTION_HEADINGS,
        "RÜCKTAUSCH",
        "Rücktauschrechte",
        "Rücktausch des Token",
        "Rücktauschplan",
        "I.06 Recht auf Rücktausch",
        "D.31 Form des Rücktauschs",
        # A DOCX heading with a line break in it.
        "Recht auf\nRücktausch",
    ],
)
def test_german_redemption_heading_names_the_redemption(heading: str):
    assert match_heading_to_section(heading) == ["redemption"]


# A heading that also names a risk heads a risk. It opens the sections it opened before, so the redemption
# section takes no risk text from the risks.
@pytest.mark.parametrize(
    "heading,keys",
    [
        ("Rücktauschrisiko", []),
        ("Risiko der Aussetzung des Rücktauschs", []),
        ("Liquiditätsrisiken der dauerhaften Rücktauschrechte", []),
        ("Risiken im Zusammenhang mit dem Rücktausch", ["risks"]),
    ],
)
def test_german_risk_heading_that_names_the_redemption_does_not_name_it(heading: str, keys: list[str]):
    assert match_heading_to_section(heading) == keys


@pytest.mark.parametrize("suffix", sorted(ALL_WRITERS))
@pytest.mark.parametrize("heading", GERMAN_REDEMPTION_HEADINGS)
def test_german_redemption_heading_opens_its_section(tmp_path: Path, suffix: str, heading: str):
    write, load = ALL_WRITERS[suffix]
    path = write(
        tmp_path / f"quillmoor{suffix}", _german_draft(heading, GERMAN_REDEMPTION, GERMAN_ART_SUMMARY)
    )

    sections = load(path)
    assert sorted(sections) == ["redemption", "summary", "technology"]
    assert _words(sections["redemption"]) == GERMAN_REDEMPTION
    finding = _findings(path, WhitepaperType.ART)["ANNEX_II.I"]
    assert finding.rule.severity is Severity.BLOCKER
    assert (finding.status, finding.word_count) == ("review", len(GERMAN_REDEMPTION.split()))


@pytest.mark.parametrize("suffix", sorted(ALL_WRITERS))
def test_german_rights_heading_that_names_the_redemption_opens_both_sections(tmp_path: Path, suffix: str):
    # A heading names every section whose pattern it matches, so the rights section keeps its heading and
    # the redemption check reads the same chapter.
    heading = "Rechte und Pflichten, Rücktausch"
    rights = (
        "Die Inhaber können ihre Ansprüche aus den Bedingungen des QAR-Token vor den Gerichten in Luxemburg "
        f"geltend machen. {GERMAN_REDEMPTION}"
    )
    write, load = ALL_WRITERS[suffix]
    path = write(tmp_path / f"quillmoor{suffix}", _german_draft(heading, rights, GERMAN_ART_SUMMARY))

    assert match_heading_to_section(heading) == ["rights_and_obligations", "redemption"]
    sections = load(path)
    assert sorted(sections) == ["redemption", "rights_and_obligations", "summary", "technology"]
    assert _words(sections["rights_and_obligations"]) == _words(sections["redemption"]) == rights
    findings = _findings(path, WhitepaperType.ART)
    for rule_id in ("ANNEX_II.D", "ANNEX_II.I"):
        assert findings[rule_id].rule.severity is Severity.BLOCKER
        assert findings[rule_id].word_count == len(rights.split())


# A risk section whose paragraphs sit under headings that each name one risk.
GERMAN_RISK_PARTS = [
    (
        "Marktrisiko",
        "Der Marktpreis des QAR-Token kann vom Wert der Referenzwerte abweichen, und Inhaber können beim "
        "Verkauf Verluste erleiden.",
    ),
    (None, "Verzögert sich der Rücktausch, erhalten die Inhaber ihr Geld später als erwartet."),
    (
        "Liquiditätsrisiko",
        "Bei wenigen Käufern kann die Liquidität des QAR-Token ausbleiben, sodass Inhaber nicht zu einem "
        "fairen Preis verkaufen können.",
    ),
]


@pytest.mark.parametrize("suffix", sorted(WRITERS))
@pytest.mark.parametrize(
    "heading",
    [
        "Rücktauschrisiko",
        "Risiko der Aussetzung des Rücktauschs",
        "Risiken im Zusammenhang mit dem Rücktausch",
    ],
)
def test_german_risk_sub_heading_that_names_the_redemption_stays_in_the_risks(
    tmp_path: Path, suffix: str, heading: str
):
    # The second risk heading names the redemption. Its paragraph and the one after it stay in the risks
    # (a PDF keeps a bold label that names no section as text of its section), and the redemption section
    # holds only the chapter under "Rücktausch".
    write, load = WRITERS[suffix]
    blocks = [("large", "Zusammenfassung"), ("body", GERMAN_ART_SUMMARY), ("large", "Risikofaktoren")]
    for label, text in GERMAN_RISK_PARTS:
        blocks += [("subheading", label or heading), ("body", text)]
    blocks += [("large", "Rücktausch"), ("body", GERMAN_REDEMPTION)]
    path = write(tmp_path / f"quillmoor{suffix}", blocks)

    sections = load(path)
    assert sorted(sections) == ["redemption", "risks", "summary"]
    risks = _words(sections["risks"])
    assert all(_words(text) in risks for _, text in GERMAN_RISK_PARTS)
    assert _words(sections["redemption"]) == GERMAN_REDEMPTION
    findings = _findings(path, WhitepaperType.ART)
    assert findings["ANNEX_II.I"].rule.severity is Severity.BLOCKER
    assert (findings["ANNEX_II.I"].status, findings["ANNEX_II.I"].word_count) == (
        "review",
        len(GERMAN_REDEMPTION.split()),
    )


def test_xhtml_german_redemption_heading_anchors_its_findings(tmp_path: Path):
    # Finding anchors and SARIF lines read the headings as the loader does: the risk heading before the
    # chapter names the risks only, so the redemption findings point at "Rücktausch", the third heading,
    # on line 8.
    blocks = [
        ("heading", "Zusammenfassung"),
        ("body", GERMAN_ART_SUMMARY),
        ("heading", "Risiken im Zusammenhang mit dem Rücktausch"),
        ("body", GERMAN_RISK_PARTS[1][1]),
        ("heading", "Rücktausch"),
        ("body", GERMAN_REDEMPTION),
    ]
    path = _write_xhtml(tmp_path / "quillmoor.xhtml", blocks)

    identifiers = _xhtml_section_identifiers(path)
    assert (identifiers["risks"], identifiers["redemption"]) == ("xhtml:heading-2", "xhtml:heading-3")
    assert _xhtml_heading_lines(path)["redemption"] == 8


@pytest.mark.parametrize("answer", ["N/A", "n/a", "entfällt", "nicht anwendbar", "k. A."])
def test_pdf_numbered_heading_after_a_not_applicable_answer_opens_its_section(tmp_path: Path, answer: str):
    # An answer that says the item does not apply, in English or German, continues no sentence, so the
    # heading above it stays a heading even in lower case, and the numbered heading after it opens the
    # technology section instead of only ending the risks.
    blocks = [
        ("heading", "Zusammenfassung"),
        ("body", GERMAN_SUMMARY),
        ("heading", "5. Risikofaktoren"),
        ("line", answer),
        ("heading", "6. Technologie"),
        ("body", GERMAN_TECHNOLOGY),
    ]
    sections = load_from_pdf(_write_pdf(tmp_path / "quillmoor.pdf", blocks))

    assert sorted(sections) == ["risks", "summary", "technology"]
    assert sections["risks"] == answer
    assert _words(sections["technology"]) == GERMAN_TECHNOLOGY


# The answer to field 03 of the 2024/2984 summary, in the wording of the template.
COMPLIANCE = (
    "This crypto-asset white paper complies with Title II of Regulation (EU) 2023/1114 and, to the best of "
    "the knowledge of the management body, the information presented in the crypto-asset white paper is "
    "fair, clear and not misleading and the crypto-asset white paper makes no omission likely to affect "
    "its import."
)


@pytest.mark.parametrize("suffix,kind", [(".docx", "bold"), (".pdf", "line")], ids=["docx-bold", "pdf-line"])
def test_compliance_statement_field_of_the_summary_opens_the_management_statement(
    tmp_path: Path, suffix: str, kind: str
):
    # Field 03 "Compliance statement in accordance with Article 6(6)" holds the statement of the
    # management body, so its answer is the management statement and passes its check.
    write, load = WRITERS[suffix]
    blocks = [
        ("large", "Summary"),
        ("body", KEY_SUMMARY),
        (kind, "03 Compliance statement in accordance with Article 6(6) of Regulation (EU) 2023/1114"),
        ("body", COMPLIANCE),
        (kind, "04 Statement in accordance with Article 6(5) of Regulation (EU) 2023/1114"),
        (
            "body",
            "The crypto-asset may lose its value in part or in full and may not always be transferable.",
        ),
        ("large", f"Part I {DASH} Information on risks"),
        ("body", RISK_TEXT),
    ]
    path = write(tmp_path / f"quillmoor{suffix}", blocks)

    sections = load(path)
    assert sorted(sections) == ["management_statement", "risks", "summary"]
    # The PDF writer may wrap "crypto-asset" at its hyphen, so compare the parts of the statement.
    statement = _words(sections["management_statement"])
    assert "complies with Title II of Regulation (EU) 2023/1114" in statement
    assert "fair, clear and not misleading" in statement
    assert _words(sections["summary"]) == KEY_SUMMARY
    findings = _findings(path)
    assert findings["COMMON.MANAGEMENT_STATEMENT"].rule.severity is Severity.BLOCKER
    assert findings["COMMON.MANAGEMENT_STATEMENT"].status == "pass"


# A technology section that passes ANNEX_I.H only as a whole: the first half names the protocol and the
# consensus, the second half the smart contracts and the audit.
TECHNOLOGY_FIRST = (
    "Transfers settle on a permissioned ledger based on a Byzantine fault tolerant consensus protocol run "
    "by five validating nodes. Each node is operated by a different partner firm in a different Member State."
)
TECHNOLOGY_SECOND = (
    "Smart contracts implement the unit register and the fee settlement. The source code is audited once a "
    "year by an independent firm, and the audit reports are published on the website."
)
RUNNING_HEADER = f"Quillmoor Utility Token (QMT) {DASH} Crypto-Asset White Paper"


@pytest.mark.parametrize("kind", ["large", "heading"], ids=["14pt-bold", "body-font"])
def test_pdf_running_page_header_that_names_a_section_is_body_text(tmp_path: Path, kind: str):
    # The header names the crypto-asset section on every page; it must not open that section and take
    # the technology text that runs on to the next page.
    blocks = [
        ("line", RUNNING_HEADER),
        (kind, "Summary"),
        ("body", SUMMARY),
        (kind, "Information on the Underlying Technology"),
        ("body", TECHNOLOGY_FIRST),
        ("page", ""),
        ("line", RUNNING_HEADER),
        ("body", TECHNOLOGY_SECOND),
        (kind, "Risk Factors"),
        ("body", RISK_TEXT),
    ]
    path = _write_pdf(tmp_path / "quillmoor.pdf", blocks)

    sections = load_from_pdf(path)
    assert sorted(sections) == ["risks", "summary", "technology"]
    assert _words(TECHNOLOGY_SECOND) in _words(sections["technology"])
    findings = _findings(path)
    assert findings["ANNEX_I.F"].rule.severity is Severity.BLOCKER
    assert findings["ANNEX_I.F"].status == "missing"
    assert findings["ANNEX_I.H"].rule.severity is Severity.BLOCKER
    assert findings["ANNEX_I.H"].status == "pass"


def test_pdf_body_font_line_is_no_heading_where_headings_have_a_heading_font(tmp_path: Path):
    # The draft sets its headings in 14 pt bold, so the title-case line "Technology Roadmap" in the body
    # font is no heading: the roadmap must not pass the thin technology section.
    blocks = [
        ("large", "Summary"),
        ("body", SUMMARY),
        ("large", "Information about the Project"),
        ("body", PROJECT),
        ("line", "Technology Roadmap"),
        ("body", f"{TECHNOLOGY_FIRST} {TECHNOLOGY_SECOND}"),
        ("large", "Information on the Underlying Technology"),
        ("body", THIN_TECHNOLOGY),
    ]
    path = _write_pdf(tmp_path / "quillmoor.pdf", blocks)

    sections = load_from_pdf(path)
    assert sorted(sections) == ["project", "summary", "technology"]
    assert _words(sections["project"]) == PROJECT
    assert _words(sections["technology"]) == _words(THIN_TECHNOLOGY)
    findings = _findings(path)
    assert findings["ANNEX_I.H"].rule.severity is Severity.BLOCKER
    assert (findings["ANNEX_I.H"].status, findings["ANNEX_I.H"].word_count) == (
        "review",
        len(THIN_TECHNOLOGY.split()),
    )


def test_docx_style_based_on_itself_through_a_cycle_loads(tmp_path: Path):
    doc = docx.Document()
    loop_a = doc.styles.add_style("Loop A", WD_STYLE_TYPE.PARAGRAPH)
    loop_b = doc.styles.add_style("Loop B", WD_STYLE_TYPE.PARAGRAPH)
    loop_a.base_style = loop_b
    loop_b.base_style = loop_a
    doc.add_heading("Summary", level=1)
    doc.add_paragraph(SUMMARY, style="Loop A")
    path = tmp_path / "quillmoor.docx"
    doc.save(path)

    # A daemon thread, so a loader that never returns fails the test instead of hanging the run.
    loaded: dict[str, str] = {}
    worker = threading.Thread(target=lambda: loaded.update(load_from_docx(path)), daemon=True)
    worker.start()
    worker.join(timeout=10)
    assert not worker.is_alive(), "load_from_docx did not return for a cyclic style chain"
    assert loaded == {"summary": SUMMARY}


def test_docx_styles_based_on_a_numbering_style_load(tmp_path: Path):
    # python-docx gives a numbering (list) style neither a base style nor a font.
    doc = docx.Document()
    numbering = doc.styles.add_style("Quillmoor List", WD_STYLE_TYPE.LIST)
    body = doc.styles.add_style("Quillmoor Body", WD_STYLE_TYPE.PARAGRAPH)
    body.base_style = numbering
    strong = doc.styles.add_style("Quillmoor Strong", WD_STYLE_TYPE.CHARACTER)
    strong.base_style = numbering
    doc.add_heading("Summary", level=1)
    doc.add_paragraph(SUMMARY, style="Quillmoor Body")
    doc.add_paragraph(style="Quillmoor Body").add_run("Risk Factors", style="Quillmoor Strong").bold = True
    doc.add_paragraph(RISK_TEXT, style="Quillmoor Body")
    path = tmp_path / "quillmoor.docx"
    doc.save(path)

    sections = load_from_docx(path)
    assert sorted(sections) == ["risks", "summary"]
    assert _words(sections["risks"]) == _words(RISK_TEXT)
    findings = _findings(path)
    assert findings["ANNEX_I.I"].rule.severity is Severity.BLOCKER
    assert findings["ANNEX_I.I"].status == "pass"
