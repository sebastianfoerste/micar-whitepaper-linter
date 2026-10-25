"""The first-page statement that no competent authority has approved the white paper.

The English fixture text is the statement as Art. 6(3) and Art. 51(3) MiCAR word it,
and as Table 2 field 02 and Table 4 field I.02 of the Annex to Implementing Regulation
(EU) 2024/2984 print it ("Predefined alphanumerical text"). Source read: the EUR-Lex
texts CELEX 32023R1114 and 32024R2984 as mirrored in legalize-dev/legalize-eu. Art. 19 MiCAR
prescribes no such statement for an ART white paper (Table 3 field I.02 of 2024/2984
holds the Art. 19(4) warning instead), so the ART ruleset carries no such rule. That is
not because every ART white paper is approved: that of an authorised issuer is deemed
approved (Art. 21(1)) and that of a credit institution is approved (Art. 17(1)(a)), but an
issuer exempt under Art. 16(2) only notifies it.

The German wording has not been checked against the German Official Journal text. German
drafts are therefore referred to review and never pass; the German sentence below is only
a routing fixture, not a statement of the statutory wording.
"""

import textwrap
from pathlib import Path
from unittest.mock import MagicMock, patch

import docx
import pytest
import test_heading_detection
from test_heading_detection import _write_docx, _write_pdf

from micar_linter.linter import Linter, lint_whitepaper
from micar_linter.rules import RULESETS
from micar_linter.rules.base import Severity
from micar_linter.whitepaper import Whitepaper, WhitepaperType, load_whitepaper

ROOT = Path(__file__).resolve().parents[1]
TITLE_II = "ANNEX_I.NOT_APPROVED_STATEMENT"
EMT = "ANNEX_III.NOT_APPROVED_STATEMENT"

APPROVAL = (
    "This crypto-asset white paper has not been approved by any competent authority "
    "in any Member State of the European Union."
)


def _statement(subject: str) -> str:
    return (
        f"{APPROVAL} {subject} of the crypto-asset is solely responsible for the content "
        "of this crypto-asset white paper."
    )


def _finding(wp_type: WhitepaperType, rule_id: str, sections: dict[str, str], **metadata):
    whitepaper = Whitepaper(title="Statement test", type=wp_type, sections=sections, metadata=metadata)
    for finding in Linter(RULESETS[wp_type]).lint(whitepaper):
        if finding.rule.rule_id == rule_id:
            return finding
    raise AssertionError(f"{rule_id} was not applied")


@pytest.mark.parametrize(
    "subject",
    ["The offeror", "The person seeking admission to trading", "The operator of the trading platform"],
)
def test_title_ii_statement_passes_for_each_person_art_6_3_names(subject: str):
    finding = _finding(WhitepaperType.OTHER, TITLE_II, {"not_approved_statement": _statement(subject)})
    assert finding.status == "pass", finding.issues
    assert finding.rule.citation == "Art. 6 Abs. 3 MiCAR"


def test_title_ii_statement_naming_another_person_is_referred_to_review():
    # Art. 6(3) MiCAR allows only the offeror, the person seeking admission to trading or
    # the operator of the trading platform; Table 2 field 02 of 2024/2984 prints all three.
    finding = _finding(WhitepaperType.OTHER, TITLE_II, {"not_approved_statement": _statement("The issuer")})
    assert finding.status == "review"
    assert any("operator of the trading platform" in issue for issue in finding.issues)


def test_emt_statement_passes_when_it_names_the_issuer():
    finding = _finding(WhitepaperType.EMT, EMT, {"not_approved_statement": _statement("The issuer")})
    assert finding.status == "pass", finding.issues
    assert finding.rule.citation == "Art. 51 Abs. 3 MiCAR"


def test_emt_statement_naming_the_offeror_is_referred_to_review():
    finding = _finding(WhitepaperType.EMT, EMT, {"not_approved_statement": _statement("The offeror")})
    assert finding.status == "review"
    assert any("the issuer of the crypto-asset is solely responsible" in issue for issue in finding.issues)


@pytest.mark.parametrize("wp_type, rule_id", [(WhitepaperType.OTHER, TITLE_II), (WhitepaperType.EMT, EMT)])
def test_absent_statement_fails_closed_as_a_blocker(wp_type: WhitepaperType, rule_id: str):
    finding = _finding(wp_type, rule_id, {"summary": "A summary without the statement."})
    assert finding.status == "missing"
    assert finding.rule.severity is Severity.BLOCKER


@pytest.mark.parametrize("wp_type, rule_id", [(WhitepaperType.OTHER, TITLE_II), (WhitepaperType.EMT, EMT)])
def test_paraphrased_statement_is_referred_to_review(wp_type: WhitepaperType, rule_id: str):
    paraphrase = (
        "No supervisor has reviewed this document. The issuer of the crypto-asset is solely "
        "responsible for the content of this crypto-asset white paper."
    )
    finding = _finding(wp_type, rule_id, {"not_approved_statement": paraphrase})
    assert finding.status == "review"
    assert any("has not been approved by any competent authority" in issue for issue in finding.issues)


def test_art_ruleset_has_no_not_approved_statement():
    art_rules = {rule.rule_id for rule in RULESETS[WhitepaperType.ART]}
    assert not {rule_id for rule_id in art_rules if rule_id.endswith("NOT_APPROVED_STATEMENT")}
    assert TITLE_II in {rule.rule_id for rule in RULESETS[WhitepaperType.OTHER]}
    assert EMT in {rule.rule_id for rule in RULESETS[WhitepaperType.EMT]}


GERMAN_ROUTING_FIXTURE = (
    "Dieses Kryptowerte-Whitepaper wurde von keiner zuständigen Behörde eines Mitgliedstaats "
    "der Europäischen Union genehmigt. Der Anbieter des Kryptowerts trägt die alleinige "
    "Verantwortung für den Inhalt dieses Kryptowerte-Whitepapers."
)


@pytest.mark.parametrize("wp_type, rule_id", [(WhitepaperType.OTHER, TITLE_II), (WhitepaperType.EMT, EMT)])
def test_german_statement_is_referred_to_review_until_the_wording_is_verified(
    wp_type: WhitepaperType, rule_id: str
):
    finding = _finding(wp_type, rule_id, {"not_approved_statement": GERMAN_ROUTING_FIXTURE}, language="de")
    assert finding.status == "review"
    assert finding.rule.severity is Severity.BLOCKER
    assert any("nicht gegen die deutsche Fassung im Amtsblatt geprüft" in issue for issue in finding.issues)


def test_absent_german_statement_still_fails_closed():
    finding = _finding(WhitepaperType.OTHER, TITLE_II, {"summary": "Zusammenfassung."}, language="de")
    assert finding.status == "missing"


@pytest.mark.parametrize("value", ["true", "TRUE", "false", "Yes"])
def test_a_yes_no_field_value_is_not_read_as_the_statement(value: str):
    finding = _finding(WhitepaperType.EMT, EMT, {"not_approved_statement": value})
    assert finding.status == "review"
    assert any("yes/no value" in issue for issue in finding.issues)


def test_rule_labels_say_that_page_placement_is_not_checked():
    for wp_type, rule_id in [(WhitepaperType.OTHER, TITLE_II), (WhitepaperType.EMT, EMT)]:
        rule = next(r for r in RULESETS[wp_type] if r.rule_id == rule_id)
        assert "first-page placement not checked" in rule.label


@pytest.mark.parametrize("filename", ["other-crypto-asset.json", "emt-token.json"])
def test_canonical_examples_carry_the_statement(filename: str):
    report = lint_whitepaper(load_whitepaper(ROOT / "examples" / filename))
    statement = next(f for f in report.findings if f.rule.rule_id.endswith("NOT_APPROVED_STATEMENT"))
    assert statement.status == "pass", statement.issues


def test_numbered_field_02_feeds_the_title_ii_check():
    whitepaper = load_whitepaper(ROOT / "examples" / "numbered-other-draft.xhtml", WhitepaperType.OTHER)
    assert whitepaper.section("not_approved_statement").startswith(
        "Statement in accordance with Article 6(3) of Regulation (EU) 2023/1114:"
    )
    report = lint_whitepaper(whitepaper)
    statement = next(f for f in report.findings if f.rule.rule_id == TITLE_II)
    assert statement.status == "pass", statement.issues


def test_numbered_field_i_02_feeds_the_emt_check(tmp_path: Path):
    rows = "".join(
        f"<tr><td>{number}</td><td>{label}</td><td>{content}</td></tr>"
        for number, label, content in [
            ("I.01", "Date of notification", "2026-03-02"),
            (
                "I.02",
                "Statement in accordance with Article 51(3) of Regulation (EU) 2023/1114",
                _statement("The issuer"),
            ),
            ("A.1", "Statutory name", "Example EMI GmbH"),
            ("B.1", "Name of the e-money token", "Example EUR"),
            ("D.1", "Holder's rights and obligations", "Redemption at any time and at par value."),
        ]
    )
    path = tmp_path / "emt.xhtml"
    path.write_text(f"<html><body><table>{rows}</table></body></html>", encoding="utf-8")
    whitepaper = load_whitepaper(path, WhitepaperType.EMT)
    assert APPROVAL in whitepaper.section("not_approved_statement")
    statement = next(f for f in lint_whitepaper(whitepaper).findings if f.rule.rule_id == EMT)
    assert statement.status == "pass", statement.issues


def test_numbered_art_field_i_02_stays_the_risk_warning(tmp_path: Path):
    rows = "".join(
        f"<tr><td>{number}</td><td>{label}</td><td>{content}</td></tr>"
        for number, label, content in [
            ("I.02", "Statement in accordance with Article 19(4)", "The token may lose its value."),
            ("A.1", "Statutory name", "Example Issuer AG"),
            ("B.1", "Name", "Example ART"),
            ("G.1", "Description of the reserve of assets", "Bank deposits."),
        ]
    )
    path = tmp_path / "art.xhtml"
    path.write_text(f"<html><body><table>{rows}</table></body></html>", encoding="utf-8")
    whitepaper = load_whitepaper(path, WhitepaperType.ART)
    assert "not_approved_statement" not in whitepaper.sections
    assert "may lose its value" in whitepaper.section("risk_warning")


def test_statement_heading_routes_prose_drafts(tmp_path: Path):
    path = tmp_path / "draft.md"
    path.write_text(
        "# Statement in accordance with Article 6(3) of Regulation (EU) 2023/1114\n"
        f"{_statement('The offeror')}\n\n# Summary\nKey information.\n",
        encoding="utf-8",
    )
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    assert APPROVAL in whitepaper.section("not_approved_statement")
    assert "not been approved" not in whitepaper.section("summary")


# Regression: the loaders keep line breaks, and Art. 6(3) gives the statement no heading.
WRAPPED_STATEMENT = """This crypto-asset white paper has not been approved by any competent
authority in any Member State of the European Union. The offeror of the crypto-asset
is solely responsible for the content of this crypto-asset white paper."""


def _statement_from_markdown(tmp_path: Path, content: str) -> str:
    path = tmp_path / "draft.md"
    path.write_text(content, encoding="utf-8")
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    statement = next(f for f in lint_whitepaper(whitepaper).findings if f.rule.rule_id == TITLE_II)
    return statement


def test_hard_wrapped_markdown_statement_passes(tmp_path: Path):
    content = (
        "# Statement in accordance with Article 6(3) of Regulation (EU) 2023/1114\n"
        f"{WRAPPED_STATEMENT}\n\n# Summary\nKey information.\n"
    )
    statement = _statement_from_markdown(tmp_path, content)
    assert statement.status == "pass", statement.issues


def test_heading_less_first_paragraph_is_routed_to_the_statement(tmp_path: Path):
    content = f"{WRAPPED_STATEMENT}\n\n# Summary\nKey information about the offer.\n"
    statement = _statement_from_markdown(tmp_path, content)
    assert statement.status == "pass", statement.issues


def test_statement_under_an_unrelated_heading_is_routed_by_content(tmp_path: Path):
    content = f"# Important notice\n{WRAPPED_STATEMENT}\n\n# Summary\nKey information.\n"
    statement = _statement_from_markdown(tmp_path, content)
    assert statement.status == "pass", statement.issues


@pytest.mark.parametrize(
    "heading",
    [
        "Statement in accordance with Art. 6(3) of Regulation (EU) 2023/1114",
        "Erklärung gemäß Artikel 6 Absatz 3 der Verordnung (EU) 2023/1114",
    ],
)
def test_abbreviated_and_german_headings_route_to_the_statement(tmp_path: Path, heading: str):
    path = tmp_path / "draft.md"
    path.write_text(f"# {heading}\nStatement text.\n", encoding="utf-8")
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    assert whitepaper.section("not_approved_statement") == "Statement text."


def test_german_first_paragraph_is_routed_and_referred_to_review(tmp_path: Path):
    path = tmp_path / "entwurf.md"
    path.write_text(
        f"{GERMAN_ROUTING_FIXTURE}\n\n# Zusammenfassung\nDie wesentlichen Informationen.\n", encoding="utf-8"
    )
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    statement = next(f for f in lint_whitepaper(whitepaper).findings if f.rule.rule_id == TITLE_II)
    assert statement.status == "review"
    assert statement.word_count > 0


@patch("pypdf.PdfReader")
def test_wrapped_pdf_statement_passes(mock_reader_cls, tmp_path: Path):
    # The continuation lines contain "crypto-asset" and "offeror", which the PDF heading
    # heuristic would otherwise read as new section headings.
    page = MagicMock()
    page.extract_text.return_value = (
        "Example Token White Paper\n"
        "This crypto-asset white paper has not been approved by any competent authority in any\n"
        "Member State of the European Union. The offeror of the crypto-asset is solely\n"
        "responsible for the content of this crypto-\n"
        "asset white paper.\n"
        "Summary\n"
        "Key information about the offer.\n"
    )
    mock_reader_cls.return_value = MagicMock(pages=[page])
    path = tmp_path / "draft.pdf"
    path.write_bytes(b"%PDF-1.4")
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    statement = next(f for f in lint_whitepaper(whitepaper).findings if f.rule.rule_id == TITLE_II)
    assert statement.status == "pass", statement.issues
    assert "not been approved" not in whitepaper.section("summary")


def test_docx_statement_paragraph_without_heading_passes(tmp_path: Path):
    path = tmp_path / "draft.docx"
    document = docx.Document()
    document.add_paragraph(_statement("The offeror"))
    document.add_heading("Summary", level=1)
    document.add_paragraph("Key information about the offer.")
    document.save(path)
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    statement = next(f for f in lint_whitepaper(whitepaper).findings if f.rule.rule_id == TITLE_II)
    assert statement.status == "pass", statement.issues


def test_xhtml_statement_paragraph_without_heading_is_routed(tmp_path: Path):
    path = tmp_path / "draft.xhtml"
    path.write_text(
        "<html><body><p>"
        + _statement("The offeror")
        + "</p><h1>Summary</h1><p>Key information.</p></body></html>",
        encoding="utf-8",
    )
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    assert APPROVAL in whitepaper.section("not_approved_statement")
    assert "not been approved" not in whitepaper.section("summary")


# Regression: an incomplete statement swallowed the sections that followed it, so a draft
# that has a summary and a risk warning got MISSING blockers for both.
SUMMARY_BODY = " ".join(["The summary gives key information about the offer and the token."] * 15)
RISK_BODY = "The crypto-asset may lose its value in part or in full."


def _sections_after_incomplete_statement(whitepaper):
    findings = {f.rule.rule_id: f for f in lint_whitepaper(whitepaper).findings}
    assert findings["COMMON.SUMMARY"].status == "pass", findings["COMMON.SUMMARY"].issues
    assert findings["COMMON.RISK_WARNING"].status != "missing"
    assert findings[TITLE_II].status == "review"
    assert "key information" not in whitepaper.section("not_approved_statement")


def test_incomplete_docx_statement_does_not_swallow_the_next_sections(tmp_path: Path):
    path = tmp_path / "draft.docx"
    document = docx.Document()
    document.add_paragraph(APPROVAL)
    document.add_heading("Summary", level=1)
    document.add_paragraph(SUMMARY_BODY)
    document.add_heading("Risk warning", level=1)
    document.add_paragraph(RISK_BODY)
    document.save(path)
    _sections_after_incomplete_statement(load_whitepaper(path, WhitepaperType.OTHER))


def test_docx_heading_ends_the_statement_even_before_a_responsibility_sentence(tmp_path: Path):
    path = tmp_path / "draft.docx"
    document = docx.Document()
    document.add_paragraph(APPROVAL)
    document.add_heading("Summary", level=1)
    document.add_paragraph("The issuer is responsible for the token. " + SUMMARY_BODY)
    document.save(path)
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    assert whitepaper.section("not_approved_statement") == APPROVAL
    assert whitepaper.section("summary").startswith("The issuer is responsible for the token.")


def test_incomplete_markdown_statement_does_not_swallow_the_next_sections(tmp_path: Path):
    path = tmp_path / "draft.md"
    path.write_text(
        f"{APPROVAL}\n\nSummary\n{SUMMARY_BODY}\n\n# Summary\n{SUMMARY_BODY}\n\n# Risk warning\n{RISK_BODY}\n",
        encoding="utf-8",
    )
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    assert whitepaper.section("not_approved_statement") == APPROVAL
    _sections_after_incomplete_statement(whitepaper)


def _pdf_whitepaper(mock_reader_cls, tmp_path: Path, text: str):
    # The mocked reader reports no fonts, so a PDF heading must be in title case or capitals
    # ("Risk Warning"), as for any PDF without font information.
    page = MagicMock()
    page.extract_text.return_value = text
    mock_reader_cls.return_value = MagicMock(pages=[page])
    path = tmp_path / "draft.pdf"
    path.write_bytes(b"%PDF-1.4")
    return load_whitepaper(path, WhitepaperType.OTHER)


@patch("pypdf.PdfReader")
def test_incomplete_pdf_statement_does_not_swallow_the_next_sections(mock_reader_cls, tmp_path: Path):
    whitepaper = _pdf_whitepaper(
        mock_reader_cls,
        tmp_path,
        "This crypto-asset white paper has not been approved by any competent authority in any\n"
        "Member State of the European Union.\n"
        f"Summary\n{SUMMARY_BODY}\nRisk Warning\n{RISK_BODY}\n",
    )
    _sections_after_incomplete_statement(whitepaper)


@patch("pypdf.PdfReader")
def test_pdf_statement_cut_off_mid_sentence_ends_at_the_heading(mock_reader_cls, tmp_path: Path):
    whitepaper = _pdf_whitepaper(
        mock_reader_cls,
        tmp_path,
        "This crypto-asset white paper has not been approved\n"
        f"Summary\n{SUMMARY_BODY}\nRisk Warning\n{RISK_BODY}\n",
    )
    assert (
        whitepaper.section("not_approved_statement") == "This crypto-asset white paper has not been approved"
    )
    _sections_after_incomplete_statement(whitepaper)


@patch("pypdf.PdfReader")
def test_statement_wrapped_in_a_narrow_pdf_column_passes(mock_reader_cls, tmp_path: Path):
    # "State of the European Union. The offeror" and "of the crypto-asset is solely
    # responsible" are short and name sections, so a heading check alone would cut the
    # statement off.
    whitepaper = _pdf_whitepaper(
        mock_reader_cls,
        tmp_path,
        "This crypto-asset white paper has not been approved\n"
        "by any competent authority in any Member\n"
        "State of the European Union. The offeror\n"
        "of the crypto-asset is solely responsible\n"
        "for the content of this crypto-asset\n"
        "white paper.\n"
        f"Summary\n{SUMMARY_BODY}\n",
    )
    statement = next(f for f in lint_whitepaper(whitepaper).findings if f.rule.rule_id == TITLE_II)
    assert statement.status == "pass", statement.issues
    assert whitepaper.section("summary").startswith("The summary gives key information")


# The statement next to the headings that the DOCX and PDF loaders read from Heading styles, bold or
# capitals labels, heading fonts and template numbering: an incomplete statement stops at such a heading
# even where the next sentence names who is responsible.
@pytest.mark.parametrize("label", [("bold", "Summary"), ("plain", "SUMMARY")])
def test_docx_label_ends_the_statement_even_before_a_responsibility_sentence(tmp_path: Path, label):
    path = _write_docx(
        tmp_path / "draft.docx",
        [
            ("plain", APPROVAL),
            label,
            ("plain", "The issuer is responsible for the token. " + SUMMARY_BODY),
            ("bold", "Risk Warning"),
            ("plain", RISK_BODY),
        ],
    )
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    assert whitepaper.section("not_approved_statement") == APPROVAL
    assert whitepaper.section("summary").startswith("The issuer is responsible for the token.")
    assert whitepaper.section("risk_warning") == RISK_BODY


def test_docx_statement_set_in_bold_passes(tmp_path: Path):
    path = _write_docx(
        tmp_path / "draft.docx",
        [
            ("bold", APPROVAL),
            (
                "bold",
                "The offeror of the crypto-asset is solely responsible for the content of this crypto-asset white paper.",
            ),
            ("large", "Summary"),
            ("plain", SUMMARY_BODY),
        ],
    )
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    statement = next(f for f in lint_whitepaper(whitepaper).findings if f.rule.rule_id == TITLE_II)
    assert statement.status == "pass", statement.issues


def test_pdf_heading_font_ends_the_statement_even_before_a_responsibility_sentence(tmp_path: Path):
    path = _write_pdf(
        tmp_path / "draft.pdf",
        [
            ("body", APPROVAL),
            ("large", "Summary"),
            ("body", "The issuer is responsible for the token. " + SUMMARY_BODY),
            ("large", "Risk Warning"),
            ("body", RISK_BODY),
        ],
    )
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    assert " ".join(whitepaper.section("not_approved_statement").split()) == APPROVAL
    assert whitepaper.section("summary").startswith("The issuer is responsible for the token.")
    _sections_after_incomplete_statement(whitepaper)


def test_pdf_statement_set_in_bold_passes(tmp_path: Path):
    # Each line of the statement is in the bold font, so a short one ("The offeror of the crypto-asset
    # is solely") must not be read as a heading that ends it.
    path = _write_pdf(
        tmp_path / "draft.pdf",
        [
            ("bold", APPROVAL),
            ("bold", "The offeror of the crypto-asset is solely"),
            ("bold", "responsible for the content of this crypto-asset white paper."),
            ("large", "Summary"),
            ("body", SUMMARY_BODY),
        ],
    )
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    statement = next(f for f in lint_whitepaper(whitepaper).findings if f.rule.rule_id == TITLE_II)
    assert statement.status == "pass", statement.issues
    assert whitepaper.section("summary").startswith("The summary gives key information")


@pytest.mark.parametrize(
    "cut",
    [
        "This crypto-asset white paper has not been approved by the",
        "This crypto-asset white paper has not been approved by any compe-",
        "This crypto-asset white paper has not been approved by any competent authority in",
        "This crypto-asset white paper has not been approved by any competent authority,",
    ],
)
@patch("pypdf.PdfReader")
def test_pdf_statement_cut_off_after_a_connective_ends_at_the_heading(mock_reader_cls, tmp_path: Path, cut):
    # The heading the statement stops at is read as a heading, not as the rest of the cut sentence.
    whitepaper = _pdf_whitepaper(
        mock_reader_cls, tmp_path, f"{cut}\nSummary\n{SUMMARY_BODY}\nRisk Warning\n{RISK_BODY}\n"
    )
    assert whitepaper.section("not_approved_statement") == cut
    assert whitepaper.section("summary").startswith("The summary gives key information")
    _sections_after_incomplete_statement(whitepaper)


# A statement field left as a placeholder is reported missing, as any placeholder-only section is,
# also in a German draft, whose real statement is referred to review.
@pytest.mark.parametrize("text", ["TBD", "[\u25cf]", "noch offen"])
@pytest.mark.parametrize(
    ("wp_type", "rule_id"), [(WhitepaperType.OTHER, TITLE_II), (WhitepaperType.EMT, EMT)]
)
def test_placeholder_statement_in_a_german_draft_is_missing(wp_type, rule_id, text):
    finding = _finding(wp_type, rule_id, {"not_approved_statement": text}, language="de")
    assert (finding.status, finding.word_count) == ("missing", 0)
    assert finding.issues == ("Abschnitt enthält nur Platzhalter wie TBD oder TBC.",)


# A statement set apart from the text around it, in bold, capitals or a larger font, carries the
# marks of a heading. Its own lines must not end it, and the sentence after it stays body text.
def _section_findings(whitepaper):
    return {f.rule.rule_id: f for f in lint_whitepaper(whitepaper).findings}


RUN_ON = [
    "This crypto-asset white paper has not been approved by any competent authority in any Member",
    "State of the European Union. The offeror of the crypto-asset is solely responsible for the",
    "content of this crypto-asset white paper. This summary should be read as an introduction to the",
    "Crypto-Asset White Paper and any decision to purchase",
    "Quillmoor Utility Tokens should be based on a consideration of the Crypto-Asset White Paper as a",
    "whole.",
]
GERMAN_RUN_ON = (
    "Dieses Kryptowerte-Whitepaper wurde von keiner zuständigen Behörde eines Mitgliedstaats der "
    "Europäischen Union genehmigt. Der Anbieter des Kryptowerts trägt die alleinige Verantwortung für "
    "den Inhalt dieses Kryptowerte-Whitepapers. Anleger sollten vor dem Kauf die Angaben zum Anbieter "
    "und zu den Risiken lesen, die in diesem Kryptowerte-Whitepaper beschrieben sind."
)


@patch("pypdf.PdfReader")
def test_pdf_sentence_after_a_complete_statement_runs_on_as_body_text(mock_reader_cls, tmp_path: Path):
    text = "\n".join(["Summary", *RUN_ON, *textwrap.wrap(SUMMARY_BODY, 90), "Risk Warning", RISK_BODY])
    whitepaper = _pdf_whitepaper(mock_reader_cls, tmp_path, text + "\n")
    findings = _section_findings(whitepaper)
    assert findings[TITLE_II].status == "pass", findings[TITLE_II].issues
    assert findings["COMMON.SUMMARY"].status == "pass", findings["COMMON.SUMMARY"].issues
    assert whitepaper.section("summary").startswith("Crypto-Asset White Paper and any decision")


def test_pdf_with_fonts_sentence_after_a_complete_statement_runs_on_as_body_text(tmp_path: Path):
    blocks = [("large", "Summary"), *[("plain", line) for line in RUN_ON], ("body", SUMMARY_BODY)]
    path = _write_pdf(tmp_path / "draft.pdf", [*blocks, ("large", "Risk Warning"), ("body", RISK_BODY)])
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    assert _section_findings(whitepaper)["COMMON.SUMMARY"].status == "pass"
    assert whitepaper.section("summary").startswith("Crypto-Asset White Paper and any decision")


@pytest.mark.parametrize("width", [68, 70])
@patch("pypdf.PdfReader")
def test_pdf_german_sentence_after_a_complete_statement_runs_on_as_body_text(
    mock_reader_cls, tmp_path: Path, width
):
    lines = ["Zusammenfassung", *textwrap.wrap(GERMAN_RUN_ON, width), *textwrap.wrap(SUMMARY_BODY, 90)]
    whitepaper = _pdf_whitepaper(
        mock_reader_cls, tmp_path, "\n".join([*lines, "Warnhinweis", RISK_BODY]) + "\n"
    )
    assert _section_findings(whitepaper)["COMMON.SUMMARY"].status != "missing"
    assert "Anleger sollten" in whitepaper.section("not_approved_statement") + whitepaper.section("summary")


FULL_WIDTH = [
    "The offeror of the crypto-asset is solely responsible for the content of this Crypto-Asset",
    "White Paper.",
]
NARROW = [
    "The offeror of the crypto-asset is solely",
    "responsible for the content of this crypto-asset",
    "white paper.",
]


@pytest.mark.parametrize("kind", ["bold", "large"])
@pytest.mark.parametrize(
    ("approval", "responsibility"),
    [
        (
            [
                "This crypto-asset white paper has not been approved by any competent authority in any Member",
                "State of the European Union.",
            ],
            FULL_WIDTH,
        ),
        (
            [
                "This crypto-asset white paper has not been approved",
                "by any competent authority in any Member State of",
                "the European Union.",
            ],
            NARROW,
        ),
    ],
    ids=["full-width", "narrow"],
)
def test_pdf_responsibility_sentence_set_apart_stays_in_the_statement(
    tmp_path: Path, kind, approval, responsibility
):
    blocks = [("plain", line) for line in approval] + [(kind, line) for line in responsibility]
    path = _write_pdf(
        tmp_path / "draft.pdf",
        [
            *blocks,
            ("large", "Summary"),
            ("body", SUMMARY_BODY),
            ("large", "Risk Warning"),
            ("body", RISK_BODY),
        ],
    )
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    statement = _section_findings(whitepaper)[TITLE_II]
    assert statement.status == "pass", statement.issues
    assert whitepaper.section("offeror") == whitepaper.section("crypto_asset") == ""


@pytest.mark.parametrize("header", ["small", "plain"])
def test_pdf_part_header_on_the_next_page_does_not_end_the_statement(tmp_path: Path, monkeypatch, header):
    # A page header that names the current part ("Part A ...") changes from page to page, so it is not
    # a running header; between the statement's two sentences it is no heading.
    monkeypatch.setitem(test_heading_detection.PDF_FONTS, "small", "/F1 8 Tf")
    project = "The ledger service runs for logistics firms in three Member States. " * 12
    path = _write_pdf(
        tmp_path / "draft.pdf",
        [
            ("large", "Information about the Project"),
            ("body", project),
            (
                "plain",
                "This crypto-asset white paper has not been approved by any competent authority in any Member",
            ),
            ("plain", "State of the European Union."),
            ("plain", "Page 1"),
            ("page", ""),
            (header, "Part A Information about the Offeror"),
            (
                "plain",
                "The offeror of the crypto-asset is solely responsible for the content of this crypto-asset",
            ),
            ("plain", "white paper."),
            ("large", "Summary"),
            ("body", SUMMARY_BODY),
            ("plain", "Page 2"),
        ],
    )
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    statement = _section_findings(whitepaper)[TITLE_II]
    assert statement.status == "pass", statement.issues
    assert not whitepaper.section("offeror").startswith("The offeror of the crypto-asset is solely")


STATEMENT_LINES = [
    "This crypto-asset white paper has not been approved by any competent",
    "authority in any Member State of the European Union. The offeror of the",
    "crypto-asset is solely responsible for the content of this crypto-asset",
    "white paper.",
]


@pytest.mark.parametrize(
    ("kind", "lines", "wp_type", "rule_id"),
    [
        ("plain", [line.upper() for line in STATEMENT_LINES], WhitepaperType.OTHER, TITLE_II),
        ("bold", STATEMENT_LINES, WhitepaperType.OTHER, TITLE_II),
        (
            "plain",
            [line.upper().replace("OFFEROR", "ISSUER") for line in STATEMENT_LINES],
            WhitepaperType.EMT,
            EMT,
        ),
    ],
    ids=["capitals", "bold", "capitals-emt"],
)
def test_docx_statement_set_a_paragraph_per_line_passes(tmp_path: Path, kind, lines, wp_type, rule_id):
    blocks = [(kind, line) for line in lines] + [("large", "Summary"), ("plain", SUMMARY_BODY)]
    whitepaper = load_whitepaper(_write_docx(tmp_path / "draft.docx", blocks), wp_type)
    statement = _section_findings(whitepaper)[rule_id]
    assert statement.status == "pass", statement.issues
    assert whitepaper.section("not_approved_statement").splitlines() == lines


RESPONSIBILITY_WITHOUT_STOP = (
    "The offeror of the crypto-asset is solely responsible for the content of this crypto-asset white paper"
)


@pytest.mark.parametrize(
    "blocks",
    [
        [("bold", APPROVAL[:-1]), ("bold", RESPONSIBILITY_WITHOUT_STOP)],
        [("plain", APPROVAL[:-1].upper()), ("plain", RESPONSIBILITY_WITHOUT_STOP.upper())],
    ],
    ids=["bold", "capitals"],
)
def test_docx_responsibility_sentence_without_full_stop_stays_in_the_statement(tmp_path: Path, blocks):
    path = _write_docx(tmp_path / "draft.docx", [*blocks, ("large", "Summary"), ("plain", SUMMARY_BODY)])
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    statement = _section_findings(whitepaper)[TITLE_II]
    assert statement.status == "pass", statement.issues
    assert "solely responsible" in whitepaper.section("not_approved_statement").lower()


def test_pdf_responsibility_sentence_without_full_stop_stays_in_the_statement(tmp_path: Path):
    blocks = [
        ("body", APPROVAL[:-1]),
        ("bold", RESPONSIBILITY_WITHOUT_STOP),
        ("large", "Summary"),
        ("body", SUMMARY_BODY),
    ]
    whitepaper = load_whitepaper(_write_pdf(tmp_path / "draft.pdf", blocks), WhitepaperType.OTHER)
    statement = _section_findings(whitepaper)[TITLE_II]
    assert statement.status == "pass", statement.issues


# The statutory responsibility sentence is recognised by its words, so its lines stay in the statement
# whatever their layout; every other line the loaders mark as a heading still ends the statement.
APPROVAL_LINES = [
    "This crypto-asset white paper has not been approved by any competent authority in any Member",
    "State of the European Union.",
]
OFFEROR_BODY = "Example Ledger S.A. is a public limited company registered in Luxembourg. " * 4


@pytest.mark.parametrize("kind", ["bold", "large"])
@pytest.mark.parametrize("responsibility", [FULL_WIDTH, NARROW], ids=["full-width", "narrow"])
def test_pdf_responsibility_sentence_set_apart_with_wide_line_spacing_stays_in_the_statement(
    tmp_path: Path, kind, responsibility
):
    # A "space" block between the lines puts them about two and a half lines apart, as double spacing does.
    spaced = [block for line in responsibility for block in ((kind, line), ("space", ""))]
    path = _write_pdf(
        tmp_path / "draft.pdf",
        [
            *[("plain", line) for line in APPROVAL_LINES],
            *spaced,
            ("large", "Summary"),
            ("body", SUMMARY_BODY),
        ],
    )
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    statement = _section_findings(whitepaper)[TITLE_II]
    assert statement.status == "pass", statement.issues
    assert whitepaper.section("offeror") == whitepaper.section("crypto_asset") == ""


def test_pdf_heading_between_the_two_sentences_ends_the_statement(tmp_path: Path):
    path = _write_pdf(
        tmp_path / "draft.pdf",
        [
            ("body", APPROVAL),
            ("large", "Summary"),
            ("body", _statement("The offeror").removeprefix(APPROVAL + " ") + " " + SUMMARY_BODY),
        ],
    )
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    assert " ".join(whitepaper.section("not_approved_statement").split()) == APPROVAL
    assert _section_findings(whitepaper)["COMMON.SUMMARY"].status == "pass"


def test_pdf_wrapped_heading_ends_an_incomplete_statement(tmp_path: Path):
    path = _write_pdf(
        tmp_path / "draft.pdf",
        [
            ("body", APPROVAL),
            ("large", "Summary"),
            ("large", "and Key Information"),
            ("body", "The issuer is responsible for the token. " + SUMMARY_BODY),
        ],
    )
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    assert " ".join(whitepaper.section("not_approved_statement").split()) == APPROVAL
    # main reads the lower-case second line of the heading as summary text, so only its place is checked.
    assert "The issuer is responsible for the token." in whitepaper.section("summary")


def test_pdf_part_title_smaller_than_the_statement_ends_it_on_the_same_page(tmp_path: Path, monkeypatch):
    monkeypatch.setitem(test_heading_detection.PDF_FONTS, "plain11", "/F1 11 Tf")
    path = _write_pdf(
        tmp_path / "draft.pdf",
        [
            *[("plain11", line) for line in APPROVAL_LINES],
            ("bold", "Part A Information about the Offeror"),
            ("body", "The offeror is responsible for the offer. " + OFFEROR_BODY),
            ("large", "Summary"),
            ("body", SUMMARY_BODY),
        ],
    )
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    assert " ".join(whitepaper.section("not_approved_statement").split()) == APPROVAL
    assert _section_findings(whitepaper)["ANNEX_I.A"].status != "missing"


@pytest.mark.parametrize("repeated", [True, False], ids=["page-header", "heading"])
def test_pdf_bold_line_at_the_top_of_the_next_page(tmp_path: Path, repeated):
    # A bold page header that repeats the heading its page opens does not end the statement; a bold
    # heading that opens the page does.
    page_two = [("bold", "Part A Information about the Offeror")]
    if repeated:
        page_two += [
            (
                "plain",
                "The offeror of the crypto-asset is solely responsible for the content of this crypto-asset",
            ),
            ("plain", "white paper."),
            ("large", "Part A Information about the Offeror"),
        ]
    else:
        page_two += [("plain", "The offeror is responsible for the offer.")]
    blocks = [
        ("large", "Summary"),
        ("body", SUMMARY_BODY),
        *[("plain", line) for line in APPROVAL_LINES],
        ("plain", "Page 1"),
        ("page", ""),
        *page_two,
        ("body", OFFEROR_BODY),
        ("plain", "Page 2"),
    ]
    whitepaper = load_whitepaper(_write_pdf(tmp_path / "draft.pdf", blocks), WhitepaperType.OTHER)
    findings = _section_findings(whitepaper)
    assert findings[TITLE_II].status == ("pass" if repeated else "review"), findings[TITLE_II].issues
    assert findings["ANNEX_I.A"].status != "missing"


@pytest.mark.parametrize("kind", ["subheading", "outline"])
def test_docx_heading_in_lower_case_ends_an_incomplete_statement(tmp_path: Path, kind):
    path = _write_docx(
        tmp_path / "draft.docx",
        [
            ("plain", APPROVAL),
            (kind, "a) Summary"),
            ("plain", "The issuer is responsible for the token. " + SUMMARY_BODY),
        ],
    )
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    assert whitepaper.section("not_approved_statement") == APPROVAL
    assert whitepaper.section("summary").startswith("The issuer is responsible for the token.")


def test_docx_long_heading_naming_the_persons_responsible_ends_an_incomplete_statement(tmp_path: Path):
    heading = "Information about the offeror and the persons responsible for drawing up this white paper"
    path = _write_docx(
        tmp_path / "draft.docx",
        [("plain", APPROVAL), ("heading", heading), ("plain", "The offeror is responsible. " + OFFEROR_BODY)],
    )
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    assert whitepaper.section("not_approved_statement") == APPROVAL
    assert _section_findings(whitepaper)["ANNEX_I.A"].status != "missing"


@pytest.mark.parametrize(
    ("kind", "responsibility"),
    [
        (
            "bold",
            [
                "The offeror of the crypto-asset is solely responsible for",
                "the content of this crypto-asset",
                "white paper.",
            ],
        ),
        (
            "plain",
            [
                "THE OFFEROR OF THE CRYPTO-ASSET IS SOLELY",
                "RESPONSIBLE FOR THE CONTENT OF THIS CRYPTO-ASSET",
                "WHITE PAPER.",
            ],
        ),
    ],
    ids=["bold", "capitals"],
)
def test_docx_responsibility_sentence_set_apart_a_paragraph_per_line_passes(
    tmp_path: Path, kind, responsibility
):
    blocks = [("plain", line) for line in APPROVAL_LINES] + [(kind, line) for line in responsibility]
    path = _write_docx(tmp_path / "draft.docx", [*blocks, ("large", "Summary"), ("plain", SUMMARY_BODY)])
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    statement = _section_findings(whitepaper)[TITLE_II]
    assert statement.status == "pass", statement.issues


# The end of the statement: a full stop inside closing quotes, emphasis or before a note mark ends the
# sentence, a numbered heading after the statement stays a heading, and a part title that names the
# offeror is not the person the responsibility sentence names.
STATUTORY = _statement("The offeror")
SUMMARY_TEXT = "This summary should be read as an introduction to the crypto-asset white paper. " * 3
RISK_TEXT = "The crypto-asset may lose its value in part or in full and may not be transferable."


def _markdown(tmp_path: Path, statement: str, heading: str = "## 1. Summary") -> Whitepaper:
    path = tmp_path / "draft.md"
    path.write_text(
        f"# Example Token White Paper\n\n{statement}\n\n{heading}\n\n{SUMMARY_TEXT}\n\n## 2. Risk Warning\n\n{RISK_TEXT}\n",
        encoding="utf-8",
    )
    return load_whitepaper(path, WhitepaperType.OTHER)


@pytest.mark.parametrize(
    "statement",
    [f"**{STATUTORY}**", f"*{STATUTORY}*", f"\u201c{STATUTORY}\u201d", f"{STATUTORY}[^1]", STATUTORY[:-1]],
    ids=["bold", "italic", "quoted", "footnote", "no-final-stop"],
)
@pytest.mark.parametrize("heading", ["## 1. Summary", "## A. Summary", "## I. Summary"])
def test_markdown_numbered_heading_after_the_statement_stays_a_heading(tmp_path: Path, statement, heading):
    whitepaper = _markdown(tmp_path, statement, heading)
    findings = _section_findings(whitepaper)
    assert findings[TITLE_II].status == "pass", findings[TITLE_II].issues
    assert "Summary" not in whitepaper.section("not_approved_statement")
    assert whitepaper.section("summary").startswith("This summary should be read")


def test_markdown_hard_wrapped_bold_statement_before_a_numbered_heading(tmp_path: Path):
    whitepaper = _markdown(tmp_path, "\n".join(textwrap.wrap(f"**{STATUTORY}**", 90)))
    assert _section_findings(whitepaper)[TITLE_II].status == "pass"
    assert whitepaper.section("summary").startswith("This summary should be read")


@pytest.mark.parametrize(
    "statement", [f"\u201c{STATUTORY}\u201d", STATUTORY[:-1]], ids=["quoted", "no-final-stop"]
)
def test_docx_numbered_heading_after_the_statement_stays_a_heading(tmp_path: Path, statement):
    path = _write_docx(
        tmp_path / "draft.docx", [("plain", statement), ("heading", "1. Summary"), ("plain", SUMMARY_TEXT)]
    )
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    assert _section_findings(whitepaper)[TITLE_II].status == "pass"
    assert whitepaper.section("summary").startswith("This summary should be read")


def test_pdf_quoted_statement_before_a_heading_passes_whole(tmp_path: Path):
    path = _write_pdf(
        tmp_path / "draft.pdf",
        [
            ("body", f"\u201c{STATUTORY}\u201d"),
            ("large", "Summary"),
            ("body", SUMMARY_BODY),
            ("large", "Risk Warning"),
            ("body", RISK_BODY),
        ],
    )
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    findings = _section_findings(whitepaper)
    assert findings[TITLE_II].status == "pass", findings[TITLE_II].issues
    assert whitepaper.section("not_approved_statement").endswith("white paper.\u201d")
    assert findings["COMMON.SUMMARY"].status == "pass"


@patch("pypdf.PdfReader")
def test_pdf_statement_with_note_marks_before_a_numbered_heading(mock_reader_cls, tmp_path: Path):
    marked_statement = STATUTORY.replace("European Union.", "European Union.1") + "2"
    lines = [*textwrap.wrap(marked_statement, 90), "1. Summary", *textwrap.wrap(SUMMARY_BODY, 90)]
    whitepaper = _pdf_whitepaper(mock_reader_cls, tmp_path, "\n".join(lines) + "\n")
    assert _section_findings(whitepaper)[TITLE_II].status == "pass"
    assert "1. Summary" not in whitepaper.section("not_approved_statement")


PART_A = "Part A \u2013 Information about the offeror or the person seeking admission to trading"
PART_A_FIELDS = "A.1 Name: Example Ledger AG. A.2 Legal form: public limited company. " * 4


@pytest.mark.parametrize(
    "sentence",
    [
        "Example Ledger AG is solely responsible for the content of this crypto-asset white paper.",
        "We are solely responsible for the content of this crypto-asset white paper.",
    ],
    ids=["company", "we"],
)
def test_part_title_naming_the_offeror_ends_an_incomplete_statement(tmp_path: Path, sentence):
    path = tmp_path / "draft.md"
    path.write_text(f"{APPROVAL}\n\n## {PART_A}\n\n{sentence} {PART_A_FIELDS}\n", encoding="utf-8")
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    assert whitepaper.section("not_approved_statement") == APPROVAL
    assert _section_findings(whitepaper)["ANNEX_I.A"].status != "missing"

    docx_path = _write_docx(
        tmp_path / "draft.docx",
        [("plain", APPROVAL), ("heading", PART_A), ("plain", f"{sentence} {PART_A_FIELDS}")],
    )
    whitepaper = load_whitepaper(docx_path, WhitepaperType.OTHER)
    assert whitepaper.section("not_approved_statement") == APPROVAL
    assert _section_findings(whitepaper)["ANNEX_I.A"].status != "missing"


def test_docx_capitals_responsibility_sentence_split_before_solely_passes(tmp_path: Path):
    lines = [
        "THE OFFEROR OF THE CRYPTO-ASSET IS",
        "SOLELY RESPONSIBLE FOR THE CONTENT OF THIS",
        "CRYPTO-ASSET WHITE PAPER.",
    ]
    blocks = [("plain", line) for line in APPROVAL_LINES] + [("plain", line) for line in lines]
    path = _write_docx(tmp_path / "draft.docx", [*blocks, ("large", "Summary"), ("plain", SUMMARY_BODY)])
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    assert _section_findings(whitepaper)[TITLE_II].status == "pass"


# A responsibility sentence with every statutory word ends the statement with its last word, whatever
# follows: no stop, a note mark, a heading in any form, body text or a run-on.
RESPONSIBILITY = (
    "The offeror of the crypto-asset is solely responsible for the content of this crypto-asset white paper."
)


@pytest.mark.parametrize(
    "tail", ["", ".\u2020", ".[1]", ".(1)"], ids=["no-stop", "dagger", "bracket", "paren"]
)
@pytest.mark.parametrize(
    ("wp_type", "rule_id", "subject"),
    [(WhitepaperType.OTHER, TITLE_II, "offeror"), (WhitepaperType.EMT, EMT, "issuer")],
)
def test_pdf_tail_on_its_own_line_before_a_heading_stays_in_the_statement(
    tmp_path: Path, tail, wp_type, rule_id, subject
):
    path = _write_pdf(
        tmp_path / "draft.pdf",
        [
            (
                "plain",
                "This crypto-asset white paper has not been approved by any competent authority in any Member State",
            ),
            (
                "plain",
                f"of the European Union. The {subject} of the crypto-asset is solely responsible for the content of this",
            ),
            ("plain", f"crypto-asset white paper{tail}"),
            ("large", "Summary"),
            ("body", SUMMARY_BODY),
        ],
    )
    whitepaper = load_whitepaper(path, wp_type)
    findings = _section_findings(whitepaper)
    assert findings[rule_id].status == "pass", findings[rule_id].issues
    assert whitepaper.section("not_approved_statement").endswith(f"crypto-asset white paper{tail}")
    assert findings["COMMON.SUMMARY"].status != "missing"


@patch("pypdf.PdfReader")
def test_pdf_note_marked_statement_without_final_stop_before_a_heading(mock_reader_cls, tmp_path: Path):
    statement = RESPONSIBILITY[:-1].join([APPROVAL.replace("Union.", "Union.1") + " ", ""])
    lines = [*textwrap.wrap(statement, 90), "Summary", *textwrap.wrap(SUMMARY_BODY, 90)]
    whitepaper = _pdf_whitepaper(mock_reader_cls, tmp_path, "\n".join(lines) + "\n")
    findings = _section_findings(whitepaper)
    assert findings[TITLE_II].status == "pass", findings[TITLE_II].issues
    assert findings["COMMON.SUMMARY"].status == "pass"


@pytest.mark.parametrize("heading", ["## A.1 Name of the offeror", "## II.3 Summary"])
def test_markdown_field_number_after_a_statement_without_stop_stays_a_heading(tmp_path: Path, heading):
    path = tmp_path / "draft.md"
    body = "Example Ledger AG is a public limited company registered in Luxembourg. " * 3
    path.write_text(f"{APPROVAL[:-1]}\n\n{heading}\n\n{body}\n", encoding="utf-8")
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    assert heading.removeprefix("## ").split()[0] not in whitepaper.section("not_approved_statement")
    assert whitepaper.section("offeror" if "A.1" in heading else "summary").startswith("Example Ledger AG")


@pytest.mark.parametrize("mark", ["**", "*", "\u201c"], ids=["bold", "italic", "quoted"])
@pytest.mark.parametrize(
    "after",
    [
        "## Summary\n\n" + SUMMARY_TEXT,
        "**Part A \u2013 Information about the offeror**\n\n" + OFFEROR_BODY,
        SUMMARY_TEXT,
    ],
    ids=["heading", "label", "body"],
)
def test_markdown_two_paragraph_statement_without_final_stop(tmp_path: Path, mark, after):
    close = "\u201d" if mark == "\u201c" else mark
    path = tmp_path / "draft.md"
    path.write_text(
        f"{mark}{APPROVAL}{close}\n\n{mark}{RESPONSIBILITY[:-1]}{close}\n\n{after}\n", encoding="utf-8"
    )
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    statement = _section_findings(whitepaper)[TITLE_II]
    assert statement.status == "pass", statement.issues
    assert "Summary" not in whitepaper.section("not_approved_statement")
    if after.startswith("## Summary"):
        assert whitepaper.section("summary").startswith("This summary should be read")


def test_markdown_responsibility_sentence_running_on_past_the_statutory_words(tmp_path: Path):
    run_on = (
        RESPONSIBILITY[:-1]
        + " and any supplements thereto, as notified to the competent authority of the home Member State."
    )
    path = tmp_path / "draft.md"
    path.write_text(f"**{APPROVAL} {run_on}**\n\n## Summary\n\n{SUMMARY_TEXT}\n", encoding="utf-8")
    whitepaper = load_whitepaper(path, WhitepaperType.OTHER)
    assert _section_findings(whitepaper)[TITLE_II].status == "pass"
    assert whitepaper.section("summary").startswith("This summary should be read")


@pytest.mark.parametrize(
    ("kind", "lines", "wp_type", "rule_id"),
    [
        (
            "bold",
            [
                "The person seeking admission to trading of the crypto-",
                "asset is solely responsible for the content of this",
                "crypto-asset white paper.",
            ],
            WhitepaperType.OTHER,
            TITLE_II,
        ),
        (
            "plain",
            [
                "THE OFFEROR OF THE CRYPTO-",
                "ASSET IS SOLELY RESPONSIBLE",
                "FOR THE CONTENT OF THIS",
                "CRYPTO-ASSET WHITE PAPER.",
            ],
            WhitepaperType.OTHER,
            TITLE_II,
        ),
        (
            "bold",
            [
                "The issuer of the crypto-",
                "asset is solely responsible for the",
                "content of this crypto-asset white paper.",
            ],
            WhitepaperType.EMT,
            EMT,
        ),
        (
            "plain",
            [
                "The offeror of the crypto-asset is solely",
                "responsible for the content of this",
                "crypto-asset white paper",
            ],
            WhitepaperType.OTHER,
            TITLE_II,
        ),
        (
            "bold",
            [
                "The offeror of the crypto-asset is solely",
                "responsible for the content of this",
                "crypto-asset white paper",
            ],
            WhitepaperType.OTHER,
            TITLE_II,
        ),
    ],
    ids=["hyphen-bold", "hyphen-capitals", "hyphen-emt", "no-stop-plain", "no-stop-bold"],
)
def test_docx_responsibility_sentence_a_paragraph_per_line_ends_with_its_last_line(
    tmp_path: Path, kind, lines, wp_type, rule_id
):
    blocks = [("plain", line) for line in APPROVAL_LINES] + [(kind, line) for line in lines]
    path = _write_docx(tmp_path / "draft.docx", [*blocks, ("bold", "Summary"), ("plain", SUMMARY_BODY)])
    whitepaper = load_whitepaper(path, wp_type)
    statement = _section_findings(whitepaper)[rule_id]
    assert statement.status == "pass", statement.issues
    assert whitepaper.section("not_approved_statement").splitlines()[-1] == lines[-1]
    assert whitepaper.section("summary").startswith("The summary gives key information")
